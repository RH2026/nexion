"""
JYPESA | REPORTES // REPORTEADOR
Ubicación: REPORTES/REPORTEADOR  (archivo único: reporteador.py)

Dos pestañas en una sola página:

  1. REPORTES DINÁMICOS   (antes "Consultas")
     Filtras (año, mes, cobro regreso/destino, fletera, transporte, destino, cliente, concepto, cajas, entrega,
     semáforo, % logístico, búsqueda), armas tu reporte marcando bloques, lo ves en pantalla y lo imprimes en PDF o Excel.

  2. REPORTES PREDEFINIDOS  (antes "Reporte Mensual PDF")
     Reporte mensual completo para Dirección, % logístico por concepto, salud de pedidos pequeños y
     detalle pedido por pedido (PDF / Excel).

Todos los botones usan el color de Consultas (#628290).
Cada pestaña conserva su propia lógica; los nombres que chocaban entre ambos archivos se aislaron con el prefijo "mp_"
en la parte de reportes predefinidos.
"""
import hashlib
import html as _html
import io
import json
import os
import time
import unicodedata
from datetime import datetime
from io import BytesIO

import numpy as np
import pandas as pd
import requests
import streamlit as st

from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.shapes import Circle, Drawing, Line, Rect, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import (BaseDocTemplate, CondPageBreak, Frame, KeepTogether, NextPageTemplate, PageBreak,
                                PageTemplate, Paragraph, Spacer, Table, TableStyle)

from components.layout import render_layout

try:
    from muestras_common import obtener_datos_github, precios as PRECIOS_MUESTRAS
    MUESTRAS_OK, MUESTRAS_ERR = True, ""
except Exception as _e_m:   # el reporte se genera igual, sin la sección de muestras
    obtener_datos_github, PRECIOS_MUESTRAS = None, {}
    MUESTRAS_OK, MUESTRAS_ERR = False, str(_e_m)


# ============================================================
# 0. PÁGINA Y LAYOUT (una sola vez para todo el reporteador)
# ============================================================
st.set_page_config(page_title="JYPESA | Reportes - Reporteador", layout="wide", initial_sidebar_state="collapsed")
render_layout(modulo_actual="REPORTES", submodulo_actual="REPORTEADOR")




# ============================================================
# PESTAÑA 1 · REPORTES DINÁMICOS
# ============================================================

# ============================================================
# 2. CONSTANTES
# ============================================================
GITHUB_USER, GITHUB_REPO, BRANCH = "RH2026", "nexion", "main"
ARCHIVO_MATRIZ = "Matriz_Excel_Dashboard.csv"
ARCHIVO_LOGO = "n1.png"

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO",
         "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
TARGET = 7.5                    # % meta de costo logístico
EXTRA_REGEX = "RECOLECCI|MANIOBRA"
CLASES = ["SALUDABLE", "EN ALERTA", "CRÍTICO", "PÉRDIDA", "SIN FACTURACIÓN", "SIN DATOS"]
CLASES_FUERA = ["EN ALERTA", "CRÍTICO", "PÉRDIDA", "SIN FACTURACIÓN"]
ESTATUS = ["ENTREGADO A TIEMPO", "ENTREGADO CON RETRASO", "EN TRÁNSITO EN TIEMPO", "EN TRÁNSITO RETRASADO", "EN TRÁNSITO"]
SCREEN_MAX, PDF_MAX = 1500, 1200   # filas máximas por tabla en pantalla / PDF
TOPS = [5, 10, 15, 20, 30, 50, 100, "TODOS"]

HEX = dict(teal="#00A3A3", green="#2E9E6B", red="#D64545", blue="#3B82F6", gold="#FFC000", purple="#7C5CBF",
           orange="#E8833A", gray="#6B7A86", slate="#2B343B")
C_NAVY, C_SLATE, C_TEAL, C_GOLD = (colors.HexColor(x) for x in ("#384A52", "#2B343B", "#00A3A3", "#FFC000"))
C_GREEN, C_RED, C_GRAY = colors.HexColor("#2E9E6B"), colors.HexColor("#D64545"), colors.HexColor("#6B7A86")
C_LIGHT, C_BORDER, C_TEXT = colors.HexColor("#F3F6F8"), colors.HexColor("#D5DDE2"), colors.HexColor("#1F2D35")
C_AMBER = colors.HexColor("#C98A00")
TONO_COLOR = {"good": C_GREEN, "bad": C_RED, "warn": C_AMBER, "mute": C_GRAY, "": C_TEXT}

# Dimensiones para agrupar / cruzar: nombre -> columna interna
DIMS = {"Fletera": "_FLE", "Transporte": "_TR", "Destino": "_DE", "Cliente": "_CLI", "Modalidad (cobro)": "_MOD",
        "Concepto": "_CONC", "Mes": "_MESN", "Nº de cajas": "_CAJB", "Semáforo": "_CLASE",
        "Estatus de entrega": "_EST", "Año": "_ANIOT"}
NATURALES = {"_MESN", "_CAJB", "_CLASE", "_ANIOT"}      # dimensiones con orden propio (no por métrica)

# Métricas por las que se puede ordenar / graficar: etiqueta -> (clave, tipo de formato)
METRICAS = {"Costo logístico ($)": ("costo", "money"), "% logístico": ("pct_log", "pct2"),
            "Facturación": ("fact", "money"), "Pedidos": ("n", "int"), "Cajas": ("cajas", "int"),
            "Costo por caja": ("costo_caja", "money2"), "Exceso sobre target ($)": ("exceso", "money")}

SECCIONES = [  # (clave, nombre, descripción corta)
    ("kpi", "Tarjetas resumen", "Pedidos, facturación, costo, % logístico, extras, entregas"),
    ("peq", "Pequeños vs grandes", "Comparativo 1 a N cajas contra más de N (ignora el filtro de cajas)"),
    ("flet", "Por fletera", "Costo, % logístico y exceso por fletera"),
    ("tran", "Por transporte", "Por tipo de transporte"),
    ("dest", "Por destino", "Por destino"),
    ("clie", "Por cliente", "Por cliente"),
    ("moda", "Cobro regreso vs cobro destino", "Por modalidad de cobro"),
    ("conc", "Por concepto", "Por concepto del envío"),
    ("mes", "Por mes (tendencia)", "Mes a mes"),
    ("caja", "Por número de cajas", "1, 2, 3, 4, 5 o más…"),
    ("sema", "Semáforo de salud", "Saludable / alerta / crítico / pérdida"),
    ("rank", "Pedido por pedido", "Lista ordenada de mayor a menor costo logístico (o lo que elijas)"),
    ("entr", "Efectividad de entrega", "A tiempo, retrasos, días promedio e incidencias por fletera"),
    ("extr", "Costos extras", "Recolecciones y maniobras (no son pedidos)"),
    ("mues", "Costos de muestras", "Envíos de muestra: producto, flete e inversión"),
    ("cruz", "Tabla cruzada", "Elige filas, columnas y métrica"),
]
PRESETS = {"Ejecutivo": ["kpi", "peq", "flet", "mes", "sema", "extr"],
           "Costos": ["kpi", "flet", "tran", "dest", "clie", "rank", "extr"],
           "Todo": [k for k, _, _ in SECCIONES]}


# ============================================================
# 3. UTILIDADES
# ============================================================
def esc(t):
    return _html.escape("" if t is None else str(t))


def trunc(t, n):
    t = "" if t is None else str(t)
    return t if len(t) <= n else t[: n - 1] + "."


def sin_acentos(t):
    return "".join(c for c in unicodedata.normalize("NFD", str(t)) if unicodedata.category(c) != "Mn")


def isnum(v):
    try:
        return v is not None and not pd.isna(v)
    except Exception:
        return False


def parse_fecha_segura(serie):
    serie = pd.Series(serie)
    if pd.api.types.is_datetime64_any_dtype(serie):
        return serie
    s_txt = serie.fillna("").astype(str).str.strip()
    es_iso = s_txt.str.match(r"^\d{4}-\d{1,2}-\d{1,2}")

    def _conv(x, **kw):
        try:
            return pd.to_datetime(x, errors="coerce", format="mixed", **kw)
        except (TypeError, ValueError):
            return pd.to_datetime(x, errors="coerce", **kw)
    out = pd.Series(pd.NaT, index=s_txt.index, dtype="datetime64[ns]")
    if es_iso.any():
        out.loc[es_iso] = _conv(s_txt[es_iso]).astype("datetime64[ns]")
    if (~es_iso).any():
        out.loc[~es_iso] = _conv(s_txt[~es_iso], dayfirst=True).astype("datetime64[ns]")
    return out


def limpiar_moneda(serie):
    return pd.to_numeric(serie.astype(str).str.replace(r"[^\d\.\-]", "", regex=True).replace("", "0"),
                         errors="coerce").fillna(0.0)


def fmt(kind, v):
    """Formato único para pantalla y PDF. Devuelve (texto, tono) con tono en '', good, bad, warn, mute."""
    if kind in ("txt", "cls", "est"):
        s = "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v)
        tono = ""
        if kind == "cls":
            tono = {"SALUDABLE": "good", "EN ALERTA": "warn", "CRÍTICO": "bad", "PÉRDIDA": "bad"}.get(s, "mute")
        elif kind == "est":
            tono = "good" if s == "ENTREGADO A TIEMPO" else ("bad" if "RETRASO" in s or "RETRASADO" in s else "")
        return s, tono
    if kind == "date":
        return ("-", "mute") if not isnum(v) else (pd.Timestamp(v).strftime("%d/%m/%Y"), "")
    if not isnum(v):
        return "-", "mute"
    v = float(v)
    if kind == "int":
        return f"{v:,.0f}", ""
    if kind == "num1":
        return f"{v:,.1f}", ""
    if kind == "money":
        return f"${v:,.0f}", ""
    if kind == "money2":
        return f"${v:,.2f}", ""
    if kind == "pct":
        return f"{v:.1f}%", ""
    if kind == "pct2":
        return f"{v:.2f}%", ""
    if kind == "pctlog":
        return f"{v:.2f}%", ("good" if v <= TARGET else "bad")
    if kind == "vs":
        return f"{v:+.2f} pp", ("good" if v <= 0 else "bad")
    if kind == "pctok":   # efectividad: más es mejor
        return f"{v:.1f}%", ("good" if v >= 90 else ("warn" if v >= 80 else "bad"))
    return str(v), ""


# ============================================================
# 4. CARGA Y PREPARACIÓN (mismas reglas que reporte_mensual_pdf.py)
# ============================================================
def _headers_github():
    try:
        token = st.secrets.get("GITHUB_TOKEN", None)
    except Exception:
        token = None
    return {"Authorization": f"token {token}"} if token else {}


@st.cache_data(ttl=300)
def cargar_matriz():
    url = f"https://raw.githubusercontent.com/{GITHUB_USER}/{GITHUB_REPO}/{BRANCH}/{ARCHIVO_MATRIZ}?v={int(time.time())}"
    try:
        r = requests.get(url, headers=_headers_github(), timeout=40)
        r.raise_for_status()
        df = pd.read_csv(BytesIO(r.content), encoding="utf-8-sig")
    except Exception:
        df = pd.read_csv(url, encoding="utf-8-sig")
    df.columns = [str(c).strip() for c in df.columns]
    return df


@st.cache_data(ttl=3600)
def obtener_logo_bytes():
    try:
        import os
        if os.path.exists(ARCHIVO_LOGO):
            with open(ARCHIVO_LOGO, "rb") as f:
                return f.read()
        url = f"https://raw.githubusercontent.com/{GITHUB_USER}/{GITHUB_REPO}/{BRANCH}/{ARCHIVO_LOGO}"
        r = requests.get(url, headers=_headers_github(), timeout=20)
        return r.content if r.status_code == 200 else None
    except Exception:
        return None


def preparar_base(df_raw):
    df = df_raw.copy()
    df.columns = [str(c).strip() for c in df.columns]
    for c in ["FECHA DE ENVÍO", "PROMESA DE ENTREGA", "FECHA DE ENTREGA REAL", "EMISION"]:
        df[c] = parse_fecha_segura(df[c]) if c in df.columns else pd.Series(pd.NaT, index=df.index, dtype="datetime64[ns]")
    for c in ["FLETERA", "FORMA DE ENVIO", "TRANSPORTE", "DESTINO", "NOMBRE DEL CLIENTE", "INCIDENCIAS", "MES", "NÚMERO DE PEDIDO"]:
        if c not in df.columns:
            df[c] = ""
        df[c] = df[c].fillna("").astype(str).str.strip()
        df.loc[df[c].str.lower() == "nan", c] = ""
    df["MES"] = df["MES"].str.upper()
    for c in ["COSTO DE LA GUÍA", "FACTURACION", "VALUACION", "COSTOS ADICIONALES"]:
        df[c] = limpiar_moneda(df[c]) if c in df.columns else 0.0
    df["CAJAS"] = pd.to_numeric(df["CAJAS"], errors="coerce").fillna(0) if "CAJAS" in df.columns else 0.0
    col_c = next((c for c in df.columns if "CONCEPTO" in str(c).upper()), None)
    if col_c:
        conc = df[col_c].fillna("").astype(str).map(sin_acentos).str.strip().str.upper()
        df["_CONC"] = conc.replace("", "SIN CONCEPTO")
        df["ES_EXTRA"] = conc.str.contains(EXTRA_REGEX, regex=True).astype(bool)
    else:
        df["_CONC"], df["ES_EXTRA"] = "SIN CONCEPTO", False
    return df


@st.cache_data(ttl=300)
def datos_base():
    return preparar_base(cargar_matriz())


def _col_factura(df):
    exactos = {"FACTURA", "NUMERO DE FACTURA", "NUMERO FACTURA", "NO FACTURA", "NO. FACTURA", "NUM FACTURA",
               "NUM. FACTURA", "FOLIO FACTURA", "FOLIO DE FACTURA", "N FACTURA"}
    for c in df.columns:
        if sin_acentos(c).strip().upper() in exactos:
            return c
    return "NÚMERO DE PEDIDO"


def preparar_universo(df, incluir_adic, hoy, max_peq):
    """Agrega a TODA la matriz las columnas de análisis (un renglón = un envío). Vectorizado."""
    u = df.copy()
    u["_CAJ"] = u["CAJAS"].astype(float).round()  # solo columna CAJAS (no usar CANTIDAD DE CAJAS)
    u["_COSTO"] = u["COSTO DE LA GUÍA"] + (u["COSTOS ADICIONALES"] if incluir_adic else 0.0)
    fact, costo = u["FACTURACION"], u["_COSTO"]
    u["_PCT"] = costo / fact.where(fact > 0) * 100
    conds = [(fact <= 0) & (costo <= 0), fact <= 0, u["_PCT"] >= 100, u["_PCT"] > 2 * TARGET, u["_PCT"] > TARGET]
    u["_CLASE"] = np.select(conds, ["SIN DATOS", "SIN FACTURACIÓN", "PÉRDIDA", "CRÍTICO", "EN ALERTA"], default="SALUDABLE")
    u["_EXC"] = (costo - TARGET / 100 * fact).clip(lower=0)
    u["_FUERA"] = u["_CLASE"].isin(CLASES_FUERA)
    u["_CPC"] = costo / u["_CAJ"].where(u["_CAJ"] > 0)
    u["_MOD"] = u["FORMA DE ENVIO"].str.upper().map(
        lambda x: "COBRO REGRESO" if "REGRESO" in x else ("COBRO DESTINO" if "DESTINO" in x else (x or "SIN FORMA")))
    u["_DE"] = u["DESTINO"].str.upper().replace("", "SIN DESTINO")
    u["_TR"] = u["TRANSPORTE"].str.upper().replace("", "SIN TRANSPORTE")
    u["_FLE"] = u["FLETERA"].str.upper().replace("", "SIN ASIGNAR")
    u["_CLI"] = u["NOMBRE DEL CLIENTE"].replace("", "SIN CLIENTE")
    col_f = _col_factura(u)
    u["_FAC"] = u[col_f].fillna("").astype(str).str.strip().replace("nan", "")
    # cajas en etiquetas con orden natural
    cap = max_peq + 1
    n = np.minimum(u["_CAJ"], cap).astype(int)
    u["_CAJN"] = n
    u["_CAJB"] = np.where(n < 1, "SIN CAJAS", np.where(n >= cap, f"{cap}+ CAJAS", n.astype(str) + np.where(n == 1, " CAJA", " CAJAS")))
    # fechas / mes / año
    u["_ANIO"] = u["FECHA DE ENVÍO"].dt.year.fillna(0).astype(int)
    u["_ANIOT"] = np.where(u["_ANIO"] > 0, u["_ANIO"].astype(str), "S/F")
    u["_MESN"] = u["MES"].replace("", "SIN MES").str.title()
    # entrega
    ent = u["FECHA DE ENTREGA REAL"].notna()
    prom = u["PROMESA DE ENTREGA"]
    u["_ENT"] = ent
    u["_ATIEMPO"] = ent & (prom.isna() | (u["FECHA DE ENTREGA REAL"] <= prom))
    u["_RETR_TR"] = (~ent) & prom.notna() & (prom < hoy)
    u["_EST"] = np.select(
        [ent & u["_ATIEMPO"], ent, (~ent) & prom.notna() & (prom >= hoy), u["_RETR_TR"]],
        ESTATUS[:4], default="EN TRÁNSITO")
    d_ent = (u["FECHA DE ENTREGA REAL"] - u["FECHA DE ENVÍO"]).dt.days
    u["_DIAS"] = d_ent.where(ent & (d_ent >= 0))
    u["_INC"] = ~u["INCIDENCIAS"].astype(str).str.strip().str.upper().isin(["", "OK"])
    return u


# ============================================================
# 5. FILTROS
# ============================================================
def aplicar_filtros(u, f):
    """u: universo preparado. f: dict de filtros. Devuelve (d, ext, d_sin_cajas)."""
    ok_anio = u["FECHA DE ENVÍO"].isna() | u["_ANIO"].isin(f["anios"])
    ok_mes = u["MES"].isin([MESES[m - 1] for m in f["meses"]]) if f["meses"] else (u["MES"] != "")
    m = ok_anio & ok_mes
    if f["modalidad"] != "TODAS":
        m &= u["_MOD"] == f["modalidad"]
    for k, col in (("fleteras", "_FLE"), ("transportes", "_TR"), ("destinos", "_DE"), ("conceptos", "_CONC")):
        if f[k]:
            m &= u[col].isin(f[k])
    ext = u[m & u["ES_EXTRA"]].copy()

    p = m & ~u["ES_EXTRA"]
    if f["clientes"]:
        p &= u["_CLI"].isin(f["clientes"])
    if f["estatus"]:
        p &= u["_EST"].isin(f["estatus"])
    if f["clases"]:
        p &= u["_CLASE"].isin(f["clases"])
    if f["pct_min"] and f["pct_min"] > 0:
        p &= u["_PCT"] > f["pct_min"]
    if f["solo_fuera"]:
        p &= u["_FUERA"]
    if f["buscar"]:
        b = sin_acentos(f["buscar"]).upper()
        txt = (u["_FAC"] + " " + u["_CLI"] + " " + u["NÚMERO DE PEDIDO"]).map(lambda x: sin_acentos(x).upper())
        p &= txt.str.contains(b, regex=False)
    d0 = u[p].copy()

    modo, n = f["cajas_modo"], f["max_peq"]
    if modo == "PEQUEÑOS":
        mc = (d0["_CAJ"] >= 1) & (d0["_CAJ"] <= n)
    elif modo == "GRANDES":
        mc = d0["_CAJ"] > n
    elif modo == "RANGO":
        mc = d0["_CAJ"] >= f["caj_min"]
        if f["caj_max"] > 0:
            mc &= d0["_CAJ"] <= f["caj_max"]
    else:
        mc = pd.Series(True, index=d0.index)
    return d0[mc].copy(), ext, d0


def texto_filtros(f):
    per = ("AÑO " + ", ".join(str(a) for a in f["anios"])) if not f["meses"] else \
        f"{', '.join(MESES[m - 1].title() for m in f['meses'])} {', '.join(str(a) for a in f['anios'])}"
    if not f["meses"]:
        per += " (acumulado)"

    def lst(x):
        return ", ".join(x) if len(x) <= 3 else f"{len(x)} seleccionados"
    p = [per, "Todas las modalidades" if f["modalidad"] == "TODAS" else f["modalidad"].title()]
    n = f["max_peq"]
    cj = {"PEQUEÑOS": f"1 a {n} cajas", "GRANDES": f"{n + 1} o más cajas",
          "RANGO": f"{f['caj_min']} a {f['caj_max'] or 'más'} cajas"}.get(f["cajas_modo"])
    if cj:
        p.append(cj)
    for lab, k in (("Fletera", "fleteras"), ("Transporte", "transportes"), ("Destino", "destinos"),
                   ("Cliente", "clientes"), ("Concepto", "conceptos"), ("Entrega", "estatus"), ("Semáforo", "clases")):
        if f[k]:
            p.append(f"{lab}: {lst(f[k])}")
    if f["pct_min"]:
        p.append(f"% log. > {f['pct_min']}")
    if f["solo_fuera"]:
        p.append("solo fuera de target")
    if f["buscar"]:
        p.append(f"búsqueda: {f['buscar']}")
    p.append("costo = " + ("guía + adicionales" if f["incluir_adic"] else "solo guía"))
    return "  ·  ".join(p)


# ============================================================
# 6. CÁLCULOS: agrupaciones
# ============================================================
def _derivar(r, costo_total=None):
    r = r.copy()
    r["pct_log"] = r["costo"] / r["fact"].where(r["fact"] > 0) * 100
    r["vs"] = r["pct_log"] - TARGET
    r["costo_caja"] = r["costo"] / r["cajas"].where(r["cajas"] > 0)
    r["costo_ped"] = r["costo"] / r["n"].where(r["n"] > 0)
    tot = costo_total if costo_total is not None else r["costo"].sum()
    r["pct_gasto"] = r["costo"] / tot * 100 if tot else np.nan
    r["pct_fuera"] = r["fuera"] / r["n"].where(r["n"] > 0) * 100
    r["pct_tiempo"] = r["atiempo"] / r["ent"].where(r["ent"] > 0) * 100
    r["ticket"] = r["fact"] / r["n"].where(r["n"] > 0)
    r["ticket_min"] = r["costo_ped"] / (TARGET / 100)
    return r


def agrupar(d, col):
    g = d.groupby(col, dropna=False).agg(
        n=("_COSTO", "size"), cajas=("_CAJ", "sum"), fact=("FACTURACION", "sum"), guia=("COSTO DE LA GUÍA", "sum"),
        adic=("COSTOS ADICIONALES", "sum"), costo=("_COSTO", "sum"), exceso=("_EXC", "sum"), fuera=("_FUERA", "sum"),
        ent=("_ENT", "sum"), atiempo=("_ATIEMPO", "sum"), retr=("_RETR_TR", "sum"), inc=("_INC", "sum"),
        dias=("_DIAS", "mean")).reset_index().rename(columns={col: "K"})
    return _derivar(g)


def total_de(r):
    s = {c: r[c].sum() for c in ("n", "cajas", "fact", "guia", "adic", "costo", "exceso", "fuera", "ent", "atiempo", "retr", "inc")}
    t = _derivar(pd.DataFrame([s]), costo_total=s["costo"]).iloc[0].to_dict()
    t["K"] = "TOTAL"
    return t


def _filas(df):
    return df.replace({np.nan: None}).to_dict("records")


# ============================================================
# 7. BLOQUES (elementos neutrales que se dibujan en pantalla y en PDF)
#    ("kpis", [(etiqueta, valor_texto, tono_color, subtexto)])
#    ("table", dict(title, cols=[(hdr,key,kind,align,wide_only)], rows, total, note, landscape))
#    ("bars", dict(title, items=[(label, valor, color_tono)], kind))
#    ("note", texto)
# ============================================================
def E_kpis(items):
    return ("kpis", items)


def E_table(title, cols, rows, total=None, note="", landscape=False):
    return ("table", dict(title=title, cols=cols, rows=rows, total=total, note=note, landscape=landscape))


def E_bars(title, items, kind):
    return ("bars", dict(title=title, items=items, kind=kind))


def E_note(t):
    return ("note", t)


def _cols_grupo(nombre_dim, kind_k="txt"):
    return [(nombre_dim, "K", kind_k, "L", False), ("PEDIDOS", "n", "int", "R", False), ("CAJAS", "cajas", "int", "R", False),
            ("FACTURACIÓN", "fact", "money", "R", False), ("GUÍA", "guia", "money", "R", True),
            ("ADIC.", "adic", "money", "R", True), ("COSTO", "costo", "money", "R", False),
            ("% LOG.", "pct_log", "pctlog", "R", False), ("VS TARGET", "vs", "vs", "R", False),
            ("COSTO/CAJA", "costo_caja", "money2", "R", False), ("COSTO/PED.", "costo_ped", "money2", "R", True),
            ("% GASTO", "pct_gasto", "pct", "R", True), ("EXCESO $", "exceso", "money", "R", False),
            ("% FUERA", "pct_fuera", "pct", "R", True)]


def _orden_tabla(r, col, ctx):
    """Ordena un agrupado: orden natural para meses/cajas/semáforo, o por la métrica elegida."""
    o = ctx["opts"]
    if col == "_MESN":
        orden = {m.title(): i for i, m in enumerate(MESES)}
        return r.assign(_o=r["K"].map(orden).fillna(99)).sort_values("_o").drop(columns="_o")
    if col == "_CAJB":
        return r.assign(_o=r["K"].map(lambda x: 0 if x == "SIN CAJAS" else int(str(x).split("+")[0].split(" ")[0]))).sort_values("_o").drop(columns="_o")
    if col == "_CLASE":
        return r.assign(_o=r["K"].map(lambda x: CLASES.index(x) if x in CLASES else 99)).sort_values("_o").drop(columns="_o")
    if col == "_ANIOT":
        return r.sort_values("K")
    key = METRICAS[o["orden_met"]][0]
    return r.sort_values(key, ascending=(o["orden_dir"] == "MENOR A MAYOR"), na_position="last")


def sec_dim(ctx, titulo, dim, d=None, con_top=True):
    d = ctx["d"] if d is None else d
    col = DIMS[dim]
    if d.empty:
        return [E_note("Sin pedidos con los filtros elegidos.")]
    r = agrupar(d, col)
    tot = total_de(r)
    r = _orden_tabla(r, col, ctx)
    nat = col in NATURALES
    top = ctx["opts"]["top"]
    mostrados = r if (nat or top == "TODOS" or not con_top) else r.head(int(top))
    nota = f"Mostrando {len(mostrados)} de {len(r)} · ordenado por {ctx['opts']['orden_met'].lower()}" if len(mostrados) < len(r) else ""
    kind_k = "cls" if col == "_CLASE" else ("est" if col == "_EST" else "txt")
    out = [E_table(titulo, _cols_grupo(dim.upper(), kind_k), _filas(mostrados), total=tot, note=nota)]
    # gráfica de la métrica elegida (top 12 o natural)
    mkey, mkind = METRICAS[ctx["opts"]["orden_met"]]
    gb = mostrados.head(12) if not nat or len(mostrados) <= 12 else mostrados.head(12)
    items = []
    for _, x in gb.iterrows():
        v = x[mkey]
        if not isnum(v):
            continue
        if mkey == "pct_log":
            tono = "green" if v <= TARGET else "red"
        elif col == "_CLASE":
            tono = {"SALUDABLE": "green", "EN ALERTA": "gold", "CRÍTICO": "orange", "PÉRDIDA": "red"}.get(x["K"], "gray")
        else:
            tono = "teal"
        items.append((str(x["K"]), float(v), tono))
    if items:
        out.append(E_bars(f"{ctx['opts']['orden_met']} · {dim}", items, mkind))
    return out


def sec_kpi(ctx):
    d, ext = ctx["d"], ctx["ext"]
    if d.empty and ext.empty:
        return [E_note("Sin información con los filtros elegidos.")]
    n, cajas, fact = len(d), d["_CAJ"].sum(), d["FACTURACION"].sum()
    guia, adic, costo = d["COSTO DE LA GUÍA"].sum(), d["COSTOS ADICIONALES"].sum(), d["_COSTO"].sum()
    c_ext = (ext["COSTO DE LA GUÍA"] + ext["COSTOS ADICIONALES"]).sum() if len(ext) else 0.0
    pct = costo / fact * 100 if fact > 0 else np.nan
    pct_t = (costo + c_ext) / fact * 100 if fact > 0 else np.nan
    ent, at = int(d["_ENT"].sum()), int(d["_ATIEMPO"].sum())
    fuera, exc = int(d["_FUERA"].sum()), d["_EXC"].sum()
    tono = lambda v: "green" if isnum(v) and v <= TARGET else "red"
    txt = lambda k, v: fmt(k, v)[0]
    fila1 = [("Pedidos", txt("int", n), "blue", f"{txt('int', cajas)} cajas"),
             ("Facturación", txt("money", fact), "green", f"ticket prom. {txt('money', fact / n if n else np.nan)}"),
             ("Costo logístico", txt("money", costo), "slate", f"guía {txt('money', guia)} · adic. {txt('money', adic)}"),
             (f"% logístico (target {TARGET}%)", txt("pct2", pct), tono(pct), f"{pct - TARGET:+.2f} pp vs target" if isnum(pct) else "")]
    fila2 = [("Costo por caja", txt("money2", costo / cajas if cajas else np.nan), "orange", ""),
             ("Costo por pedido", txt("money2", costo / n if n else np.nan), "purple", ""),
             ("Costos extras", txt("money", c_ext), "gold", f"{len(ext)} registros · % log. total {txt('pct2', pct_t)}"),
             ("Fuera de target", f"{fuera} de {n}", "red" if fuera else "green", f"exceso pagado {txt('money', exc)}"),
             ("Entregados", txt("int", ent), "teal", f"{(ent / n * 100 if n else 0):.1f}% del total"),
             ("Entregas a tiempo", txt("pct", at / ent * 100 if ent else np.nan), "green" if ent and at / ent >= .9 else "gold",
              f"{at} de {ent}")]
    return [E_kpis(fila1), E_kpis(fila2)]


def sec_peq(ctx):
    d0, n = ctx["d_nocaj"], ctx["f"]["max_peq"]
    d0 = d0[d0["_CAJ"] >= 1]
    if d0.empty:
        return [E_note("Sin pedidos con cajas registradas.")]
    peq, gra = d0[d0["_CAJ"] <= n], d0[d0["_CAJ"] > n]
    filas, tot_costo = [], d0["_COSTO"].sum()
    for lab, g in ((f"1 A {n} CAJAS (PEQUEÑOS)", peq), (f"{n + 1} O MÁS CAJAS (GRANDES)", gra), ("TOTAL", d0)):
        if g.empty:
            continue
        r = _derivar(pd.DataFrame([dict(n=len(g), cajas=g["_CAJ"].sum(), fact=g["FACTURACION"].sum(), guia=g["COSTO DE LA GUÍA"].sum(),
                                        adic=g["COSTOS ADICIONALES"].sum(), costo=g["_COSTO"].sum(), exceso=g["_EXC"].sum(),
                                        fuera=g["_FUERA"].sum(), ent=g["_ENT"].sum(), atiempo=g["_ATIEMPO"].sum(), retr=0, inc=0, dias=0)]),
                     costo_total=tot_costo).iloc[0].to_dict()
        r["K"] = lab
        r["pct_ped"] = r["n"] / len(d0) * 100
        r["pct_bajo"] = (g["FACTURACION"] < r["ticket_min"]).mean() * 100
        filas.append(r)
    cols = [("GRUPO", "K", "txt", "L", False), ("PEDIDOS", "n", "int", "R", False), ("% PEDIDOS", "pct_ped", "pct", "R", False),
            ("CAJAS", "cajas", "int", "R", False), ("FACTURACIÓN", "fact", "money", "R", False), ("COSTO", "costo", "money", "R", False),
            ("% LOG.", "pct_log", "pctlog", "R", False), ("COSTO/CAJA", "costo_caja", "money2", "R", False),
            ("COSTO/PED.", "costo_ped", "money2", "R", False), ("TICKET PROM.", "ticket", "money", "R", True),
            ("TICKET MÍN. TARGET", "ticket_min", "money", "R", True), ("% BAJO MÍNIMO", "pct_bajo", "pct", "R", True),
            ("EXCESO $", "exceso", "money", "R", False)]
    out = []
    if not peq.empty:
        p = filas[0]
        out.append(E_kpis([("Pedidos pequeños", fmt("int", p["n"])[0], "blue", f"{p['pct_ped']:.1f}% de los pedidos"),
                           ("% logístico pequeños", fmt("pct2", p["pct_log"])[0], "green" if isnum(p["pct_log"]) and p["pct_log"] <= TARGET else "red",
                            f"target {TARGET}%"),
                           ("Ticket mínimo para estar en target", fmt("money", p["ticket_min"])[0], "gold", f"ticket actual {fmt('money', p['ticket'])[0]}"),
                           ("Exceso pagado (pequeños)", fmt("money", p["exceso"])[0], "red", "sobre el target")]))
    out.append(E_table("Pequeños vs grandes", cols, [x for x in filas[:-1]], total=filas[-1],
                       note=f"Pequeño = de 1 a {n} cajas. El filtro de cajas no aplica a este bloque; los demás filtros sí."))
    out += sec_dim(ctx, "Detalle por número de cajas", "Nº de cajas", d=d0)
    return out


COLS_RANK = [("FACTURA", "_FAC", "txt", "L", False), ("FECHA", "FECHA DE ENVÍO", "date", "L", False),
             ("CLIENTE", "_CLI", "txt", "L", False), ("DESTINO", "_DE", "txt", "L", False), ("CJ.", "_CAJ", "int", "R", False),
             ("COBRO", "_MODC", "txt", "L", False), ("TRANSPORTE", "_TR", "txt", "L", False), ("FLETERA", "_FLE", "txt", "L", True),
             ("FACTURACIÓN", "FACTURACION", "money", "R", False), ("GUÍA", "COSTO DE LA GUÍA", "money", "R", True),
             ("ADIC.", "COSTOS ADICIONALES", "money", "R", True), ("COSTO", "_COSTO", "money", "R", False),
             ("% LOG.", "_PCT", "pctlog", "R", False), ("VS TGT", "_VS", "vs", "R", False),
             ("EXCESO $", "_EXC", "money", "R", False), ("SEMÁFORO", "_CLASE", "cls", "L", False),
             ("ENTREGA", "_EST", "est", "L", True)]
ORDEN_RANK = {"% logístico": "_PCT", "Costo logístico ($)": "_COSTO", "Exceso sobre target ($)": "_EXC", "Facturación": "FACTURACION",
              "Costo por caja": "_CPC", "Cajas": "_CAJ", "Fecha de envío": "FECHA DE ENVÍO"}


def sec_rank(ctx):
    d, o = ctx["d"], ctx["opts"]
    if d.empty:
        return [E_note("Sin pedidos con los filtros elegidos.")]
    r = d.copy()
    r["_MODC"] = r["_MOD"].map({"COBRO REGRESO": "REGRESO", "COBRO DESTINO": "DESTINO"}).fillna(r["_MOD"])
    r["_VS"] = r["_PCT"] - TARGET
    r = r.sort_values(ORDEN_RANK[o["rank_orden"]], ascending=(o["rank_dir"] == "MENOR A MAYOR"), na_position="last")
    total_n = len(r)
    top = o["rank_top"]
    sel = r if top == "TODOS" else r.head(int(top))
    sel = sel.head(SCREEN_MAX)
    rows = sel[[c[1] for c in COLS_RANK]].replace({np.nan: None, pd.NaT: None}).to_dict("records")
    tot = dict(_FAC="TOTAL", _CAJ=sel["_CAJ"].sum(), FACTURACION=sel["FACTURACION"].sum(), **{"COSTO DE LA GUÍA": sel["COSTO DE LA GUÍA"].sum(),
               "COSTOS ADICIONALES": sel["COSTOS ADICIONALES"].sum()}, _COSTO=sel["_COSTO"].sum(),
               _PCT=(sel["_COSTO"].sum() / sel["FACTURACION"].sum() * 100) if sel["FACTURACION"].sum() > 0 else None,
               _EXC=sel["_EXC"].sum())
    tot["_VS"] = tot["_PCT"] - TARGET if tot["_PCT"] is not None else None
    nota = (f"Mostrando {len(sel)} de {total_n} pedidos · ordenado por {o['rank_orden'].lower()} ({o['rank_dir'].lower()}). "
            f"Pedidos = envíos sin recolecciones/maniobras.")
    return [E_table(f"Pedido por pedido · {o['rank_orden']}", COLS_RANK, rows, total=tot, note=nota, landscape=True)]


def sec_entr(ctx):
    d = ctx["d"]
    if d.empty:
        return [E_note("Sin pedidos con los filtros elegidos.")]
    r = agrupar(d, "_FLE")
    tot = total_de(r)
    r["trans"], tot["trans"] = r["n"] - r["ent"], tot["n"] - tot["ent"]
    r["retraso"], tot["retraso"] = r["ent"] - r["atiempo"], tot["ent"] - tot["atiempo"]
    r["pct_inc"], tot["pct_inc"] = r["inc"] / r["n"] * 100, tot["inc"] / tot["n"] * 100 if tot["n"] else None
    tot["dias"] = d["_DIAS"].mean()
    r = r.sort_values("pct_tiempo", ascending=False, na_position="last")
    cols = [("FLETERA", "K", "txt", "L", False), ("PEDIDOS", "n", "int", "R", False), ("ENTREGADOS", "ent", "int", "R", False),
            ("A TIEMPO", "atiempo", "int", "R", False), ("CON RETRASO", "retraso", "int", "R", False),
            ("% A TIEMPO", "pct_tiempo", "pctok", "R", False), ("EN TRÁNSITO", "trans", "int", "R", False),
            ("TRÁNSITO RETRASADO", "retr", "int", "R", True), ("DÍAS PROM.", "dias", "num1", "R", False),
            ("% INCIDENCIAS", "pct_inc", "pct", "R", False)]
    items = [(str(x.K), float(x.pct_tiempo), "green" if x.pct_tiempo >= 90 else ("gold" if x.pct_tiempo >= 80 else "red"))
             for x in r.itertuples() if isnum(x.pct_tiempo)][:12]
    out = [E_table("Efectividad de entrega por fletera", cols, _filas(r), total=tot,
                   note="Semáforo de % a tiempo: verde ≥ 90 %, ámbar ≥ 80 %, rojo menor. A tiempo = entrega real ≤ promesa.")]
    if items:
        out.append(E_bars("% de entregas a tiempo por fletera", items, "pct"))
    return out


def sec_extr(ctx):
    ext = ctx["ext"]
    if ext.empty:
        return [E_note("No hay recolecciones ni maniobras con los filtros elegidos.")]
    e = ext.assign(_ECOSTO=ext["COSTO DE LA GUÍA"] + ext["COSTOS ADICIONALES"])
    costo_ped = ctx["d"]["_COSTO"].sum()
    fact = ctx["d"]["FACTURACION"].sum()
    tot = e["_ECOSTO"].sum()
    out = [E_kpis([("Costo extra total", fmt("money", tot)[0], "gold", f"{len(e)} registros"),
                   ("% del costo logístico total", fmt("pct", tot / (tot + costo_ped) * 100 if tot + costo_ped else np.nan)[0], "orange", ""),
                   ("% logístico solo pedidos", fmt("pct2", costo_ped / fact * 100 if fact > 0 else np.nan)[0], "teal", ""),
                   ("% logístico con extras", fmt("pct2", (costo_ped + tot) / fact * 100 if fact > 0 else np.nan)[0],
                    "green" if fact > 0 and (costo_ped + tot) / fact * 100 <= TARGET else "red", f"target {TARGET}%")])]
    for titulo, col in (("Costos extras por concepto", "_CONC"), ("Costos extras por fletera", "_FLE")):
        g = e.groupby(col).agg(n=("_ECOSTO", "size"), guia=("COSTO DE LA GUÍA", "sum"), adic=("COSTOS ADICIONALES", "sum"),
                               costo=("_ECOSTO", "sum")).reset_index().sort_values("costo", ascending=False)
        g["pct"] = g["costo"] / tot * 100 if tot else np.nan
        cols = [(col.strip("_").replace("CONC", "CONCEPTO").replace("FLE", "FLETERA"), col, "txt", "L", False), ("REGISTROS", "n", "int", "R", False),
                ("GUÍA", "guia", "money", "R", False), ("ADIC.", "adic", "money", "R", False), ("COSTO", "costo", "money", "R", False),
                ("% DEL EXTRA", "pct", "pct", "R", False)]
        out.append(E_table(titulo, cols, _filas(g), total=dict(**{col: "TOTAL"}, n=len(e), guia=e["COSTO DE LA GUÍA"].sum(),
                                                               adic=e["COSTOS ADICIONALES"].sum(), costo=tot, pct=100.0)))
    return out


# --- Muestras ---------------------------------------------------------------
@st.cache_data(ttl=300)
def datos_muestras():
    """Devuelve (df_preparado | None, mensaje)."""
    if not MUESTRAS_OK:
        return None, f"No se pudo importar muestras_common: {MUESTRAS_ERR}"
    try:
        df_m, _sha = obtener_datos_github()
    except Exception as e:
        return None, f"No se pudo leer el registro de muestras: {e}"
    if df_m is None or df_m.empty:
        return None, "No hay registros de muestras cargados."
    d = df_m.copy()
    for col in ["PAQUETERIA_NOMBRE", "NUMERO_GUIA", "COSTO_GUIA", "CANTIDAD_TOTAL", "COSTO_TOTAL", "ESTATUS", "SOLICITO",
                "NOMBRE DEL HOTEL", "FOLIO", "FECHA"]:
        if col not in d.columns:
            d[col] = "NO SURTIDO" if col == "ESTATUS" else (0.0 if col in ("COSTO_GUIA", "CANTIDAD_TOTAL", "COSTO_TOTAL") else "")
    d["FECHA_DT"] = parse_fecha_segura(d["FECHA"].astype(str).str.strip())
    d["COSTO_TOTAL"] = pd.to_numeric(d["COSTO_TOTAL"], errors="coerce").fillna(0)
    d["COSTO_GUIA"] = pd.to_numeric(d["COSTO_GUIA"], errors="coerce").fillna(0)
    d["INVERSION"] = d["COSTO_TOTAL"] + d["COSTO_GUIA"]
    d["_SOL"] = d["SOLICITO"].astype(str).str.upper().replace("", "SIN DATO")
    col_paq = "PAQUETERIA_NOMBRE" if "PAQUETERIA_NOMBRE" in d.columns else "PAQUETERIA"
    d["_PAQ"] = d[col_paq].astype(str).str.upper().replace({"": "SIN ASIGNAR", "0.0": "SIN ASIGNAR", "NAN": "SIN ASIGNAR"})
    d["_HOT"] = d["NOMBRE DEL HOTEL"].astype(str).str.upper().replace("", "SIN DATO")
    d["_STAT"] = d["ESTATUS"].astype(str).str.upper()
    return d, ""


def sec_mues(ctx):
    dm, msg = datos_muestras()
    if dm is None:
        return [E_note(f"Costos de muestras sin datos: {msg}")]
    f, o = ctx["f"], ctx["opts"]
    m = dm["FECHA_DT"].dt.year.isin(f["anios"])
    if f["meses"]:
        m &= dm["FECHA_DT"].dt.month.isin(f["meses"])
    if o["mu_sol"]:
        m &= dm["_SOL"].isin(o["mu_sol"])
    d = dm[m].copy()
    if d.empty:
        return [E_note("No hay muestras en el periodo elegido (se usan los mismos mes y año del filtro general).")]
    n, prod, flete, inv = len(d), d["COSTO_TOTAL"].sum(), d["COSTO_GUIA"].sum(), d["INVERSION"].sum()
    desp = (d["_STAT"] == "DESPACHADO").mean() * 100
    out = [E_kpis([("Envíos de muestra", fmt("int", n)[0], "blue", f"{desp:.0f}% despachado"),
                   ("Costo de producto", fmt("money", prod)[0], "purple", ""),
                   ("Costo de flete", fmt("money", flete)[0], "orange", ""),
                   ("Inversión total", fmt("money", inv)[0], "red", f"prom. {fmt('money', inv / n)[0]} por envío")])]

    def grp(col, titulo, nombre):
        g = d.groupby(col).agg(envios=("INVERSION", "size"), prod=("COSTO_TOTAL", "sum"), flete=("COSTO_GUIA", "sum"),
                               inv=("INVERSION", "sum")).reset_index().sort_values("inv", ascending=False)
        g["prom"] = g["inv"] / g["envios"]
        g["pct"] = g["inv"] / inv * 100 if inv else np.nan
        top = o["top"]
        gm = g if top == "TODOS" else g.head(int(top))
        cols = [(nombre, col, "txt", "L", False), ("ENVÍOS", "envios", "int", "R", False), ("PRODUCTO", "prod", "money", "R", False),
                ("FLETE", "flete", "money", "R", False), ("INVERSIÓN", "inv", "money", "R", False),
                ("PROM./ENVÍO", "prom", "money", "R", False), ("% DEL TOTAL", "pct", "pct", "R", False)]
        tot = dict(**{col: "TOTAL"}, envios=n, prod=prod, flete=flete, inv=inv, prom=inv / n, pct=100.0)
        res = [E_table(titulo, cols, _filas(gm), total=tot, note=f"Mostrando {len(gm)} de {len(g)}" if len(gm) < len(g) else "")]
        return res, [(str(x[col]), float(x["inv"]), "red") for _, x in gm.head(12).iterrows()]

    t1, b1 = grp("_SOL", "Muestras por solicitante", "SOLICITÓ")
    out += t1
    if b1:
        out.append(E_bars("Inversión en muestras por solicitante", b1, "money"))
    out += grp("_PAQ", "Muestras por paquetería", "PAQUETERÍA")[0]
    out += grp("_HOT", "Muestras por destino / hotel", "HOTEL")[0]
    d["_MM"] = d["FECHA_DT"].dt.year.astype(int).astype(str) + "-" + d["FECHA_DT"].dt.month.astype(int).map("{:02d}".format)
    mm = d.groupby("_MM").agg(envios=("INVERSION", "size"), prod=("COSTO_TOTAL", "sum"), flete=("COSTO_GUIA", "sum"),
                              inv=("INVERSION", "sum")).reset_index().sort_values("_MM")
    mm["prom"] = mm["inv"] / mm["envios"]
    mm["pct"] = mm["inv"] / inv * 100 if inv else np.nan
    out.append(E_table("Muestras por mes", [("MES", "_MM", "txt", "L", False), ("ENVÍOS", "envios", "int", "R", False),
                                           ("PRODUCTO", "prod", "money", "R", False), ("FLETE", "flete", "money", "R", False),
                                           ("INVERSIÓN", "inv", "money", "R", False), ("PROM./ENVÍO", "prom", "money", "R", False),
                                           ("% DEL TOTAL", "pct", "pct", "R", False)], _filas(mm)))
    if o["mu_detalle"]:
        det = d.sort_values("INVERSION", ascending=False)
        det = det.head(SCREEN_MAX if o["top"] == "TODOS" else max(int(o["top"]), 20))
        det = det.assign(_FOL=det["FOLIO"].astype(str))
        cols = [("FOLIO", "_FOL", "txt", "L", False), ("FECHA", "FECHA_DT", "date", "L", False), ("SOLICITÓ", "_SOL", "txt", "L", False),
                ("HOTEL / DESTINO", "_HOT", "txt", "L", False), ("PAQUETERÍA", "_PAQ", "txt", "L", True), ("ESTATUS", "_STAT", "txt", "L", True),
                ("PRODUCTO", "COSTO_TOTAL", "money", "R", False), ("FLETE", "COSTO_GUIA", "money2", "R", False), ("INVERSIÓN", "INVERSION", "money", "R", False)]
        out.append(E_table("Detalle de muestras (mayor inversión primero)", cols,
                           det[[c[1] for c in cols]].replace({np.nan: None, pd.NaT: None}).to_dict("records"), landscape=True))
    return out


# --- Tabla cruzada ----------------------------------------------------------
CRUCE_MET = {"Costo logístico ($)": ("costo", "money"), "% logístico": ("pct_log", "pctlog"), "Facturación": ("fact", "money"),
             "Pedidos": ("n", "int"), "Cajas": ("cajas", "int"), "Costo por caja": ("costo_caja", "money2"),
             "Exceso sobre target ($)": ("exceso", "money")}


def sec_cruz(ctx):
    d, o = ctx["d"], ctx["opts"]
    rdim, cdim = o["cz_filas"], o["cz_cols"]
    if d.empty:
        return [E_note("Sin pedidos con los filtros elegidos.")]
    if rdim == cdim:
        return [E_note("Elige dimensiones distintas para filas y columnas.")]
    rcol, ccol = DIMS[rdim], DIMS[cdim]
    mkey, mkind = CRUCE_MET[o["cz_met"]]
    x = d[[rcol, ccol, "_COSTO", "FACTURACION", "_CAJ", "_EXC"]].copy()
    x["n"] = 1
    # columnas: máx. 12 (el resto va a OTROS)
    cnat = ccol in NATURALES
    orden_c = x.groupby(ccol)["_COSTO"].sum().sort_values(ascending=False).index.tolist()
    if ccol == "_MESN":
        orden_c = [m.title() for m in MESES if m.title() in set(x[ccol])] + [c for c in orden_c if c not in [m.title() for m in MESES]]
    elif ccol == "_CAJB":
        orden_c = sorted(orden_c, key=lambda s: 0 if s == "SIN CAJAS" else int(str(s).split("+")[0].split(" ")[0]))
    elif ccol == "_CLASE":
        orden_c = [c for c in CLASES if c in orden_c]
    elif ccol == "_ANIOT":
        orden_c = sorted(orden_c)
    if len(orden_c) > 12:
        keep = orden_c[:11]
        x[ccol] = np.where(x[ccol].isin(keep), x[ccol], "OTROS")
        orden_c = keep + ["OTROS"]
    sums = {k: pd.pivot_table(x, index=rcol, columns=ccol, values=v, aggfunc="sum", fill_value=0)
            for k, v in (("costo", "_COSTO"), ("fact", "FACTURACION"), ("cajas", "_CAJ"), ("exceso", "_EXC"), ("n", "n"))}
    sums = {k: v.reindex(columns=orden_c, fill_value=0) for k, v in sums.items()}

    def valor(s):   # s: dict de sumas (escalar/Series) -> métrica
        if mkey == "pct_log":
            return s["costo"] / s["fact"] * 100 if np.all(np.asarray(s["fact"]) > 0) else None
        if mkey == "costo_caja":
            return s["costo"] / s["cajas"] if np.all(np.asarray(s["cajas"]) > 0) else None
        return s[mkey]

    filas_idx = sums["costo"].sum(axis=1).sort_values(ascending=False).index.tolist()
    if rcol == "_MESN":
        filas_idx = [m.title() for m in MESES if m.title() in set(x[rcol])] + [r for r in filas_idx if r not in [m.title() for m in MESES]]
    elif rcol == "_CAJB":
        filas_idx = sorted(filas_idx, key=lambda s: 0 if s == "SIN CAJAS" else int(str(s).split("+")[0].split(" ")[0]))
    elif rcol == "_CLASE":
        filas_idx = [c for c in CLASES if c in filas_idx]
    elif rcol == "_ANIOT":
        filas_idx = sorted(filas_idx)
    elif o["top"] != "TODOS":
        filas_idx = filas_idx[: int(o["top"])]

    def fila(idx):
        r = {"K": idx}
        for i, c in enumerate(orden_c):
            s = {k: float(v.at[idx, c]) if idx in v.index else 0.0 for k, v in sums.items()}
            r[f"c{i}"] = (valor(s) if (mkey in ("pct_log", "costo_caja") or s["n"] > 0) else None)
        s = {k: float(v.loc[idx].sum()) for k, v in sums.items()}
        r["tot"] = valor(s)
        return r
    rows = [fila(i) for i in filas_idx]
    stot = {k: float(v.loc[filas_idx].sum().sum()) for k, v in sums.items()}
    tr = {"K": "TOTAL", "tot": valor(stot)}
    for i, c in enumerate(orden_c):
        s = {k: float(v.loc[filas_idx, c].sum()) for k, v in sums.items()}
        tr[f"c{i}"] = valor(s)
    cols = [(rdim.upper(), "K", "txt", "L", False)] + [(trunc(str(c), 14), f"c{i}", mkind, "R", False) for i, c in enumerate(orden_c)] + \
           [("TOTAL", "tot", mkind, "R", False)]
    return [E_table(f"Tabla cruzada · {o['cz_met']} · {rdim} × {cdim}", cols, rows, total=tr, landscape=len(cols) > 9,
                    note="Las columnas se limitan a 12; lo demás se agrupa en OTROS.")]


def construir(ctx, activas):
    """Devuelve [(clave, titulo, subtitulo, [elementos])] en el orden de SECCIONES."""
    dims = {"flet": ("Por fletera", "Fletera"), "tran": ("Por transporte", "Transporte"), "dest": ("Por destino", "Destino"),
            "clie": ("Por cliente", "Cliente"), "moda": ("Cobro regreso vs cobro destino", "Modalidad (cobro)"),
            "conc": ("Por concepto", "Concepto"), "mes": ("Por mes", "Mes"), "caja": ("Por número de cajas", "Nº de cajas"),
            "sema": ("Semáforo de salud", "Semáforo")}
    out = []
    for key, nombre, desc in SECCIONES:
        if key not in activas:
            continue
        if key in dims:
            els = sec_dim(ctx, dims[key][0], dims[key][1])
        else:
            els = {"kpi": sec_kpi, "peq": sec_peq, "rank": sec_rank, "entr": sec_entr, "extr": sec_extr,
                   "mues": sec_mues, "cruz": sec_cruz}[key](ctx)
        out.append((key, nombre.upper(), desc, els))
    return out


# ============================================================
# 8. RENDER EN PANTALLA (tarjetas premium + tablas con encabezado sticky)
# ============================================================
CSS = """
<style>
.jq-sec{display:flex;align-items:baseline;gap:12px;margin:34px 0 14px 0;padding:12px 16px;border-radius:10px;
 background:linear-gradient(90deg,#2B343B 0%,#232B31 100%);border-left:5px solid #00A3A3;box-shadow:0 4px 14px rgba(0,0,0,.25)}
.jq-sec .n{color:#FFC000;font-weight:800;font-size:13px;letter-spacing:1px}
.jq-sec .t{color:#fff;font-weight:800;font-size:15px;letter-spacing:1.5px}
.jq-sec .s{color:#9FB3BF;font-size:11px;margin-left:auto}
.jq-cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(210px,1fr));gap:14px;margin:6px 0 14px 0}
.jq-card{position:relative;overflow:hidden;padding:16px 18px 15px 18px;border-radius:14px;
 background:linear-gradient(145deg,#2F3A42 0%,#1F282D 100%);border:1px solid rgba(255,255,255,.07);
 box-shadow:0 8px 22px rgba(0,0,0,.30),inset 0 1px 0 rgba(255,255,255,.05);transition:transform .2s,box-shadow .2s}
.jq-card:hover{transform:translateY(-3px);box-shadow:0 12px 28px rgba(0,0,0,.40),inset 0 1px 0 rgba(255,255,255,.07)}
.jq-card:before{content:"";position:absolute;left:0;top:0;right:0;height:3px;background:var(--c)}
.jq-card:after{content:"";position:absolute;right:-30px;top:-30px;width:110px;height:110px;border-radius:50%;
 background:radial-gradient(circle,var(--c) 0%,transparent 70%);opacity:.14}
.jq-card .l{color:#9FB3BF;font-size:10px;font-weight:800;letter-spacing:1.4px;text-transform:uppercase}
.jq-card .v{color:var(--c);font-size:26px;font-weight:800;margin:7px 0 3px 0;line-height:1.1}
.jq-card .v.sm{font-size:19px}
.jq-card .sub{color:#8497A3;font-size:11px;min-height:14px}
.jq-h{color:#E8EEF1;font-weight:800;font-size:12px;letter-spacing:1.4px;text-transform:uppercase;margin:20px 0 8px 2px}
.jq-tw{max-height:520px;overflow:auto;border-radius:12px;border:1px solid rgba(255,255,255,.08);
 box-shadow:0 8px 22px rgba(0,0,0,.28);background:#1B2328}
.jq-t{border-collapse:separate;border-spacing:0;width:100%;font-size:12.5px;color:#DCE5EA}
.jq-t th{position:sticky;top:0;z-index:3;background:#2B343B;color:#fff;font-size:10.5px;letter-spacing:1px;font-weight:800;
 padding:10px 12px;border-bottom:2px solid #00A3A3;white-space:nowrap;text-transform:uppercase}
.jq-t td{padding:8px 12px;border-bottom:1px solid rgba(255,255,255,.05);white-space:nowrap;background:#1B2328}
.jq-t tbody tr:nth-child(even) td{background:#202A30}
.jq-t tbody tr:hover td{background:#2A3840}
.jq-t th:first-child,.jq-t td:first-child{position:sticky;left:0;z-index:2}
.jq-t th:first-child{z-index:4}
.jq-t td:first-child{font-weight:700;color:#fff;box-shadow:2px 0 0 rgba(0,0,0,.25);max-width:260px;overflow:hidden;text-overflow:ellipsis}
.jq-t tfoot td{position:sticky;bottom:0;z-index:3;background:#2B343B!important;color:#fff;font-weight:800;border-top:2px solid #FFC000}
.jq-t tfoot td:first-child{z-index:4}
.jq-t .r{text-align:right;font-variant-numeric:tabular-nums}
.jq-t .good{color:#4FD1A0;font-weight:700}.jq-t .bad{color:#FF6B6B;font-weight:700}
.jq-t .warn{color:#FFC000;font-weight:700}.jq-t .mute{color:#6F808B}
.jq-pill{display:inline-block;padding:2px 9px;border-radius:20px;font-size:10.5px;font-weight:800;letter-spacing:.5px}
.jq-pill.good{background:rgba(79,209,160,.15)}.jq-pill.bad{background:rgba(255,107,107,.15)}
.jq-pill.warn{background:rgba(255,192,0,.15)}.jq-pill.mute{background:rgba(111,128,139,.18)}
.jq-nota{color:#8497A3;font-size:11px;margin:6px 2px 0 2px;font-style:italic}
.jq-aviso{padding:12px 16px;border-radius:10px;background:#2B343B;border-left:4px solid #FFC000;color:#DCE5EA;font-size:13px;margin:8px 0}
.jq-bars{background:linear-gradient(145deg,#2F3A42,#1F282D);border:1px solid rgba(255,255,255,.07);border-radius:14px;padding:14px 18px;
 box-shadow:0 8px 22px rgba(0,0,0,.28)}
.jq-b{display:grid;grid-template-columns:minmax(120px,230px) 1fr 100px;gap:12px;align-items:center;margin:7px 0;font-size:12px;color:#DCE5EA}
.jq-b .lab{overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.jq-b .tr{height:13px;background:rgba(255,255,255,.06);border-radius:8px;overflow:hidden}
.jq-b .fi{height:100%;border-radius:8px;background:linear-gradient(90deg,var(--c),var(--c2))}
.jq-b .val{text-align:right;font-weight:800;color:#fff;font-variant-numeric:tabular-nums}
.jq-chips{display:flex;flex-wrap:wrap;gap:8px;margin:4px 0 10px 0}
.jq-chip{padding:5px 12px;border-radius:20px;background:#2B343B;border:1px solid rgba(0,163,163,.45);color:#DCE5EA;font-size:11px}
</style>
"""


def _s(t):
    """Texto seguro para st.markdown (escapa HTML y el signo $ que Streamlit toma como LaTeX)."""
    return esc(t).replace("$", "&#36;")


def html_kpis(items):
    out = ['<div class="jq-cards">']
    for lab, val, tono, sub in items:
        c = HEX.get(tono, HEX["teal"])
        if tono == "slate":
            c = "#DCE5EA"  # slate = mismo color que la tarjeta; en pantalla se ve gris claro
        sm = " sm" if len(str(val)) > 12 else ""
        out.append(f'<div class="jq-card" style="--c:{c}"><div class="l">{_s(lab)}</div><div class="v{sm}">{_s(val)}</div>'
                   f'<div class="sub">{_s(sub)}</div></div>')
    out.append("</div>")
    return "".join(out)


def _td(kind, align, v):
    txt, tono = fmt(kind, v)
    cls = ("r " if align == "R" else "") + tono
    if kind in ("cls", "est") and txt:
        return f'<td class="{cls.strip()}"><span class="jq-pill {tono}">{_s(txt)}</span></td>'
    return f'<td class="{cls.strip()}">{_s(txt)}</td>'


def html_table(t):
    cols, rows = t["cols"], t["rows"][:SCREEN_MAX]
    out = []
    if t["title"]:
        out.append(f'<div class="jq-h">{_s(t["title"])}</div>')
    out.append('<div class="jq-tw"><table class="jq-t"><thead><tr>')
    out += [f'<th class="{"r" if c[3] == "R" else ""}">{_s(c[0])}</th>' for c in cols]
    out.append("</tr></thead><tbody>")
    for r in rows:
        out.append("<tr>" + "".join(_td(c[2], c[3], r.get(c[1])) for c in cols) + "</tr>")
    out.append("</tbody>")
    if t["total"]:
        out.append("<tfoot><tr>" + "".join(
            f'<td class="{"r" if c[3] == "R" else ""}">{_s(fmt(c[2], t["total"].get(c[1]))[0]) if c[2] not in ("cls", "est") else _s(t["total"].get(c[1]) or "")}</td>'
            for c in cols) + "</tr></tfoot>")
    out.append("</table></div>")
    if t["note"]:
        out.append(f'<div class="jq-nota">{_s(t["note"])}</div>')
    return "".join(out)


def html_bars(b):
    items = b["items"]
    vmax = max([v for _, v, _ in items] + [1e-9])
    out = [f'<div class="jq-h">{_s(b["title"])}</div><div class="jq-bars">']
    for lab, v, tono in items:
        c = HEX.get(tono, HEX["teal"])
        if tono == "slate":
            c = "#9FB3BF"
        w = max(v / vmax * 100, 1.5) if v > 0 else 0
        out.append(f'<div class="jq-b" style="--c:{c};--c2:{c}99"><div class="lab" title="{esc(lab)}">{_s(lab)}</div>'
                   f'<div class="tr"><div class="fi" style="width:{w:.1f}%"></div></div><div class="val">{_s(fmt(b["kind"], v)[0])}</div></div>')
    out.append("</div>")
    return "".join(out)


def render_pantalla(secciones, resumen_filtros, f):
    st.markdown(CSS, unsafe_allow_html=True)
    chips = "".join(f'<span class="jq-chip">{_s(p.strip())}</span>' for p in resumen_filtros.split("  ·  "))
    st.markdown(f'<div class="jq-chips">{chips}</div>', unsafe_allow_html=True)
    for i, (key, nombre, desc, els) in enumerate(secciones, 1):
        st.markdown(f'<div class="jq-sec"><span class="n">{i:02d}</span><span class="t">{_s(nombre)}</span>'
                    f'<span class="s">{_s(desc)}</span></div>', unsafe_allow_html=True)
        for kind, payload in els:
            if kind == "kpis":
                st.markdown(html_kpis(payload), unsafe_allow_html=True)
            elif kind == "table":
                st.markdown(html_table(payload), unsafe_allow_html=True)
            elif kind == "bars":
                st.markdown(html_bars(payload), unsafe_allow_html=True)
            else:
                st.markdown(f'<div class="jq-aviso">{_s(payload)}</div>', unsafe_allow_html=True)


# ============================================================
# 9. PDF DE ANÁLISIS (mismo estilo que el reporte mensual)
# ============================================================
def _estilos():
    b = dict(fontName="Helvetica", textColor=C_TEXT)
    return {"sec_t": ParagraphStyle("sec_t", fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=colors.white),
            "sec_s": ParagraphStyle("sec_s", fontName="Helvetica", fontSize=8, leading=10, textColor=colors.HexColor("#B7C4CC")),
            "sub": ParagraphStyle("sub", fontName="Helvetica-Bold", fontSize=9, leading=12, textColor=C_SLATE, spaceBefore=6, spaceAfter=4),
            "kpi_l": ParagraphStyle("kpi_l", fontName="Helvetica-Bold", fontSize=6.5, leading=8, textColor=C_GRAY),
            "kpi_v": ParagraphStyle("kpi_v", fontName="Helvetica-Bold", fontSize=15, leading=18),
            "kpi_v_s": ParagraphStyle("kpi_v_s", fontName="Helvetica-Bold", fontSize=10.5, leading=13),
            "kpi_s": ParagraphStyle("kpi_s", fontName="Helvetica", fontSize=6.5, leading=8, textColor=C_GRAY),
            "nota": ParagraphStyle("nota", fontName="Helvetica-Oblique", fontSize=7.5, leading=10, textColor=C_GRAY),
            "body": ParagraphStyle("body", fontSize=8.5, leading=11.5, **b)}


ST = _estilos()


class _Canvas(rl_canvas.Canvas):
    titulo_pie, pw = "", letter[0]

    def __init__(self, *a, **k):
        super().__init__(*a, **k)
        self._saved = []

    def showPage(self):
        self._saved.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        n = len(self._saved)
        for s_ in self._saved:
            self.__dict__.update(s_)
            self.setStrokeColor(C_BORDER)
            self.setLineWidth(0.5)
            self.line(36, 34, self.pw - 36, 34)
            self.setFont("Helvetica", 7)
            self.setFillColor(C_GRAY)
            self.drawString(36, 23, f"JYPESA | Logística  -  {self.titulo_pie}  -  Generado e impreso con Nexion Smart Logistic")
            self.drawRightString(self.pw - 36, 23, f"Página {self._pageNumber} de {n}")
            super().showPage()
        super().save()


def _pdf_sec(num_, titulo, sub, W):
    cont = [Paragraph(f"{num_:02d}&nbsp;&nbsp;{esc(titulo)}", ST["sec_t"])]
    if sub:
        cont.append(Paragraph(esc(sub), ST["sec_s"]))
    t = Table([[cont]], colWidths=[W])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), C_SLATE), ("LINEBEFORE", (0, 0), (0, 0), 5, C_TEAL),
                           ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 12),
                           ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    return [CondPageBreak(110), t, Spacer(1, 9)]


def _pdf_kpis(items, W, gap=8):
    n = len(items)
    ncols = min(n, 4) if n != 5 else 3
    cw = (W - gap * (ncols - 1)) / ncols
    out = []
    for i in range(0, n, ncols):
        chunk = items[i:i + ncols]
        row, widths, style = [], [], []
        for j, (lab, val, tono, sub) in enumerate(chunk):
            c = j * 2
            col = HEX.get(tono, HEX["teal"])
            row.append([Paragraph(esc(lab).upper(), ST["kpi_l"]), Spacer(1, 3),
                        Paragraph(f'<font color="{col}">{esc(val)}</font>', ST["kpi_v"] if len(str(val)) <= 13 else ST["kpi_v_s"]),
                        Paragraph(esc(sub), ST["kpi_s"])])
            widths.append(cw)
            style += [("BACKGROUND", (c, 0), (c, 0), C_LIGHT), ("LINEABOVE", (c, 0), (c, 0), 3, colors.HexColor(col)),
                      ("BOX", (c, 0), (c, 0), 0.5, C_BORDER)]
            if j < len(chunk) - 1:
                row.append("")
                widths.append(gap)
        t = Table([row], colWidths=widths, hAlign="LEFT")
        t.setStyle(TableStyle(style + [("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 8),
                                       ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 7),
                                       ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        out += [t, Spacer(1, 8)]
    return out


def _pdf_table(t, W, wide):
    cols = [c for c in t["cols"] if wide or not c[4]]
    rows = t["rows"][:PDF_MAX]
    nc = len(cols)
    font = 7.5 if nc <= 8 else (7 if nc <= 11 else (6.3 if nc <= 14 else 5.6))
    pesos = []
    for c in cols:
        k = c[2]
        if (k == "txt" and c[1] in ("_CLI", "K", "_HOT")) or c[0] in ("CLIENTE", "HOTEL / DESTINO"):
            pesos.append(2.6)
        elif k == "date":
            pesos.append(1.7)
        elif k in ("vs", "pctlog", "pctok"):
            pesos.append(1.45)
        elif k in ("money", "money2"):
            pesos.append(1.25)
        elif k in ("cls", "est"):
            pesos.append(1.9)
        elif k == "txt":
            pesos.append(1.7)
        else:
            pesos.append(0.9 if len(c[0]) <= 4 else 1.1)
    sw = W / sum(pesos)
    widths = [p * sw for p in pesos]
    amap = {"L": 0, "C": 1, "R": 2}

    def cell(txt, a, tono="", bold=False):
        return Paragraph(esc(txt), ParagraphStyle("c", fontName="Helvetica-Bold" if bold or tono in ("good", "bad", "warn") else "Helvetica",
                                                  fontSize=font, leading=font + 2, alignment=amap[a], textColor=TONO_COLOR.get(tono, C_TEXT)))
    head = [Paragraph(esc(c[0]), ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=font - 0.3, leading=font + 1.5,
                                                 textColor=colors.white, alignment=amap[c[3]])) for c in cols]
    data = [head]
    for r in rows:
        fila = []
        for c in cols:
            txt, tono = fmt(c[2], r.get(c[1]))
            fila.append(cell(txt, c[3], tono))
        data.append(fila)
    if t["total"]:
        data.append([cell(fmt(c[2], t["total"].get(c[1]))[0] if c[2] not in ("cls", "est") else str(t["total"].get(c[1]) or ""),
                          c[3], bold=True) for c in cols])
    tb = Table(data, colWidths=widths, repeatRows=1)
    sty = [("BACKGROUND", (0, 0), (-1, 0), C_SLATE), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
           ("LINEBELOW", (0, 0), (-1, -1), 0.3, C_BORDER), ("BOX", (0, 0), (-1, -1), 0.5, C_BORDER),
           ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
           ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
    for i in range(2, len(data), 2):
        sty.append(("BACKGROUND", (0, i), (-1, i), C_LIGHT))
    if t["total"]:
        sty += [("BACKGROUND", (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor("#E3EAEE")),
                ("LINEABOVE", (0, len(data) - 1), (-1, len(data) - 1), 0.8, C_SLATE)]
    tb.setStyle(TableStyle(sty))
    out = [CondPageBreak(90)]
    if t["title"]:
        out.append(Paragraph(esc(t["title"]).upper(), ST["sub"]))
    out.append(tb)
    nota = t["note"] + (f"  ·  PDF limitado a {PDF_MAX} filas; el Excel trae todas." if len(t["rows"]) > PDF_MAX else "")
    if nota.strip():
        out.append(Paragraph(esc(nota.strip(" ·")), ST["nota"]))
    out.append(Spacer(1, 10))
    return out


def _pdf_bars(b, W, label_w=170, bar_h=11, gap=4):
    items = b["items"]
    h = len(items) * (bar_h + gap) + gap
    d = Drawing(W, h)
    vmax = max([v for _, v, _ in items] + [1e-9])
    area = W - label_w - 70
    y = h - gap - bar_h
    pal = {k: colors.HexColor(v) for k, v in HEX.items()}
    for lab, v, tono in items:
        d.add(String(label_w - 6, y + bar_h / 2 - 2.5, trunc(lab, int(label_w / 4.1)), fontName="Helvetica", fontSize=7,
                     textAnchor="end", fillColor=C_TEXT))
        w = max(area * v / vmax, 0.8) if v > 0 else 0
        d.add(Rect(label_w, y, w, bar_h, fillColor=pal.get(tono, C_TEAL), strokeColor=None))
        d.add(String(label_w + w + 4, y + bar_h / 2 - 2.5, fmt(b["kind"], v)[0], fontName="Helvetica-Bold", fontSize=7, fillColor=C_TEXT))
        y -= bar_h + gap
    return [Paragraph(esc(b["title"]).upper(), ST["sub"]), d, Spacer(1, 10)]


def generar_pdf(secciones, resumen_filtros, titulo, logo_bytes, orientacion):
    hay_ancho = any(k == "table" and p.get("landscape") for _, _, _, els in secciones for k, p in els)
    horiz = orientacion == "HORIZONTAL" or (orientacion == "AUTOMÁTICA" and hay_ancho)
    pw, ph = landscape(letter) if horiz else letter
    W = pw - 72
    ahora = datetime.now().strftime("%d/%m/%Y %H:%M")
    BH = 120

    def portada(c, doc):
        c.saveState()
        c.setFillColor(C_NAVY)
        c.rect(0, ph - BH, pw, BH, stroke=0, fill=1)
        c.setFillColor(C_TEAL)
        c.rect(0, ph - BH, pw, 5, stroke=0, fill=1)
        c.setFillColor(C_GOLD)
        c.rect(36, ph - 52, 38, 3, stroke=0, fill=1)
        c.setFillColor(colors.HexColor("#9FB3BF"))
        c.setFont("Helvetica-Bold", 9)
        c.drawString(36, ph - 42, "JYPESA  |  LOGÍSTICA Y DISTRIBUCIÓN")
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 24)
        c.drawString(36, ph - 80, "REPORTE DE ANÁLISIS")
        c.setFillColor(C_GOLD)
        c.setFont("Helvetica-Bold", 11)
        c.drawString(36, ph - 100, titulo)
        if logo_bytes:
            try:
                c.drawImage(ImageReader(BytesIO(logo_bytes)), pw - 36 - 110, ph - 84, width=102, height=48,
                            preserveAspectRatio=True, mask="auto", anchor="c")
            except Exception:
                pass
        c.restoreState()

    def normal(c, doc):
        c.saveState()
        c.setFillColor(C_NAVY)
        c.rect(0, ph - 30, pw, 30, stroke=0, fill=1)
        c.setFillColor(C_TEAL)
        c.rect(0, ph - 30, pw, 2.5, stroke=0, fill=1)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 8.5)
        c.drawString(36, ph - 19, "REPORTE DE ANÁLISIS · CONSULTAS")
        c.setFillColor(C_GOLD)
        c.drawRightString(pw - 36, ph - 19, titulo)
        c.restoreState()

    buf = BytesIO()
    doc = BaseDocTemplate(buf, pagesize=(pw, ph), leftMargin=36, rightMargin=36, topMargin=48, bottomMargin=44,
                          title="Reporte de análisis · Consultas", author="JYPESA | Nexion")
    doc.addPageTemplates([
        PageTemplate(id="portada", frames=[Frame(36, 44, W, ph - BH - 44 - 14, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)], onPage=portada),
        PageTemplate(id="normal", frames=[Frame(36, 44, W, ph - 44 - 48, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)], onPage=normal)])
    caja = Table([[Paragraph("<b>FILTROS APLICADOS</b><br/>" + esc(resumen_filtros) + f"<br/><font color='#6B7A86'>Generado el {ahora}</font>", ST["body"])]],
                 colWidths=[W])
    caja.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), C_LIGHT), ("BOX", (0, 0), (-1, -1), 0.5, C_BORDER),
                              ("LINEBEFORE", (0, 0), (0, 0), 4, C_GOLD), ("LEFTPADDING", (0, 0), (-1, -1), 12),
                              ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story = [NextPageTemplate("normal"), caja, Spacer(1, 12)]
    for i, (key, nombre, desc, els) in enumerate(secciones, 1):
        story += _pdf_sec(i, nombre, desc, W)
        for kind, p in els:
            if kind == "kpis":
                story += _pdf_kpis(p, W)
            elif kind == "table":
                story += _pdf_table(p, W, wide=horiz)
            elif kind == "bars":
                story += _pdf_bars(p, W)
            else:
                story += [Paragraph(esc(p), ST["nota"]), Spacer(1, 8)]
    if not secciones:
        story.append(Paragraph("No se eligió ningún bloque.", ST["nota"]))
    cv = type("CV", (_Canvas,), {"titulo_pie": f"Consultas · {titulo}", "pw": pw})
    doc.build(story, canvasmaker=cv)
    buf.seek(0)
    return buf


def generar_excel(secciones, resumen_filtros):
    buf = BytesIO()
    usados = set()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame({"FILTROS APLICADOS": [resumen_filtros], "GENERADO": [datetime.now().strftime("%d/%m/%Y %H:%M")]}).to_excel(
            xw, sheet_name="Filtros", index=False)
        for _, nombre, _, els in secciones:
            for kind, p in els:
                if kind != "table":
                    continue
                nom = "".join(ch for ch in f"{nombre[:12]} {p['title'][:16]}" if ch not in '[]:*?/\\')[:31] or "Hoja"
                base, k = nom, 2
                while nom in usados:
                    nom = f"{base[:28]}_{k}"
                    k += 1
                usados.add(nom)
                df = pd.DataFrame(p["rows"] + ([p["total"]] if p["total"] else []))
                if df.empty:
                    continue
                df = df.reindex(columns=[c[1] for c in p["cols"]])
                df.columns = [c[0] for c in p["cols"]]
                for c in p["cols"]:
                    if c[2] == "date":
                        df[c[0]] = pd.to_datetime(df[c[0]], errors="coerce").dt.strftime("%d/%m/%Y")
                df.to_excel(xw, sheet_name=nom, index=False)
                ws = xw.sheets[nom]
                ws.freeze_panes = "B2"
                ws.auto_filter.ref = ws.dimensions
                for i, col in enumerate(df.columns, 1):
                    ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(len(str(col)) + 3, 12), 34)
    buf.seek(0)
    return buf



# ============================================================
# 10. INTERFAZ
# ============================================================
def _sel_todos(prefix, keys):
    for k in keys:
        st.session_state[f"{prefix}{k}"] = True


def _preset(nombre):
    act = set(PRESETS[nombre])
    for k, _, _ in SECCIONES:
        st.session_state[f"sec_{k}"] = k in act


def _limpiar():
    for k, _, _ in SECCIONES:
        st.session_state[f"sec_{k}"] = False


def ui_dinamicos():
    try:
        import pytz
        hoy = pd.Timestamp(datetime.now(pytz.timezone("America/Mexico_City")).date())
    except Exception:
        hoy = pd.Timestamp(datetime.now().date())

    with st.spinner("Cargando base de envíos..."):
        try:
            base = datos_base()
        except Exception as e:
            st.error(f"No se pudo cargar la matriz de envíos: {e}")
            return

    anios_disp = sorted({int(y) for y in base["FECHA DE ENVÍO"].dt.year.dropna().unique()} | {hoy.year}, reverse=True)
    mes_prev = hoy.month - 1 or 12
    anio_prev = hoy.year if hoy.month > 1 else hoy.year - 1

    # ---------- 1. FILTROS ----------
    st.markdown("##### 1 · FILTROS")
    c1, c2, c3, c4 = st.columns([1, 1.6, 1.2, 1.2])
    anios = c1.multiselect("AÑO(S)", anios_disp, default=[anio_prev if anio_prev in anios_disp else anios_disp[0]], key="q_anios")
    meses_sel = c2.multiselect("MES(ES) (vacío = todo el año)", MESES, default=[MESES[mes_prev - 1]], key="q_meses")
    modalidad = c3.selectbox("COBRO", ["TODAS", "COBRO REGRESO", "COBRO DESTINO"], key="q_mod")
    costo_sel = c4.selectbox("COSTO A CONSIDERAR", ["Guía + adicionales", "Solo guía"], key="q_costo")
    if not anios:
        st.warning("Elige al menos un año.")
        return
    incluir_adic = costo_sel == "Guía + adicionales"

    d1, d2, d3, d4 = st.columns([1.3, 1, 1, 1])
    cajas_modo_txt = d1.selectbox("PEDIDOS POR NÚMERO DE CAJAS", ["TODOS", "PEQUEÑOS (1 a N)", "GRANDES (más de N)", "RANGO PERSONALIZADO"], key="q_cajas")
    max_peq = d2.selectbox("PEQUEÑO = HASTA (CAJAS)", [2, 3, 4, 5, 6, 8, 10], index=2, key="q_maxpeq")
    caj_min = d3.number_input("CAJAS DESDE", min_value=0, value=1, step=1, key="q_cmin", disabled=cajas_modo_txt != "RANGO PERSONALIZADO")
    caj_max = d4.number_input("CAJAS HASTA (0 = sin límite)", min_value=0, value=0, step=1, key="q_cmax", disabled=cajas_modo_txt != "RANGO PERSONALIZADO")
    cajas_modo = {"PEQUEÑOS (1 a N)": "PEQUEÑOS", "GRANDES (más de N)": "GRANDES", "RANGO PERSONALIZADO": "RANGO"}.get(cajas_modo_txt, "TODOS")

    # Opciones dependientes del año elegido
    u = preparar_universo(base, incluir_adic, hoy, int(max_peq))
    uy = u[u["FECHA DE ENVÍO"].isna() | u["_ANIO"].isin(anios)]
    with st.expander("MÁS FILTROS: fletera, transporte, destino, cliente, concepto, entrega, semáforo, búsqueda", expanded=False):
        e1, e2, e3 = st.columns(3)
        fleteras = e1.multiselect("FLETERA (vacío = todas)", sorted(uy["_FLE"].unique()), key=f"q_fle_{'_'.join(map(str, anios))}")
        transportes = e2.multiselect("TRANSPORTE", sorted(uy["_TR"].unique()), key=f"q_tr_{'_'.join(map(str, anios))}")
        destinos = e3.multiselect("DESTINO", sorted(uy["_DE"].unique()), key=f"q_de_{'_'.join(map(str, anios))}")
        g1, g2, g3 = st.columns(3)
        clientes = g1.multiselect("CLIENTE", sorted(uy.loc[~uy["ES_EXTRA"], "_CLI"].unique()), key=f"q_cli_{'_'.join(map(str, anios))}")
        conceptos = g2.multiselect("CONCEPTO", sorted(uy["_CONC"].unique()), key=f"q_con_{'_'.join(map(str, anios))}")
        estatus = g3.multiselect("ESTATUS DE ENTREGA", ESTATUS, key="q_est")
        h1, h2, h3 = st.columns([1.3, 1, 1.2])
        clases = h1.multiselect("SEMÁFORO", CLASES, key="q_cla")
        pct_min = h2.number_input("% LOGÍSTICO MAYOR A (0 = sin filtro)", min_value=0.0, value=0.0, step=0.5, key="q_pmin")
        buscar = h3.text_input("BUSCAR FACTURA / CLIENTE", key="q_buscar")
        solo_fuera = st.checkbox("Solo pedidos fuera de target (en alerta, críticos, en pérdida o sin facturación)", key="q_fuera")

    f = dict(anios=[int(a) for a in anios], meses=[MESES.index(m) + 1 for m in meses_sel], modalidad=modalidad,
             incluir_adic=incluir_adic, cajas_modo=cajas_modo, max_peq=int(max_peq), caj_min=int(caj_min), caj_max=int(caj_max),
             fleteras=fleteras, transportes=transportes, destinos=destinos, clientes=clientes, conceptos=conceptos,
             estatus=estatus, clases=clases, pct_min=float(pct_min), buscar=buscar.strip(), solo_fuera=solo_fuera)
    d, ext, d_nocaj = aplicar_filtros(u, f)

    # ---------- 2. BLOQUES ----------
    st.markdown("##### 2 · ARMA TU REPORTE (marca los bloques que quieras)")
    for k, _, _ in SECCIONES:
        st.session_state.setdefault(f"sec_{k}", k in PRESETS["Ejecutivo"])
    b1, b2, b3, b4 = st.columns(4)
    b1.button("Ejecutivo", on_click=_preset, args=("Ejecutivo",), key="p1")
    b2.button("Costos", on_click=_preset, args=("Costos",), key="p2")
    b3.button("Todo", on_click=_preset, args=("Todo",), key="p3")
    b4.button("Limpiar", on_click=_limpiar, key="p4")
    cols_chk = st.columns(3)
    for i, (k, nombre, desc) in enumerate(SECCIONES):
        cols_chk[i % 3].checkbox(nombre, key=f"sec_{k}", help=desc)
    activas = [k for k, _, _ in SECCIONES if st.session_state.get(f"sec_{k}")]

    # ---------- 3. OPCIONES ----------
    o = dict(orden_met="Costo logístico ($)", orden_dir="MAYOR A MENOR", top=15, rank_orden="% logístico", rank_dir="MAYOR A MENOR",
             rank_top=50, cz_filas="Fletera", cz_cols="Mes", cz_met="% logístico", mu_sol=[], mu_detalle=False)
    with st.expander("3 · OPCIONES DE LOS BLOQUES (orden, top, pedido por pedido, cruce, muestras)", expanded=False):
        a1, a2, a3 = st.columns(3)
        o["orden_met"] = a1.selectbox("ORDENAR TABLAS Y GRÁFICAS POR", list(METRICAS), key="o_met")
        o["orden_dir"] = a2.selectbox("DIRECCIÓN", ["MAYOR A MENOR", "MENOR A MAYOR"], key="o_dir")
        o["top"] = a3.selectbox("TOP EN TABLAS AGRUPADAS", TOPS, index=2, key="o_top")
        if "rank" in activas:
            r1, r2, r3 = st.columns(3)
            o["rank_orden"] = r1.selectbox("PEDIDO POR PEDIDO · ORDENAR POR", list(ORDEN_RANK), key="o_ro")
            o["rank_dir"] = r2.selectbox("PEDIDO POR PEDIDO · DIRECCIÓN", ["MAYOR A MENOR", "MENOR A MAYOR"], key="o_rd")
            o["rank_top"] = r3.selectbox("PEDIDO POR PEDIDO · CUÁNTOS", [25, 50, 100, 200, 500, "TODOS"], index=1, key="o_rt")
        if "cruz" in activas:
            z1, z2, z3 = st.columns(3)
            o["cz_filas"] = z1.selectbox("CRUCE · FILAS", list(DIMS), index=0, key="o_zf")
            o["cz_cols"] = z2.selectbox("CRUCE · COLUMNAS", list(DIMS), index=list(DIMS).index("Mes"), key="o_zc")
            o["cz_met"] = z3.selectbox("CRUCE · MÉTRICA", list(CRUCE_MET), index=1, key="o_zm")
        if "mues" in activas:
            dm, _msg = datos_muestras()
            m1, m2 = st.columns([2, 1])
            sol_opts = sorted(dm["_SOL"].unique()) if dm is not None else []
            o["mu_sol"] = m1.multiselect("MUESTRAS · SOLICITANTE (vacío = todos)", sol_opts, key="o_msol")
            o["mu_detalle"] = m2.checkbox("MUESTRAS · incluir detalle por folio", key="o_mdet")

    ctx = dict(d=d, ext=ext, d_nocaj=d_nocaj, f=f, opts=o)
    resumen = texto_filtros(f)
    if not activas:
        st.info("Marca al menos un bloque para armar el reporte.")
        return
    if d.empty and ext.empty and "mues" not in activas:
        st.warning("No hay información con esos filtros. Prueba con otro mes, año o quita algún filtro.")
        return

    with st.spinner("Armando el reporte..."):
        secciones = construir(ctx, activas)

    # ---------- IMPRESIÓN ----------
    st.markdown("##### 3 · IMPRIME")
    p1, p2, p3 = st.columns([1, 1.6, 1.6], vertical_alignment="bottom")
    orient = p1.selectbox("ORIENTACIÓN DEL PDF", ["AUTOMÁTICA", "VERTICAL", "HORIZONTAL"], key="q_orient")
    firma = hashlib.md5(json.dumps([f, activas, o, orient], sort_keys=True, default=str).encode()).hexdigest()
    gen_pdf = p2.button("GENERAR PDF", key="q_pdf")
    gen_xls = p3.button("GENERAR EXCEL", key="q_xls")
    titulo = (", ".join(m.title() for m in meses_sel) if meses_sel else "Año completo") + " " + "/".join(str(a) for a in f["anios"])
    if gen_pdf:
        with st.spinner("Armando el PDF..."):
            try:
                pdf = generar_pdf(secciones, resumen, titulo, obtener_logo_bytes(), orient)
                st.session_state["q_pdf_out"] = {"bytes": pdf.getvalue(), "firma": firma,
                                                 "nombre": f"Reporte_Consultas_{titulo.replace(' ', '_').replace('/', '-').replace(',', '')}.pdf"}
            except Exception as e:
                st.error(f"No se pudo generar el PDF: {e}")
    if gen_xls:
        with st.spinner("Armando el Excel..."):
            try:
                x = generar_excel(secciones, resumen)
                st.session_state["q_xls_out"] = {"bytes": x.getvalue(), "firma": firma,
                                                 "nombre": f"Consultas_{titulo.replace(' ', '_').replace('/', '-').replace(',', '')}.xlsx"}
            except Exception as e:
                st.error(f"No se pudo generar el Excel: {e}")
    _, dl2, dl3 = st.columns([1, 1.6, 1.6])
    for clave, etiqueta, mime, col in (("q_pdf_out", "DESCARGAR PDF", "application/pdf", dl2),
                                       ("q_xls_out", "DESCARGAR EXCEL", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", dl3)):
        out = st.session_state.get(clave)
        if out and out["firma"] == firma:
            col.download_button(etiqueta, data=out["bytes"], file_name=out["nombre"], mime=mime, key=f"dl_{clave}")
        elif out:
            st.caption(f"Cambiaste filtros o bloques: vuelve a generar el {'PDF' if clave == 'q_pdf_out' else 'Excel'}.")

    st.markdown("---")
    render_pantalla(secciones, resumen, f)


# ============================================================
# PESTAÑA 2 · REPORTES PREDEFINIDOS (nombres propios con prefijo mp_)
# ============================================================

# ============================================================
# 3. CONSTANTES
# ============================================================
GITHUB_USER, GITHUB_REPO, BRANCH = "RH2026", "nexion", "main"
ARCHIVO_MATRIZ = "Matriz_Excel_Dashboard.csv"
ARCHIVO_LOGO = "n1.png"
ARCHIVO_HIST = "Historial2025.csv"   # totales mensuales de 2025: MES, COSTO DE LA GUÍA, CAJAS (cobro regreso)
ANIO_HISTORIAL = 2025

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO",
         "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

# Mismas listas que el dashboard
CARRIERS_PRINCIPALES_DC = ["TRES GUERRAS", "ONE", "TINY PACK", "PAQMEX", "PAQUETE", "SANCHEZ", "FLETES DE REGRESO", "FARMASES"]
FLETERAS_PRINCIPALES_RK = ["TRES GUERRAS", "ONE", "TINY PACK", "PAQMEX", "SANCHEZ", "FLETES DE REGRESO"]
TARGET_COSTO_LOG = 7.5   # % meta de costo logístico (igual que Análisis Mensual)
EXTRA_REGEX = "RECOLECCI|MANIOBRA"   # CONCEPTO que se trata como costo extra (sin acentos, mayúsculas): RECOLECCIONES / MANIOBRAS
CAJAS_PEQUENO_MAX = 4    # un pedido "pequeño" tiene de 1 a N cajas (editable también desde la página)
FERIADOS_24H = ['2026-01-01', '2026-02-02', '2026-03-16', '2026-05-01']   # <- agrega aquí los de otros años

# Paleta (impresión)
mp_C_NAVY = colors.HexColor("#384A52")
mp_C_SLATE = colors.HexColor("#2B343B")
mp_C_TEAL = colors.HexColor("#00A3A3")
mp_C_GOLD = colors.HexColor("#FFC000")
mp_C_GREEN = colors.HexColor("#2E9E6B")
mp_C_RED = colors.HexColor("#D64545")
C_BLUE = colors.HexColor("#3B82F6")
C_PURPLE = colors.HexColor("#7C5CBF")
C_ORANGE = colors.HexColor("#E8833A")
mp_C_GRAY = colors.HexColor("#6B7A86")
mp_C_LIGHT = colors.HexColor("#F3F6F8")
mp_C_BORDER = colors.HexColor("#D5DDE2")
mp_C_TEXT = colors.HexColor("#1F2D35")
mp_HEX = {k: v.hexval().replace("0x", "#") for k, v in dict(
    teal=mp_C_TEAL, green=mp_C_GREEN, red=mp_C_RED, blue=C_BLUE, gold=mp_C_GOLD, purple=C_PURPLE,
    orange=C_ORANGE, gray=mp_C_GRAY, slate=mp_C_SLATE).items()}

PAGE_W, PAGE_H = letter
MARGIN = 36
CONTENT_W = PAGE_W - 2 * MARGIN   # 540


# ============================================================
# 4. UTILIDADES DE FORMATO Y FECHAS
# ============================================================
def esc(t):
    return _html.escape("" if t is None else str(t))


def trunc(t, n):
    t = "" if t is None else str(t)
    return t if len(t) <= n else t[: n - 1] + "."


def money(v, dec=0):
    try:
        return f"${float(v):,.{dec}f}"
    except Exception:
        return "-"


def num(v, dec=0):
    try:
        return f"{float(v):,.{dec}f}"
    except Exception:
        return "-"


def pct(v, dec=1):
    try:
        if v is None or pd.isna(v):
            return "-"
        return f"{float(v):.{dec}f}%"
    except Exception:
        return "-"


def mp_parse_fecha_segura(serie):
    """Convierte fechas sin invertir día/mes (ISO -> año-mes-día; el resto -> día/mes/año)."""
    serie = pd.Series(serie)
    if pd.api.types.is_datetime64_any_dtype(serie):
        return serie
    s_txt = serie.fillna("").astype(str).str.strip()
    es_iso = s_txt.str.match(r"^\d{4}-\d{1,2}-\d{1,2}")

    def _conv(x, **kw):
        try:
            return pd.to_datetime(x, errors="coerce", format="mixed", **kw)
        except (TypeError, ValueError):
            return pd.to_datetime(x, errors="coerce", **kw)

    out = pd.Series(pd.NaT, index=s_txt.index, dtype="datetime64[ns]")
    if es_iso.any():
        out.loc[es_iso] = _conv(s_txt[es_iso]).astype("datetime64[ns]")
    if (~es_iso).any():
        out.loc[~es_iso] = _conv(s_txt[~es_iso], dayfirst=True).astype("datetime64[ns]")
    return out


def mp_limpiar_moneda(serie):
    return pd.to_numeric(
        serie.astype(str).str.replace(r"[^\d\.\-]", "", regex=True).replace("", "0"),
        errors="coerce",
    ).fillna(0.0)


# ============================================================
# 5. CARGA Y PREPARACIÓN DE DATOS
# ============================================================
def _headers_github():
    try:
        token = st.secrets.get("GITHUB_TOKEN", None)
    except Exception:
        token = None
    return {"Authorization": f"token {token}"} if token else {}


@st.cache_data(ttl=300)
def cargar_matriz():
    url = f"https://raw.githubusercontent.com/{GITHUB_USER}/{GITHUB_REPO}/{BRANCH}/{ARCHIVO_MATRIZ}?v={int(time.time())}"
    try:
        r = requests.get(url, headers=_headers_github(), timeout=40)
        r.raise_for_status()
        df = pd.read_csv(BytesIO(r.content), encoding="utf-8-sig")
    except Exception:
        df = pd.read_csv(url, encoding="utf-8-sig")
    df.columns = [str(c).strip() for c in df.columns]
    return df


@st.cache_data(ttl=3600)
def cargar_historial_2025():
    """Totales mensuales de 2025 (para el comparativo anual). Devuelve None si no se puede leer."""
    url = f"https://raw.githubusercontent.com/{GITHUB_USER}/{GITHUB_REPO}/{BRANCH}/{ARCHIVO_HIST}?v={int(time.time())}"
    try:
        try:
            r = requests.get(url, headers=_headers_github(), timeout=40)
            r.raise_for_status()
            df = pd.read_csv(BytesIO(r.content), encoding="utf-8-sig")
        except Exception:
            df = pd.read_csv(url, encoding="utf-8-sig")
        df.columns = ["".join(c for c in unicodedata.normalize("NFD", str(x)) if unicodedata.category(c) != "Mn").strip().upper()
                      for x in df.columns]
        df["MES"] = df["MES"].fillna("").astype(str).str.strip().str.upper()
        for c in ["COSTO DE LA GUIA", "CAJAS"]:
            df[c] = mp_limpiar_moneda(df[c]) if c in df.columns else 0.0
        return df[["MES", "COSTO DE LA GUIA", "CAJAS"]]
    except Exception:
        return None


@st.cache_data(ttl=3600)
def mp_obtener_logo_bytes():
    try:
        if os.path.exists(ARCHIVO_LOGO):
            with open(ARCHIVO_LOGO, "rb") as f:
                return f.read()
        url = f"https://raw.githubusercontent.com/{GITHUB_USER}/{GITHUB_REPO}/{BRANCH}/{ARCHIVO_LOGO}"
        r = requests.get(url, headers=_headers_github(), timeout=20)
        return r.content if r.status_code == 200 else None
    except Exception:
        return None


def mp_preparar_base(df_raw):
    """Normaliza la matriz de envíos: fechas, textos y números."""
    df = df_raw.copy()
    df.columns = [str(c).strip() for c in df.columns]

    for c in ["FECHA DE ENVÍO", "PROMESA DE ENTREGA", "FECHA DE ENTREGA REAL", "EMISION"]:
        df[c] = mp_parse_fecha_segura(df[c]) if c in df.columns else pd.Series(pd.NaT, index=df.index, dtype="datetime64[ns]")

    for c in ["FLETERA", "FORMA DE ENVIO", "TRANSPORTE", "DESTINO", "NOMBRE DEL CLIENTE",
              "INCIDENCIAS", "MES", "NÚMERO DE PEDIDO"]:
        if c not in df.columns:
            df[c] = ""
        df[c] = df[c].fillna("").astype(str).str.strip()
        df.loc[df[c].str.lower() == "nan", c] = ""
    df["MES"] = df["MES"].str.upper()

    for c in ["COSTO DE LA GUÍA", "FACTURACION", "VALUACION", "COSTOS ADICIONALES", "CANTIDAD DE CAJAS"]:
        df[c] = mp_limpiar_moneda(df[c]) if c in df.columns else 0.0
    df["CAJAS"] = pd.to_numeric(df["CAJAS"], errors="coerce").fillna(0) if "CAJAS" in df.columns else 0.0

    # Registros de recolecciones / maniobras: no son pedidos, son costos extras
    col_c = next((c for c in df.columns if "CONCEPTO" in str(c).upper()), None)
    if col_c:
        conc = df[col_c].fillna("").astype(str).map(
            lambda x: "".join(ch for ch in unicodedata.normalize("NFD", x) if unicodedata.category(ch) != "Mn")).str.strip().str.upper()
        df["ES_EXTRA"] = conc.str.contains(EXTRA_REGEX, regex=True).astype(bool)
    else:
        df["ES_EXTRA"] = False
    return df


def filtrar_mes(df, col, anio, mes):
    f = df[col]
    return df[(f.dt.year == anio) & (f.dt.month == mes)].copy()


# ============================================================
# 6. CÁLCULOS POR SECCIÓN (replican el dashboard)
# ============================================================
def calc_resumen(df_mes, hoy):
    total_p = len(df_mes)
    entregados = int(df_mes["FECHA DE ENTREGA REAL"].notna().sum())
    df_trans = df_mes[df_mes["FECHA DE ENTREGA REAL"].isna()]
    en_tiempo = int((df_trans["PROMESA DE ENTREGA"] >= hoy).sum())
    retrasados = int((df_trans["PROMESA DE ENTREGA"] < hoy).sum())

    ret = df_trans[df_trans["PROMESA DE ENTREGA"] < hoy].copy()
    ret["FLETERA"] = ret["FLETERA"].replace("", "SIN ASIGNAR")
    ret_fletera = ret.groupby("FLETERA").size().sort_values(ascending=False)

    dest = df_mes[df_mes["DESTINO"] != ""]["DESTINO"].value_counts().head(8)
    return dict(total=total_p, entregados=entregados, en_transito=len(df_trans), en_tiempo=en_tiempo,
                retrasados=retrasados, pct_entregado=(entregados / total_p * 100) if total_p else 0.0,
                ret_fletera=ret_fletera, destinos=dest)


def _kpi_24h(ini, fin, feriados):
    if pd.isna(ini) and not pd.isna(fin):
        ini = fin
    if pd.isna(ini) or pd.isna(fin):
        return "Sin Datos", None
    try:
        if fin <= ini:
            return "A Tiempo", 0
        d = int(np.busday_count(ini.date(), fin.date(), weekmask='1111100', holidays=feriados))
        if d == 0:
            return "A Tiempo", d
        if d == 1 and fin.time() <= ini.time():
            return "A Tiempo", d
        return "Fuera de Tiempo", d
    except Exception:
        return "Sin Datos", None


def calc_despachos(df_mes):
    d = df_mes.copy()
    feriados = np.array(FERIADOS_24H, dtype='datetime64[D]')
    estados, dias = [], []
    for ini, fin in zip(d["EMISION"], d["FECHA DE ENVÍO"]):
        e, n = _kpi_24h(ini, fin, feriados)
        estados.append(e)
        dias.append(n)
    d["Estado_KPI"], d["DIAS_HABILES"] = estados, dias

    validos = d[d["Estado_KPI"] != "Sin Datos"]
    tot = len(validos)
    ok = int((validos["Estado_KPI"] == "A Tiempo").sum())
    no = tot - ok

    por_dia = pd.DataFrame()
    if not validos.empty:
        v = validos.assign(_ORDEN=validos["FECHA DE ENVÍO"].dt.normalize())
        por_dia = (v.groupby(["_ORDEN", "Estado_KPI"]).size().unstack(fill_value=0).sort_index())
        for c in ["A Tiempo", "Fuera de Tiempo"]:
            if c not in por_dia.columns:
                por_dia[c] = 0
    fuera = (validos[validos["Estado_KPI"] == "Fuera de Tiempo"]
             .sort_values("DIAS_HABILES", ascending=False).head(12))
    return dict(total=tot, ok=ok, no=no, pct_ok=(ok / tot * 100) if tot else 0.0,
                pct_no=(no / tot * 100) if tot else 0.0, por_dia=por_dia, fuera=fuera,
                sin_datos=len(d) - tot)


def calc_bi(df_mes):
    d = df_mes
    fact, guias = d["FACTURACION"].sum(), d["COSTO DE LA GUÍA"].sum()
    adic, cajas = d["COSTOS ADICIONALES"].sum(), d["CANTIDAD DE CAJAS"].sum()
    por_f = pd.DataFrame()
    dd = d[d["FLETERA"] != ""]
    if not dd.empty:
        por_f = dd.groupby("FLETERA").agg(
            pedidos=("FLETERA", "size"), facturacion=("FACTURACION", "sum"),
            guias=("COSTO DE LA GUÍA", "sum"), adic=("COSTOS ADICIONALES", "sum"),
            cajas=("CANTIDAD DE CAJAS", "sum")).reset_index()
        por_f["operativo"] = por_f["guias"] + por_f["adic"]
        por_f["pct_log"] = np.where(por_f["facturacion"] > 0, por_f["guias"] / por_f["facturacion"] * 100, 0.0)
        por_f["costo_caja"] = np.where(por_f["cajas"] > 0, por_f["operativo"] / por_f["cajas"], np.nan)
        por_f = por_f.sort_values("facturacion", ascending=False)
    return dict(fact=fact, guias=guias, adic=adic, cajas=cajas,
                costo_log=(guias / fact * 100) if fact else 0.0,
                costo_caja=(guias / cajas) if cajas else 0.0, por_fletera=por_f)


def calc_top_clientes(df_mes, n=20):
    d = df_mes[df_mes["NOMBRE DEL CLIENTE"] != ""]
    total_fact = df_mes["FACTURACION"].sum()
    top = pd.DataFrame()
    if not d.empty:
        top = d.groupby("NOMBRE DEL CLIENTE").agg(
            pedidos=("NOMBRE DEL CLIENTE", "size"), cajas=("CANTIDAD DE CAJAS", "sum"),
            facturacion=("FACTURACION", "sum"), guias=("COSTO DE LA GUÍA", "sum")).reset_index()
        top = top.sort_values("facturacion", ascending=False).head(n)
        top["share"] = np.where(total_fact > 0, top["facturacion"] / total_fact * 100, 0.0)
    forma = df_mes[df_mes["FORMA DE ENVIO"] != ""]["FORMA DE ENVIO"].value_counts()
    return dict(top=top, forma=forma, total_fact=total_fact,
                share_top=(top["facturacion"].sum() / total_fact * 100) if (total_fact and not top.empty) else 0.0)


def calc_carga_regreso(df_all, anio, mes):
    """Distribución de carga, solo COBRO REGRESO (misma lógica que la pestaña del dashboard)."""
    d = df_all.copy()
    fletera_base = d["FLETERA"].str.upper().str.strip()
    sin_asignar = (d["TRANSPORTE"].str.upper().str.strip() == "COBRO REGRESO") | (fletera_base == "COBRO REGRESO")
    d = d[~sin_asignar]
    texto_carrier = d["TRANSPORTE"].str.upper() + " | " + d["FLETERA"].str.upper()
    d = d[texto_carrier.apply(lambda x: any(p in x for p in CARRIERS_PRINCIPALES_DC))]

    # Periodo: el dashboard usa la columna MES (texto); aquí además se valida el año con la fecha de envío
    f_env = d["FECHA DE ENVÍO"]
    ok_anio = f_env.isna() | (f_env.dt.year == anio)
    d = d[(d["MES"] == MESES[mes - 1]) & ok_anio & (d["TRANSPORTE"] != "")].copy()
    d = d[d["FORMA DE ENVIO"].str.contains("REGRESO", case=False, na=False)
          | d["TRANSPORTE"].str.contains("COBRO REGRESO", case=False, na=False)]
    if d.empty:
        return dict(vacio=True)

    total = d["CAJAS"].sum()
    part = d.groupby("TRANSPORTE", as_index=False)["CAJAS"].sum()
    part["PCT"] = (part["CAJAS"] / total * 100) if total else 0.0
    part = part.sort_values("CAJAS", ascending=False)
    rutas = (d.groupby(["TRANSPORTE", "DESTINO", "FORMA DE ENVIO"], as_index=False)["CAJAS"].sum())
    orden = part.set_index("TRANSPORTE")["CAJAS"]
    rutas["_O"] = rutas["TRANSPORTE"].map(orden)
    rutas = rutas.sort_values(["_O", "TRANSPORTE", "CAJAS"], ascending=[False, True, False]).drop(columns="_O")
    return dict(vacio=False, registros=len(d), total=total, part=part, rutas=rutas,
                lider=part.iloc[0], destinos=d["DESTINO"].replace("", pd.NA).nunique())


def calc_costos_regreso(df_all, anio, mes):
    """Costos de COBRO REGRESO (lo que paga JYPESA), por fletera.
    Misma lógica que la página 'Análisis Mensual': FORMA DE ENVIO contiene REGRESO,
    flete = guía + costos adicionales, costo logístico = flete / facturación.
    Periodo: igual que la sección 05 (columna MES + validación de año)."""
    d = df_all.copy()
    f_env = d["FECHA DE ENVÍO"]
    ok_anio = f_env.isna() | (f_env.dt.year == anio)
    d = d[(d["MES"] == MESES[mes - 1]) & ok_anio]
    d = d[d["FORMA DE ENVIO"].str.contains("REGRESO", case=False, na=False)].copy()
    if d.empty:
        return dict(vacio=True)

    d["FLETERA"] = d["FLETERA"].replace("", "SIN ASIGNAR")
    d["_FLETE"] = d["COSTO DE LA GUÍA"] + d["COSTOS ADICIONALES"]
    d["_EVAL"] = d["PROMESA DE ENTREGA"].notna() & d["FECHA DE ENTREGA REAL"].notna()
    d["_OK"] = d["_EVAL"] & (d["FECHA DE ENTREGA REAL"] <= d["PROMESA DE ENTREGA"])
    d["_INC"] = d["VALUACION"] > 0

    flete, fact, cajas = d["_FLETE"].sum(), d["FACTURACION"].sum(), d["CAJAS"].sum()
    val, n_eval, n_ok = d["VALUACION"].sum(), int(d["_EVAL"].sum()), int(d["_OK"].sum())

    # Desglose por concepto (informativo; usa solo el costo de guía, como el dashboard)
    consignas = fnacional = 0.0
    col_c = next((c for c in d.columns if "CONCEPTO" in str(c).upper()), None)
    if col_c:
        conc = d[col_c].fillna("").astype(str).str.strip().str.upper()
        g = d["COSTO DE LA GUÍA"]
        consignas = g[conc.str.contains("CONSIGNA", regex=True)].sum()
        fnacional = g[conc.str.contains("NACIONAL", regex=True)].sum()

    res = d.groupby("FLETERA").agg(
        envios=("FLETERA", "size"), cajas=("CAJAS", "sum"), flete=("_FLETE", "sum"),
        fact=("FACTURACION", "sum"), val=("VALUACION", "sum"),
        n_eval=("_EVAL", "sum"), n_ok=("_OK", "sum")).reset_index()
    res["pct_gasto"] = (res["flete"] / flete * 100) if flete else 0.0
    res["costo_caja"] = res["flete"] / res["cajas"].replace(0, np.nan)
    res["pct_log"] = res["flete"] / res["fact"].replace(0, np.nan) * 100
    res["efic"] = res["n_ok"] / res["n_eval"].replace(0, np.nan) * 100
    res = res.sort_values("flete", ascending=False).reset_index(drop=True)

    return dict(vacio=False, registros=len(d), flete=flete, fact=fact, cajas=cajas, val=val,
                costo_log=(flete / fact * 100) if fact else 0.0,
                costo_caja=(flete / cajas) if cajas else 0.0,
                efic=(n_ok / n_eval * 100) if n_eval else None,
                pct_inc=(d["_INC"].sum() / len(d) * 100),
                consignas=consignas, fnacional=fnacional, res=res)


def calc_comparativa_2025(G, df_hist, anio, mes):
    """Flete y cajas de cobro regreso contra el mismo mes de 2025 (misma lógica que 'Análisis Mensual':
    flete 2026 = guía + adicionales; flete 2025 = guía). None si no aplica; dict(sin_dato=True) si falta información."""
    if G.get("vacio") or anio != ANIO_HISTORIAL + 1:
        return None
    nombre = MESES[mes - 1].title()
    if df_hist is None or df_hist.empty:
        return dict(sin_dato=True, motivo=f"No se pudo cargar {ARCHIVO_HIST}; no hay comparativo contra {ANIO_HISTORIAL}.")
    f = df_hist[df_hist["MES"] == MESES[mes - 1]]
    flete25, cajas25 = float(f["COSTO DE LA GUIA"].sum()), float(f["CAJAS"].sum())
    if f.empty or (flete25 <= 0 and cajas25 <= 0):
        return dict(sin_dato=True, motivo=f"{ARCHIVO_HIST} no tiene datos de {nombre} {ANIO_HISTORIAL}.")
    return dict(sin_dato=False, flete25=flete25, cajas25=cajas25,
                dif_flete=G.get("flete_total", G["flete"]) - flete25, dif_cajas=G.get("cajas_total", G["cajas"]) - cajas25,
                var_flete=((G.get("flete_total", G["flete"]) - flete25) / flete25 * 100) if flete25 > 0 else None,
                var_cajas=((G.get("cajas_total", G["cajas"]) - cajas25) / cajas25 * 100) if cajas25 > 0 else None)


def _separar_extras(df):
    """Devuelve (pedidos, extras): extras = registros con CONCEPTO de recolecciones / maniobras."""
    if "ES_EXTRA" not in df.columns:
        return df, df.iloc[0:0]
    m = df["ES_EXTRA"].astype(bool)
    return df[~m].copy(), df[m].copy()


def _resumen_extras(d):
    """Resume registros extras ya filtrados por periodo/modalidad: costo = guía + adicionales."""
    r0 = dict(n=0, guia=0.0, adic=0.0, costo=0.0, cajas=0.0, por_concepto=pd.DataFrame())
    if d is None or d.empty:
        return r0
    d = d.copy()
    col_c = next((c for c in d.columns if "CONCEPTO" in str(c).upper()), None)
    d["_CONC"] = (d[col_c].fillna("").astype(str).map(lambda x: _sin_acentos(x).strip().upper()).replace("", "SIN CONCEPTO")
                  if col_c else "SIN CONCEPTO")
    d["_COSTO"] = d["COSTO DE LA GUÍA"] + d["COSTOS ADICIONALES"]
    pc = d.groupby("_CONC").agg(n=("_CONC", "size"), guia=("COSTO DE LA GUÍA", "sum"), adic=("COSTOS ADICIONALES", "sum"),
                                costo=("_COSTO", "sum"), cajas=("CAJAS", "sum")).reset_index().rename(columns={"_CONC": "CONCEPTO"})
    pc = pc.sort_values("costo", ascending=False).reset_index(drop=True)
    return dict(n=len(d), guia=float(d["COSTO DE LA GUÍA"].sum()), adic=float(d["COSTOS ADICIONALES"].sum()),
                costo=float(d["_COSTO"].sum()), cajas=float(d["CAJAS"].sum()), por_concepto=pc)


def calc_extras(df_extras, anio, mes, modalidad="TODAS"):
    """Costos extras (recolecciones / maniobras) del periodo y modalidad, con el mismo filtro que Análisis Mensual (MES + año)."""
    if df_extras is None or df_extras.empty:
        return _resumen_extras(None)
    return _resumen_extras(_filtrar_periodo_modalidad(df_extras, anio, mes, modalidad))


def _resumen_por_fletera(df_in):
    filas = []
    for fletera, g in df_in.groupby("FLETERA"):
        entregados = int(g["_ENTREGADO"].sum())
        a_tiempo = int(g["_A_TIEMPO"].sum())
        dias_ok = g.loc[g["_ENTREGADO"] & (g["_DIAS"] >= 0), "_DIAS"]
        cajas = g["CANTIDAD DE CAJAS"].sum()
        costo = g["_COSTO"].sum()
        filas.append({
            "FLETERA": fletera, "ENVIOS": len(g), "ENTREGADOS": entregados, "A_TIEMPO": a_tiempo,
            "RETRASO": max(entregados - a_tiempo, 0),
            "PCT_A_TIEMPO": (a_tiempo / entregados * 100) if entregados else None,
            "INCIDENCIAS_PCT": (g["_INCIDENCIA"].sum() / len(g) * 100) if len(g) else 0.0,
            "DIAS_PROM": dias_ok.mean() if not dias_ok.empty else None,
            "COSTO_PROM_ENVIO": g["_COSTO"].mean() if len(g) else 0.0,
            "COSTO_PROM_CAJA": (costo / cajas) if cajas else None,
        })
    return pd.DataFrame(filas)


def calc_ranking(df_all, anio, mes):
    d = df_all[df_all["FLETERA"].str.upper().apply(lambda x: any(p in x for p in FLETERAS_PRINCIPALES_RK))]
    d = filtrar_mes(d, "FECHA DE ENVÍO", anio, mes)
    d = d[d["FLETERA"] != ""].copy()
    if d.empty:
        return dict(vacio=True)
    d["_ENTREGADO"] = d["FECHA DE ENTREGA REAL"].notna()
    d["_A_TIEMPO"] = d["_ENTREGADO"] & (d["FECHA DE ENTREGA REAL"] <= d["PROMESA DE ENTREGA"])
    d["_INCIDENCIA"] = ~d["INCIDENCIAS"].astype(str).str.strip().str.upper().isin(["", "OK"])
    d["_DIAS"] = (d["FECHA DE ENTREGA REAL"] - d["FECHA DE ENVÍO"]).dt.days
    d["_COSTO"] = d["COSTO DE LA GUÍA"] + d["COSTOS ADICIONALES"]
    res = _resumen_por_fletera(d)

    entregados, a_tiempo = int(d["_ENTREGADO"].sum()), int(d["_A_TIEMPO"].sum())
    dias_ok = d[d["_ENTREGADO"] & (d["_DIAS"] >= 0)]["_DIAS"]
    cajas = d["CANTIDAD DE CAJAS"].sum()
    out = dict(vacio=False, df=d, res=res, registros=len(d), entregados=entregados,
               efectividad=(a_tiempo / entregados * 100) if entregados else 0.0,
               dias_prom=dias_ok.mean() if not dias_ok.empty else 0.0, dias_n=len(dias_ok),
               costo_envio=d["_COSTO"].mean() if len(d) else 0.0,
               costo_caja=(d["_COSTO"].sum() / cajas) if cajas else 0.0)

    def _ext(df_, col, fn):
        s = df_[df_[col].notna()]
        return s.loc[getattr(s[col], fn)()] if not s.empty else None
    out["top_puntual"] = _ext(res[res["ENTREGADOS"] > 0], "PCT_A_TIEMPO", "idxmax") if (res["ENTREGADOS"] > 0).any() else None
    out["top_incid"] = res.loc[res["INCIDENCIAS_PCT"].idxmax()] if not res.empty else None
    out["rapida"], out["lenta"] = _ext(res, "DIAS_PROM", "idxmin"), _ext(res, "DIAS_PROM", "idxmax")
    out["barata"], out["cara"] = _ext(res, "COSTO_PROM_ENVIO", "idxmin"), _ext(res, "COSTO_PROM_ENVIO", "idxmax")
    return out


# --- Muestras -------------------------------------------------
def preparar_muestras(df):
    d = df.copy()
    for col in ["PAQUETERIA_NOMBRE", "NUMERO_GUIA", "COSTO_GUIA", "CANTIDAD_TOTAL", "COSTO_TOTAL", "ESTATUS",
                "SOLICITO", "NOMBRE DEL HOTEL", "FOLIO", "FECHA"]:
        if col not in d.columns:
            d[col] = "NO SURTIDO" if col == "ESTATUS" else (0.0 if col in ("COSTO_GUIA", "CANTIDAD_TOTAL", "COSTO_TOTAL") else "")
    d["FECHA_DT"] = mp_parse_fecha_segura(d["FECHA"].astype(str).str.strip())
    d["COSTO_TOTAL"] = pd.to_numeric(d["COSTO_TOTAL"], errors="coerce").fillna(0)
    d["COSTO_GUIA"] = pd.to_numeric(d["COSTO_GUIA"], errors="coerce").fillna(0)
    d["INVERSION"] = d["COSTO_TOTAL"] + d["COSTO_GUIA"]
    return d


def calc_muestras(d_all, anio, mes, precios):
    d = d_all[(d_all["FECHA_DT"].dt.year == anio) & (d_all["FECHA_DT"].dt.month == mes)].copy()
    total = len(d)
    prod, flete = d["COSTO_TOTAL"].sum(), d["COSTO_GUIA"].sum()
    inv = prod + flete
    despachados = int((d["ESTATUS"].astype(str).str.upper() == "DESPACHADO").sum())
    sol = d["SOLICITO"].astype(str).str.upper()

    por_sol = pd.DataFrame()
    if total:
        por_sol = d.assign(SOL=sol).groupby("SOL").agg(
            envios=("FOLIO", "size"), prod=("COSTO_TOTAL", "sum"), flete=("COSTO_GUIA", "sum"),
            inversion=("INVERSION", "sum")).reset_index().sort_values("inversion", ascending=False).head(10)

    # Tendencia: 12 meses que terminan en el mes elegido
    fin = pd.Period(year=anio, month=mes, freq="M")
    meses_rango = [fin - i for i in range(11, -1, -1)]
    v = d_all.dropna(subset=["FECHA_DT"]).copy()
    v["P"] = v["FECHA_DT"].dt.to_period("M")
    serie = v.groupby("P")["INVERSION"].sum()
    tendencia = [(f"{p.month:02d}/{str(p.year)[2:]}", float(serie.get(p, 0.0))) for p in meses_rango]

    top_dest = d.assign(H=d["NOMBRE DEL HOTEL"].astype(str).str.upper()).groupby("H")["FOLIO"].size().sort_values(ascending=False).head(10)
    col_paq = "PAQUETERIA_NOMBRE" if "PAQUETERIA_NOMBRE" in d.columns else "PAQUETERIA"
    paq = d.assign(P=d[col_paq].replace("", "SIN ASIGNAR").replace(0.0, "SIN ASIGNAR").fillna("SIN ASIGNAR").astype(str).str.upper())
    flete_paq = paq.groupby("P")["COSTO_GUIA"].sum().sort_values(ascending=False).head(10)

    cant = {}
    for p in (precios or {}).keys():
        if p in d.columns:
            cant[p] = pd.to_numeric(d[p], errors="coerce").fillna(0).sum()
    prod_top = pd.Series(cant, dtype="float64").sort_values(ascending=False).head(10)
    prod_top = prod_top[prod_top > 0]

    return dict(total=total, prod=prod, flete=flete, inv=inv, prom=(inv / total) if total else 0.0,
                pct_desp=(despachados / total * 100) if total else 0.0,
                top_agente=(sol.value_counts().idxmax() if total else "SIN DATOS"),
                por_sol=por_sol, tendencia=tendencia, top_dest=top_dest, flete_paq=flete_paq, prod_top=prod_top)


# ============================================================
# 7. COMPONENTES VISUALES PARA EL PDF
# ============================================================
def mp_estilos():
    base = dict(fontName="Helvetica", textColor=mp_C_TEXT)
    return {
        "body": ParagraphStyle("body", fontSize=9, leading=12.5, **base),
        "bullet": ParagraphStyle("bullet", fontSize=9, leading=12.5, leftIndent=12, bulletIndent=0, spaceAfter=3, **base),
        "small": ParagraphStyle("small", fontSize=7.5, leading=10, textColor=mp_C_GRAY, fontName="Helvetica"),
        "sec_t": ParagraphStyle("sec_t", fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=colors.white),
        "sec_s": ParagraphStyle("sec_s", fontName="Helvetica", fontSize=8, leading=10, textColor=colors.HexColor("#B7C4CC")),
        "sub": ParagraphStyle("sub", fontName="Helvetica-Bold", fontSize=9, leading=12, textColor=mp_C_SLATE, spaceBefore=6, spaceAfter=4),
        "kpi_l": ParagraphStyle("kpi_l", fontName="Helvetica-Bold", fontSize=6.5, leading=8, textColor=mp_C_GRAY),
        "kpi_v": ParagraphStyle("kpi_v", fontName="Helvetica-Bold", fontSize=15, leading=18),
        "kpi_v_s": ParagraphStyle("kpi_v_s", fontName="Helvetica-Bold", fontSize=10.5, leading=13),
        "th": ParagraphStyle("th", fontName="Helvetica-Bold", fontSize=7, leading=8.5, textColor=colors.white),
        "nota": ParagraphStyle("nota", fontName="Helvetica-Oblique", fontSize=7.5, leading=10, textColor=mp_C_GRAY),
    }


mp_ST = mp_estilos()


def seccion(num_, titulo, subtitulo=""):
    cont = [Paragraph(f"{num_:02d}&nbsp;&nbsp;{esc(titulo)}", mp_ST["sec_t"])]
    if subtitulo:
        cont.append(Paragraph(esc(subtitulo), mp_ST["sec_s"]))
    t = Table([[cont]], colWidths=[CONTENT_W])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), mp_C_SLATE),
        ("LINEBEFORE", (0, 0), (0, 0), 5, mp_C_TEAL),
        ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 12),
        ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
    ]))
    return [t, Spacer(1, 10)]


def subtitulo(texto):
    return Paragraph(esc(texto).upper(), mp_ST["sub"])


def kpi_row(items, ncols=4, gap=8):
    """items: [(etiqueta, valor, color_hex)]"""
    cw = (CONTENT_W - gap * (ncols - 1)) / ncols
    out = []
    for i in range(0, len(items), ncols):
        chunk = items[i:i + ncols]
        row, widths, style = [], [], []
        for j, (lab, val, col) in enumerate(chunk):
            c = j * 2
            v_style = mp_ST["kpi_v"] if len(str(val)) <= 13 else mp_ST["kpi_v_s"]
            row.append([Paragraph(esc(lab).upper(), mp_ST["kpi_l"]),
                        Spacer(1, 3),
                        Paragraph(f'<font color="{col}">{esc(val)}</font>', v_style)])
            widths.append(cw)
            style += [("BACKGROUND", (c, 0), (c, 0), mp_C_LIGHT),
                      ("LINEABOVE", (c, 0), (c, 0), 3, colors.HexColor(col)),
                      ("BOX", (c, 0), (c, 0), 0.5, mp_C_BORDER)]
            if j < len(chunk) - 1:
                row.append("")
                widths.append(gap)
        t = Table([row], colWidths=widths, hAlign="LEFT")
        t.setStyle(TableStyle(style + [("VALIGN", (0, 0), (-1, -1), "TOP"),
                                       ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                                       ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        out += [t, Spacer(1, 8)]
    return out


def tabla(headers, rows, widths, aligns=None, font=7.5, total_row=None):
    aligns = aligns or ["L"] * len(headers)
    amap = {"L": 0, "C": 1, "R": 2}

    def cell(txt, a, bold=False, color=mp_C_TEXT):
        return Paragraph(esc(txt), ParagraphStyle("c", fontName="Helvetica-Bold" if bold else "Helvetica",
                                                  fontSize=font, leading=font + 2, alignment=amap[a], textColor=color))
    data = [[Paragraph(esc(h), ParagraphStyle("h", parent=mp_ST["th"], alignment=amap[a])) for h, a in zip(headers, aligns)]]
    for r in rows:
        data.append([cell(c, a) for c, a in zip(r, aligns)])
    if total_row:
        data.append([cell(c, a, bold=True) for c, a in zip(total_row, aligns)])
    t = Table(data, colWidths=widths, repeatRows=1)
    st_ = [("BACKGROUND", (0, 0), (-1, 0), mp_C_SLATE),
           ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
           ("LINEBELOW", (0, 0), (-1, -1), 0.3, mp_C_BORDER),
           ("BOX", (0, 0), (-1, -1), 0.5, mp_C_BORDER),
           ("TOPPADDING", (0, 0), (-1, -1), 3.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
           ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5)]
    for i in range(2, len(data), 2):
        st_.append(("BACKGROUND", (0, i), (-1, i), mp_C_LIGHT))
    if total_row:
        st_ += [("BACKGROUND", (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor("#E3EAEE")),
                ("LINEABOVE", (0, len(data) - 1), (-1, len(data) - 1), 0.8, mp_C_SLATE)]
    t.setStyle(TableStyle(st_))
    return t


def sin_datos(texto="Sin información para el periodo seleccionado."):
    t = Table([[Paragraph(esc(texto), mp_ST["nota"])]], colWidths=[CONTENT_W])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), mp_C_LIGHT), ("BOX", (0, 0), (-1, -1), 0.5, mp_C_BORDER),
                           ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
                           ("LEFTPADDING", (0, 0), (-1, -1), 10)]))
    return t


def lado_a_lado(izq, der, t_izq, t_der, w=(262, 262), gap=16):
    t = Table([[Paragraph(esc(t_izq).upper(), mp_ST["sub"]), "", Paragraph(esc(t_der).upper(), mp_ST["sub"])],
               [izq, "", der]], colWidths=[w[0], gap, w[1]], hAlign="LEFT")
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
    return t


# --- Gráficos (dibujados con reportlab) ----------------------
def mp_s(d, x, y, txt, size=7, bold=False, anchor="start", fill=mp_C_TEXT):
    d.add(String(x, y, txt, fontName="Helvetica-Bold" if bold else "Helvetica", fontSize=size,
                 textAnchor=anchor, fillColor=fill))


def chart_barras_h(items, width=CONTENT_W, color=mp_C_TEAL, fmt=num, label_w=150, bar_h=12, gap=5, colors_list=None):
    items = [(str(k), float(v)) for k, v in items]
    if not items:
        return sin_datos()
    n = len(items)
    h = n * (bar_h + gap) + gap
    d = Drawing(width, h)
    vmax = max(v for _, v in items) or 1.0
    area = width - label_w - 62
    y = h - gap - bar_h
    for i, (lab, v) in enumerate(items):
        mp_s(d, label_w - 6, y + bar_h / 2 - 2.5, trunc(lab, int(label_w / 4.1)), 7, anchor="end")
        w = max(area * v / vmax, 0.8) if v > 0 else 0
        d.add(Rect(label_w, y, w, bar_h, fillColor=(colors_list[i] if colors_list else color), strokeColor=None))
        mp_s(d, label_w + w + 4, y + bar_h / 2 - 2.5, fmt(v), 7, bold=True)
        y -= bar_h + gap
    return d


def chart_barras_v(cats, vals, width=CONTENT_W, height=150, color=mp_C_TEAL, fmt=num, rotar=False, colors_list=None):
    n = len(cats)
    if n == 0 or not any(vals):
        return sin_datos()
    d = Drawing(width, height)
    left, bottom, top = 8, 22, 14
    area_w, area_h = width - left - 6, height - bottom - top
    vmax = max(vals) or 1.0
    slot = area_w / n
    bw = min(slot * 0.62, 34)
    d.add(Line(left, bottom, left + area_w, bottom, strokeColor=mp_C_BORDER, strokeWidth=0.8))
    mostrar_val = n <= 16
    paso = 1 if n <= 16 else int(np.ceil(n / 16))
    for i, (c, v) in enumerate(zip(cats, vals)):
        x = left + i * slot + (slot - bw) / 2
        hh = area_h * v / vmax
        d.add(Rect(x, bottom, bw, hh, fillColor=(colors_list[i] if colors_list else color), strokeColor=None))
        if mostrar_val and v > 0:
            mp_s(d, x + bw / 2, bottom + hh + 3, fmt(v), 6.5, bold=True, anchor="middle")
        if i % paso == 0:
            mp_s(d, x + bw / 2, bottom - 10, str(c), 6.5, anchor="middle", fill=mp_C_GRAY)
    return d


def chart_apiladas_v(cats, series, width=CONTENT_W, height=165):
    """series: [(nombre, valores, color)]"""
    n = len(cats)
    if n == 0:
        return sin_datos()
    d = Drawing(width, height)
    left, bottom, top = 8, 22, 26
    area_w, area_h = width - left - 6, height - bottom - top
    totales = [sum(s[1][i] for s in series) for i in range(n)]
    vmax = max(totales) or 1.0
    slot = area_w / n
    bw = min(slot * 0.7, 26)
    d.add(Line(left, bottom, left + area_w, bottom, strokeColor=mp_C_BORDER, strokeWidth=0.8))
    paso = 1 if n <= 22 else int(np.ceil(n / 22))
    for i in range(n):
        x = left + i * slot + (slot - bw) / 2
        y0 = bottom
        for _, vals, col in series:
            hh = area_h * vals[i] / vmax
            if hh > 0:
                d.add(Rect(x, y0, bw, hh, fillColor=col, strokeColor=colors.white, strokeWidth=0.4))
            y0 += hh
        if totales[i] > 0 and n <= 22:
            mp_s(d, x + bw / 2, y0 + 2.5, str(int(totales[i])), 6.2, bold=True, anchor="middle")
        if i % paso == 0:
            mp_s(d, x + bw / 2, bottom - 9, trunc(cats[i], max(int(slot / 3.4), 4)), 6, anchor="middle", fill=mp_C_GRAY)
    lx = left
    for nombre, _, col in series:
        d.add(Rect(lx, height - 12, 8, 8, fillColor=col, strokeColor=None))
        mp_s(d, lx + 12, height - 11, nombre, 7)
        lx += 90
    return d


def chart_donut(items, width=262, height=130):
    """items: [(etiqueta, valor, color)]"""
    items = [(l, float(v), c) for l, v, c in items if v and v > 0]
    if not items:
        return sin_datos("Sin datos para graficar.")
    d = Drawing(width, height)
    size = min(height - 10, 110)
    pie = Pie()
    pie.x, pie.y, pie.width, pie.height = 4, (height - size) / 2, size, size
    pie.data = [v for _, v, _ in items]
    pie.labels = None
    pie.sideLabels = 0
    pie.startAngle = 90
    pie.direction = "clockwise"
    pie.slices.strokeColor = colors.white
    pie.slices.strokeWidth = 1.2
    for i, (_, _, c) in enumerate(items):
        pie.slices[i].fillColor = c
    d.add(pie)
    d.add(Circle(4 + size / 2, height / 2, size * 0.30, fillColor=colors.white, strokeColor=None))
    total = sum(v for _, v, _ in items)
    mp_s(d, 4 + size / 2, height / 2 - 4, num(total), 11, bold=True, anchor="middle")
    ly = height / 2 + len(items) * 8.5 - 6
    lx = size + 18
    for lab, v, c in items:
        d.add(Rect(lx, ly - 1, 8, 8, fillColor=c, strokeColor=None))
        mp_s(d, lx + 12, ly, f"{trunc(lab, 22)}", 7)
        mp_s(d, lx + 12, ly - 8.5, f"{num(v)}  ({v / total * 100:.0f}%)", 7, bold=True, fill=mp_C_GRAY)
        ly -= 19
    return d


def chart_scatter(puntos, width=CONTENT_W, height=200, xtitle="", ytitle=""):
    """puntos: [(etiqueta, x, y, tamaño)]  (y en porcentaje 0-100)"""
    if not puntos:
        return sin_datos("Sin datos suficientes para el mapa de valor.")
    d = Drawing(width, height)
    L, B, T, R = 48, 30, 22, 24
    aw, ah = width - L - R, height - B - T
    xs, ys = [p[1] for p in puntos], [p[2] for p in puntos]
    xmin, xmax = min(xs), max(xs)
    padx = (xmax - xmin) * 0.18 or max(abs(xmax) * 0.1, 1.0)
    xmin, xmax = max(xmin - padx, 0), xmax + padx
    # Eje Y ajustado a los datos (múltiplos de 10) para que los puntos no se amontonen
    ymin = max(np.floor((min(ys) - 8) / 10) * 10, 0)
    ymax = min(np.ceil((max(ys) + 8) / 10) * 10, 100)
    if ymax - ymin < 20:
        ymax = min(ymin + 20, 100)
        ymin = max(ymax - 20, 0)
    smax = max(p[3] for p in puntos) or 1
    pasos = 4
    for k in range(pasos + 1):
        yv = ymin + (ymax - ymin) * k / pasos
        yy = B + ah * k / pasos
        d.add(Line(L, yy, L + aw, yy, strokeColor=mp_C_BORDER, strokeWidth=0.4))
        mp_s(d, L - 5, yy - 2.5, f"{yv:.0f}%", 6.5, anchor="end", fill=mp_C_GRAY)
    for k in range(5):
        xv = xmin + (xmax - xmin) * k / 4
        xx = L + aw * k / 4
        d.add(Line(xx, B, xx, B - 3, strokeColor=mp_C_GRAY, strokeWidth=0.5))
        mp_s(d, xx, B - 11, money(xv), 6.5, anchor="middle", fill=mp_C_GRAY)
    d.add(Line(L, B, L + aw, B, strokeColor=mp_C_GRAY, strokeWidth=0.8))
    d.add(Line(L, B, L, B + ah, strokeColor=mp_C_GRAY, strokeWidth=0.8))
    mp_s(d, L + aw / 2, 4, xtitle, 7, bold=True, anchor="middle", fill=mp_C_GRAY)
    mp_s(d, L, height - 10, ytitle, 6.5, bold=True, fill=mp_C_GRAY)

    ocupados = []   # rectángulos (x0, y0, x1, y1) de etiquetas ya colocadas
    for lab, x, y, sz in sorted(puntos, key=lambda p: -p[3]):
        px_ = L + aw * (x - xmin) / (xmax - xmin)
        py_ = B + ah * (min(max(y, ymin), ymax) - ymin) / (ymax - ymin)
        r = 4 + 7 * (sz / smax) ** 0.5
        d.add(Circle(px_, py_, r, fillColor=colors.Color(0, 0.64, 0.64, alpha=0.75), strokeColor=colors.white, strokeWidth=0.8))
        txt = trunc(lab, 18)
        w_txt = 4.1 * len(txt)
        for dy in (0, 10, -10, 20, -20, 30, -30):
            x0, y0 = px_ + r + 3, py_ - 3 + dy
            if x0 + w_txt > width:           # si no cabe a la derecha, va a la izquierda
                x0 = px_ - r - 3 - w_txt
            caja = (x0, y0 - 1, x0 + w_txt, y0 + 8)
            if not any(caja[0] < o[2] and caja[2] > o[0] and caja[1] < o[3] and caja[3] > o[1] for o in ocupados):
                break
        ocupados.append(caja)
        mp_s(d, caja[0], caja[1] + 1, txt, 7, bold=True)
    return d


def chart_lineas(cats, series, width=CONTENT_W, height=175, target=None, fmt=lambda v: f"{v:.1f}%"):
    """series: [(nombre, valores (None/NaN = sin dato), color)].  target: línea punteada horizontal (mismo eje)."""
    n = len(cats)
    vals_all = [v for _, vs, _ in series for v in vs if v is not None and not pd.isna(v)]
    if n == 0 or not vals_all:
        return sin_datos("Sin datos para graficar.")
    d = Drawing(width, height)
    L, B, T, R = 40, 24, 28, 14
    aw, ah = width - L - R, height - B - T
    vmax = max(vals_all + ([target] if target else [])) * 1.15 or 1.0
    paso_x = aw / max(n - 1, 1)
    for k in range(5):
        yy = B + ah * k / 4
        d.add(Line(L, yy, L + aw, yy, strokeColor=mp_C_BORDER, strokeWidth=0.4))
        mp_s(d, L - 5, yy - 2.5, f"{vmax * k / 4:.1f}%", 6.5, anchor="end", fill=mp_C_GRAY)
    d.add(Line(L, B, L + aw, B, strokeColor=mp_C_GRAY, strokeWidth=0.8))
    if target:
        ty = B + ah * target / vmax
        d.add(Line(L, ty, L + aw, ty, strokeColor=mp_C_GREEN, strokeWidth=1, strokeDashArray=[4, 3]))
        mp_s(d, L + 3, ty + 3, f"target {target}%", 6.5, bold=True, anchor="start", fill=mp_C_GREEN)
    for i, c in enumerate(cats):
        mp_s(d, L + i * paso_x, B - 10, str(c), 6.5, anchor="middle", fill=mp_C_GRAY)
    for si, (nombre, vs, col) in enumerate(series):
        pts = [(L + i * paso_x, B + ah * v / vmax, v) if (v is not None and not pd.isna(v)) else None for i, v in enumerate(vs)]
        for a, b in zip(pts, pts[1:]):
            if a and b:
                d.add(Line(a[0], a[1], b[0], b[1], strokeColor=col, strokeWidth=1.6))
        for p in pts:
            if p:
                d.add(Circle(p[0], p[1], 2.6, fillColor=col, strokeColor=colors.white, strokeWidth=0.6))
                if n <= 14:
                    mp_s(d, p[0], p[1] + (5 if si % 2 == 0 else -9), fmt(p[2]), 6.2, bold=True, anchor="middle", fill=col)
    lx = L
    for nombre, _, col in series:
        d.add(Rect(lx, height - 12, 8, 8, fillColor=col, strokeColor=None))
        mp_s(d, lx + 12, height - 11, nombre, 7)
        lx += 110
    return d


# ============================================================
# 8. DOCUMENTO: PLANTILLAS, ENCABEZADO Y PIE
# ============================================================
class _NumberedCanvas(rl_canvas.Canvas):
    titulo_pie = ""

    def __init__(self, *a, **k):
        super().__init__(*a, **k)
        self._saved = []

    def showPage(self):
        self._saved.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        n = len(self._saved)
        for s_ in self._saved:
            self.__dict__.update(s_)
            self.setStrokeColor(mp_C_BORDER)
            self.setLineWidth(0.5)
            self.line(MARGIN, 34, PAGE_W - MARGIN, 34)
            self.setFont("Helvetica", 7)
            self.setFillColor(mp_C_GRAY)
            self.drawString(MARGIN, 23, f"JYPESA | Logística  -  {self.titulo_pie}  -  Generado e impreso con Nexion Smart Logistic")
            self.drawRightString(PAGE_W - MARGIN, 23, f"Página {self._pageNumber} de {n}")
            super().showPage()
        super().save()


def _hacer_paginas(titulo_mes, subtitulo_portada, logo_bytes):
    BANNER_H = 175

    def portada(c, doc):
        c.saveState()
        c.setFillColor(mp_C_NAVY)
        c.rect(0, PAGE_H - BANNER_H, PAGE_W, BANNER_H, stroke=0, fill=1)
        c.setFillColor(mp_C_TEAL)
        c.rect(0, PAGE_H - BANNER_H, PAGE_W, 5, stroke=0, fill=1)
        c.setFillColor(mp_C_GOLD)
        c.rect(MARGIN, PAGE_H - 70, 38, 3, stroke=0, fill=1)
        c.setFillColor(colors.HexColor("#9FB3BF"))
        c.setFont("Helvetica-Bold", 9)
        c.drawString(MARGIN, PAGE_H - 58, "JYPESA  |  LOGÍSTICA Y DISTRIBUCIÓN")
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 30)
        c.drawString(MARGIN, PAGE_H - 108, "REPORTE MENSUAL")
        c.setFont("Helvetica-Bold", 30)
        c.setFillColor(mp_C_GOLD)
        c.drawString(MARGIN, PAGE_H - 142, titulo_mes)
        c.setFillColor(colors.HexColor("#9FB3BF"))
        c.setFont("Helvetica", 8.5)
        c.drawString(MARGIN, PAGE_H - 162, subtitulo_portada)
        if logo_bytes:
            try:
                
                c.drawImage(ImageReader(BytesIO(logo_bytes)), PAGE_W - MARGIN - 110, PAGE_H - 88, width=102, height=48,
                            preserveAspectRatio=True, mask='auto', anchor='c')
            except Exception:
                pass
        c.restoreState()

    def normal(c, doc):
        c.saveState()
        c.setFillColor(mp_C_NAVY)
        c.rect(0, PAGE_H - 30, PAGE_W, 30, stroke=0, fill=1)
        c.setFillColor(mp_C_TEAL)
        c.rect(0, PAGE_H - 30, PAGE_W, 2.5, stroke=0, fill=1)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 8.5)
        c.drawString(MARGIN, PAGE_H - 19, "REPORTE MENSUAL DE LOGÍSTICA")
        c.setFillColor(mp_C_GOLD)
        c.drawRightString(PAGE_W - MARGIN, PAGE_H - 19, titulo_mes)
        c.restoreState()

    f_portada = Frame(MARGIN, 44, CONTENT_W, PAGE_H - BANNER_H - 44 - 16, id="fp", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    f_normal = Frame(MARGIN, 44, CONTENT_W, PAGE_H - 44 - 48, id="fn", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    return [PageTemplate(id="portada", frames=[f_portada], onPage=portada),
            PageTemplate(id="normal", frames=[f_normal], onPage=normal)]


# ============================================================
# 9. ARMADO DEL REPORTE
# ============================================================
def generar_reporte_pdf(df_base, df_muestras, anio, mes, hoy, precios=None, logo_bytes=None, muestras_msg="", df_hist=None):
    """Devuelve BytesIO con el PDF.  df_base: salida de preparar_base().  df_muestras: salida de preparar_muestras() o None."""
    titulo_mes = f"{MESES[mes - 1]} {anio}"
    ahora = datetime.now().strftime("%d/%m/%Y %H:%M")

    # Recolecciones / maniobras: no son pedidos, se reportan como costos extras
    df_base, df_extras = _separar_extras(df_base)
    E_all = calc_extras(df_extras, anio, mes, "TODAS")
    E_reg = calc_extras(df_extras, anio, mes, "COBRO REGRESO")

    df_mes = filtrar_mes(df_base, "FECHA DE ENVÍO", anio, mes)
    R = calc_resumen(df_mes, hoy)
    D = calc_despachos(df_mes)
    B = calc_bi(df_mes)
    T = calc_top_clientes(df_mes, 20)
    C = calc_carga_regreso(df_base, anio, mes)
    G = calc_costos_regreso(df_base, anio, mes)
    if not G.get("vacio"):
        G["extras"] = E_reg
        G["flete_total"] = G["flete"] + E_reg["costo"]          # pedidos + extras
        G["cajas_total"] = G["cajas"] + E_reg["cajas"]
        G["costo_log_total"] = (G["flete_total"] / G["fact"] * 100) if G["fact"] else 0.0
    H = calc_comparativa_2025(G, df_hist, anio, mes)
    K = calc_ranking(df_base, anio, mes)
    M = calc_muestras(df_muestras, anio, mes, precios) if df_muestras is not None else None

    story = [NextPageTemplate("normal")]

    # ---------------- PORTADA / RESUMEN EJECUTIVO ----------------
    story += [Spacer(1, 6), Paragraph("RESUMEN EJECUTIVO", ParagraphStyle("re", fontName="Helvetica-Bold", fontSize=11, textColor=mp_C_SLATE, spaceAfter=8))]
    efect = None if K.get("vacio") else K["efectividad"]
    story += kpi_row([
        ("Pedidos del mes", num(R["total"]), mp_HEX["slate"]),
        ("Entregados", f"{num(R['entregados'])} ({R['pct_entregado']:.0f}%)", mp_HEX["green"]),
        ("Facturación", money(B["fact"]), mp_HEX["teal"]),
        ("Costo logístico + extras", pct((B["guias"] + E_all["costo"]) / B["fact"] * 100 if B["fact"] else 0.0), mp_HEX["purple"]),
        ("Despacho en 24 h háb.", pct(D["pct_ok"]) if D["total"] else "-", mp_HEX["blue"]),
        ("Entregas a tiempo", pct(efect, 0) if efect is not None else "-", mp_HEX["green"]),
        ("Tránsito promedio", f"{K['dias_prom']:.1f} días" if not K.get("vacio") else "-", mp_HEX["orange"]),
        ("Inversión en muestras", money(M["inv"]) if M else "N/D", mp_HEX["gold"]),
    ], ncols=4)

    hallazgos = []
    if R["total"]:
        hallazgos.append(f"Se enviaron <b>{num(R['total'])}</b> pedidos en {titulo_mes.title()}: <b>{num(R['entregados'])}</b> ya fueron entregados "
                         f"({R['pct_entregado']:.0f}%) y <b>{num(R['en_transito'])}</b> siguen en tránsito, de los cuales <b>{num(R['retrasados'])}</b> "
                         f"presentan retraso frente a su promesa de entrega.")
    else:
        hallazgos.append(f"No se encontraron pedidos con fecha de envío en {titulo_mes.title()}.")
    if D["total"]:
        hallazgos.append(f"<b>{D['pct_ok']:.1f}%</b> de las facturas salió de almacén dentro de las 24 horas hábiles ({num(D['ok'])} de {num(D['total'])}); "
                         f"<b>{num(D['no'])}</b> quedaron fuera de meta.")
    if not K.get("vacio") and K["entregados"]:
        txt = f"La efectividad global de entregas fue <b>{K['efectividad']:.0f}%</b> con un tránsito promedio de <b>{K['dias_prom']:.1f} días</b>."
        if K["top_puntual"] is not None:
            txt += f" La más puntual fue <b>{esc(K['top_puntual']['FLETERA'])}</b> ({K['top_puntual']['PCT_A_TIEMPO']:.0f}%)"
        if K["top_incid"] is not None and K["top_incid"]["INCIDENCIAS_PCT"] > 0:
            txt += f" y la que más incidencias registró, <b>{esc(K['top_incid']['FLETERA'])}</b> ({K['top_incid']['INCIDENCIAS_PCT']:.0f}% de sus envíos)."
        else:
            txt += "."
        hallazgos.append(txt)
    if B["fact"]:
        txt_b = (f"El costo de guías equivale a <b>{B['costo_log']:.1f}%</b> de la facturación ({money(B['guias'])} sobre {money(B['fact'])}); "
                 f"costo promedio por caja <b>{money(B['costo_caja'], 2)}</b>.")
        if E_all["n"]:
            txt_b += (f" Además se pagaron <b>{money(E_all['costo'])}</b> en costos extras (recolecciones y maniobras, que no se cuentan como pedidos); "
                      f"con ellos el costo logístico sube a <b>{(B['guias'] + E_all['costo']) / B['fact'] * 100:.1f}%</b>.")
        hallazgos.append(txt_b)
    if not K.get("vacio") and K["barata"] is not None and K["cara"] is not None and K["barata"]["FLETERA"] != K["cara"]["FLETERA"]:
        hallazgos.append(f"Costo promedio por envío: la más económica fue <b>{esc(K['barata']['FLETERA'])}</b> ({money(K['barata']['COSTO_PROM_ENVIO'])}) "
                         f"y la más cara <b>{esc(K['cara']['FLETERA'])}</b> ({money(K['cara']['COSTO_PROM_ENVIO'])}).")
    if not C.get("vacio"):
        hallazgos.append(f"En cobro regreso se movieron <b>{num(C['total'])}</b> cajas; carrier dominante: <b>{esc(C['lider']['TRANSPORTE'])}</b> ({C['lider']['PCT']:.0f}%).")
    if not G.get("vacio"):
        mayor = G["res"].iloc[0]
        dentro = G["costo_log_total"] <= TARGET_COSTO_LOG
        txt_g = (f"Fletes de cobro regreso: se pagaron <b>{money(G['flete'])}</b> en pedidos ({money(G['costo_caja'], 2)} por caja)"
                 + (f" y <b>{money(E_reg['costo'])}</b> en costos extras (recolecciones y maniobras)" if E_reg["n"] else "")
                 + f", equivalente a <b>{G['costo_log_total']:.2f}%</b> de la facturación de esa modalidad ({'dentro' if dentro else 'fuera'} del target de {TARGET_COSTO_LOG}%)"
                 + (f"; solo pedidos: {G['costo_log']:.2f}%" if E_reg["n"] else "")
                 + f". El mayor gasto en pedidos fue con <b>{esc(mayor['FLETERA'])}</b> ({mayor['pct_gasto']:.0f}% del total).")
        if H and not H["sin_dato"] and H["var_flete"] is not None and H["var_cajas"] is not None:
            txt_g += (f" Contra {MESES[mes - 1].title()} {ANIO_HISTORIAL}, el gasto de flete varió <b>{H['var_flete']:+.1f}%</b> "
                      f"y el volumen de cajas <b>{H['var_cajas']:+.1f}%</b>.")
        hallazgos.append(txt_g)
    if M and M["total"]:
        hallazgos.append(f"Muestras: <b>{num(M['total'])}</b> envíos con una inversión de <b>{money(M['inv'])}</b> "
                         f"(productos {money(M['prod'])} + fletes {money(M['flete'])}); {M['pct_desp']:.0f}% despachado.")
    story.append(Paragraph("HALLAZGOS CLAVE", ParagraphStyle("hk", fontName="Helvetica-Bold", fontSize=11, textColor=mp_C_SLATE, spaceBefore=4, spaceAfter=6)))
    for h in hallazgos:
        story.append(Paragraph(h, mp_ST["bullet"], bulletText="•"))

    contenido = ["01  Resumen del mes", "02  Efectividad de envíos", "03  Inteligencia de negocio", "04  Top 20 clientes",
                 "05  Distribución de carga (cobro regreso)", "06  Cobro regreso: costo por fletera", "07  Ranking de fleteras",
                 "08  Costos de muestras"]
    story.append(Spacer(1, 8))
    story.append(KeepTogether([Paragraph("CONTENIDO", ParagraphStyle("ct", fontName="Helvetica-Bold", fontSize=9, textColor=mp_C_GRAY, spaceAfter=3))]
                              + [Paragraph(c, mp_ST["small"]) for c in contenido]))

    # ---------------- 01 RESUMEN DEL MES ----------------
    story += [PageBreak()] + seccion(1, "RESUMEN DEL MES", f"Pedidos, entregas y retrasos con fecha de envío en {titulo_mes.title()}")
    story += kpi_row([
        ("Pedidos", num(R["total"]), mp_HEX["slate"]),
        ("Entregados", num(R["entregados"]), mp_HEX["green"]),
        ("En tránsito", num(R["en_transito"]), mp_HEX["blue"]),
        ("En tiempo", num(R["en_tiempo"]), mp_HEX["gold"]),
        ("Con retraso", num(R["retrasados"]), mp_HEX["red"]),
    ], ncols=5)
    if R["total"]:
        don = chart_donut([("Entregados", R["entregados"], mp_C_GREEN), ("Tránsito en tiempo", R["en_tiempo"], mp_C_GOLD),
                           ("Tránsito con retraso", R["retrasados"], mp_C_RED)])
        rf = R["ret_fletera"]
        der = chart_barras_h(list(rf.items()), width=262, label_w=100, color=mp_C_RED, colors_list=None) if len(rf) else sin_datos("Sin pedidos con retraso en este periodo.")
        story.append(lado_a_lado(don, der, "Distribución de pedidos del mes", "Pedidos con retraso por fletera"))
        story.append(Spacer(1, 10))
        story.append(KeepTogether([subtitulo("Pedidos por destino (top 8)"),
                                   chart_barras_h(list(R["destinos"].items()), color=mp_C_GOLD, label_w=170)]))
        story += [Spacer(1, 6), Paragraph("Nota: 'En tiempo' y 'Con retraso' se evalúan sobre pedidos aún en tránsito, comparando la promesa de entrega contra la fecha de generación "
                                          f"de este reporte ({hoy.strftime('%d/%m/%Y')}).", mp_ST["nota"])]
    else:
        story.append(sin_datos())

    # ---------------- 02 EFECTIVIDAD DE ENVÍOS ----------------
    story += [PageBreak()] + seccion(2, "EFECTIVIDAD DE ENVÍOS", "Despachos de almacén dentro de 24 horas hábiles (fines de semana y feriados no cuentan)")
    if D["total"]:
        story += kpi_row([("Total facturas evaluadas", num(D["total"]), mp_HEX["slate"]),
                          ("A tiempo", f"{num(D['ok'])}  ·  {D['pct_ok']:.1f}%", mp_HEX["green"]),
                          ("Fuera de meta", f"{num(D['no'])}  ·  {D['pct_no']:.1f}%", mp_HEX["red"])], ncols=3)
        pd_ = D["por_dia"]
        cats = [i.strftime("%d/%m") for i in pd_.index]
        story.append(KeepTogether([subtitulo("Despachos por día de salida"),
                                   chart_apiladas_v(cats, [("A tiempo", pd_["A Tiempo"].astype(int).tolist(), mp_C_GREEN),
                                                           ("Fuera de tiempo", pd_["Fuera de Tiempo"].astype(int).tolist(), mp_C_RED)])]))
        story.append(Spacer(1, 8))
        if not D["fuera"].empty:
            rows = [[str(r["NÚMERO DE PEDIDO"]),
                     r["EMISION"].strftime("%d/%m/%Y %H:%M") if pd.notna(r["EMISION"]) else "S/D",
                     r["FECHA DE ENVÍO"].strftime("%d/%m/%Y %H:%M") if pd.notna(r["FECHA DE ENVÍO"]) else "S/D",
                     f"{int(r['DIAS_HABILES'])}"] for _, r in D["fuera"].iterrows()]
            story.append(KeepTogether([subtitulo("Pedidos con mayor retraso en despacho (top 12)"),
                                       tabla(["PEDIDO", "EMISIÓN", "SALIDA DE ALMACÉN", "DÍAS HÁBILES"], rows, [150, 130, 150, 110], ["L", "C", "C", "C"])]))
        else:
            story.append(Paragraph("Todos los despachos del periodo salieron dentro de la meta.", mp_ST["body"]))
        if D["sin_datos"]:
            story += [Spacer(1, 4), Paragraph(f"{num(D['sin_datos'])} registro(s) sin fecha de emisión/salida no se incluyeron en el cálculo.", mp_ST["nota"])]
    else:
        story.append(sin_datos())

    # ---------------- 03 INTELIGENCIA DE NEGOCIO ----------------
    story += [PageBreak()] + seccion(3, "INTELIGENCIA DE NEGOCIO", "Facturación, costos logísticos y comparativo por fletera")
    story += kpi_row([("Facturación", money(B["fact"]), mp_HEX["green"]), ("Costo guías", money(B["guias"]), mp_HEX["blue"]),
                      ("Costos adicionales", money(B["adic"]), mp_HEX["gold"])], ncols=3)
    story += kpi_row([("Costo logístico (guías / fact.)", pct(B["costo_log"]), mp_HEX["purple"]),
                      ("Cajas enviadas", num(B["cajas"]), mp_HEX["slate"]),
                      ("Costo promedio por caja", money(B["costo_caja"], 2), mp_HEX["red"])], ncols=3)
    if E_all["n"]:
        con_ext = (B["guias"] + E_all["costo"]) / B["fact"] * 100 if B["fact"] else 0.0
        story += kpi_row([("Costos extras (recolec. y maniobras)", money(E_all["costo"]), mp_HEX["gold"]),
                          ("Costo logístico + extras", pct(con_ext), mp_HEX["purple"]),
                          ("Registros de costos extras", num(E_all["n"]), mp_HEX["slate"])], ncols=3)
        story.append(Paragraph("Los registros con concepto de recolecciones o maniobras no se cuentan como pedidos (ni en envíos, cajas, facturación ni efectividad); "
                               "su costo (guía + adicionales) se muestra como costo extra. Costo logístico + extras = (costo de guías + costos extras) / facturación.",
                               mp_ST["nota"]))
        story.append(Spacer(1, 4))
    pf = B["por_fletera"]
    if not pf.empty:
        top_f = pf.head(8)
        izq = chart_barras_h([(r.FLETERA, r.facturacion) for r in top_f.itertuples()], width=262, label_w=92, color=mp_C_GREEN, fmt=money)
        top_o = pf.sort_values("operativo", ascending=False).head(8)
        der = chart_barras_h([(r.FLETERA, r.operativo) for r in top_o.itertuples()], width=262, label_w=92, color=C_ORANGE, fmt=money)
        story.append(lado_a_lado(izq, der, "Facturación generada por fletera", "Costo operativo por fletera (guía + adic.)"))
        story.append(Spacer(1, 10))
        rows = [[r.FLETERA, num(r.pedidos), money(r.facturacion), money(r.guias), money(r.adic), pct(r.pct_log),
                 num(r.cajas), money(r.costo_caja, 2) if pd.notna(r.costo_caja) else "-"] for r in pf.head(12).itertuples()]
        story.append(KeepTogether([subtitulo("Detalle por fletera"),
                                   tabla(["FLETERA", "PEDIDOS", "FACTURACIÓN", "COSTO GUÍAS", "COSTOS ADIC.", "% LOGÍSTICO", "CAJAS", "COSTO/CAJA"], rows,
                                         [95, 50, 80, 70, 65, 62, 50, 68], ["L", "R", "R", "R", "R", "R", "R", "R"])]))
        story += [Spacer(1, 4), Paragraph("Costo/caja = (guía + adicionales) / cajas. Costo logístico = costo de guías / facturación.", mp_ST["nota"])]
    else:
        story.append(sin_datos())

    # ---------------- 04 TOP 20 CLIENTES ----------------
    story += [PageBreak()] + seccion(4, "TOP 20 CLIENTES DE DISTRIBUCIÓN", "Clientes con mayor facturación del periodo y forma de envío")
    top = T["top"]
    if not top.empty:
        story += kpi_row([("Facturación total del periodo", money(T["total_fact"]), mp_HEX["green"]),
                          ("Concentración del Top 20", pct(T["share_top"]), mp_HEX["purple"]),
                          ("Clientes distintos", num(df_mes[df_mes['NOMBRE DEL CLIENTE'] != '']['NOMBRE DEL CLIENTE'].nunique()), mp_HEX["slate"])], ncols=3)
        rows = [[str(i + 1), trunc(r["NOMBRE DEL CLIENTE"], 46), num(r["pedidos"]), num(r["cajas"]),
                 money(r["facturacion"]), money(r["guias"]), pct(r["share"])] for i, (_, r) in enumerate(top.iterrows())]
        story.append(tabla(["#", "CLIENTE", "PEDIDOS", "CAJAS", "FACTURACIÓN", "COSTO GUÍA", "% DEL TOTAL"], rows,
                           [22, 208, 50, 45, 78, 70, 67], ["C", "L", "R", "R", "R", "R", "R"],
                           total_row=["", "TOP 20", num(top["pedidos"].sum()), num(top["cajas"].sum()), money(top["facturacion"].sum()),
                                      money(top["guias"].sum()), pct(T["share_top"])]))
        if not T["forma"].empty:
            colores = [mp_C_TEAL, C_BLUE, mp_C_GOLD, C_PURPLE, mp_C_RED, C_ORANGE, mp_C_GREEN, mp_C_GRAY]
            story += [Spacer(1, 10), KeepTogether([subtitulo("Distribución por forma de envío (pedidos)"),
                                                   chart_donut([(k, v, colores[i % len(colores)]) for i, (k, v) in enumerate(T["forma"].head(7).items())], width=CONTENT_W, height=120)])]
    else:
        story.append(sin_datos())

    # ---------------- 05 DISTRIBUCIÓN DE CARGA (COBRO REGRESO) ----------------
    story += [PageBreak()] + seccion(5, "DISTRIBUCIÓN DE CARGA  ·  COBRO REGRESO", "Volumen de cajas por carrier y destino (solo flujo de cobro regreso)")
    if not C.get("vacio"):
        story += kpi_row([("Volumen total (cajas)", num(C["total"]), mp_HEX["slate"]),
                          ("Carrier dominante", f"{C['lider']['TRANSPORTE']} · {C['lider']['PCT']:.0f}%", mp_HEX["green"]),
                          ("Destinos distintos", num(C["destinos"]), mp_HEX["teal"])], ncols=3)
        story.append(KeepTogether([subtitulo("Cajas por carrier"),
                                   chart_barras_h([(r.TRANSPORTE, r.CAJAS) for r in C["part"].itertuples()], label_w=170, color=mp_C_TEAL,
                                                  fmt=lambda v: f"{num(v)}")]))
        story.append(Spacer(1, 8))
        rutas = C["rutas"]
        MAXF = 20
        rows = [[r["TRANSPORTE"], r["DESTINO"], r["FORMA DE ENVIO"], num(r["CAJAS"])] for _, r in rutas.head(MAXF).iterrows()]
        story.append(subtitulo("Detalle de rutas (carrier / destino)"))
        story.append(tabla(["CARRIER", "DESTINO", "FORMA DE ENVÍO", "CAJAS"], rows, [190, 140, 130, 80], ["L", "L", "L", "R"],
                           total_row=["TOTAL", "", "", num(C["total"])]))
        if len(rutas) > MAXF:
            story += [Spacer(1, 3), Paragraph(f"Se muestran las {MAXF} rutas principales de {len(rutas)}. El total incluye todas.", mp_ST["nota"])]
    else:
        story.append(sin_datos(f"No se encontraron registros de COBRO REGRESO en {titulo_mes.title()}."))

    # ---------------- 06 COBRO REGRESO: COSTO POR FLETERA ----------------
    story += [PageBreak()] + seccion(6, "COBRO REGRESO  ·  COSTO POR FLETERA",
                                     "Lo que paga JYPESA por fletes de regreso: gasto, costo por caja y servicio por fletera")
    if not G.get("vacio"):
        col_log_t = mp_HEX["green"] if G["costo_log_total"] <= TARGET_COSTO_LOG else mp_HEX["red"]
        story += kpi_row([
            ("Costo de flete pedidos (guía + adic.)", money(G["flete"], 2), mp_HEX["slate"]),
            ("Cajas enviadas", num(G["cajas"]), mp_HEX["teal"]),
            ("Costo por caja (pedidos)", money(G["costo_caja"], 2), mp_HEX["orange"]),
            ("Costos extras (recolec. y maniobras)", money(E_reg["costo"], 2), mp_HEX["gold"]),
        ], ncols=4)
        story += kpi_row([
            ("Facturación cobro regreso", money(G["fact"], 2), mp_HEX["green"]),
            ("Eficiencia de entrega", pct(G["efic"]) if G["efic"] is not None else "-", mp_HEX["blue"]),
            ("Valuación de incidencias", money(G["val"], 2), mp_HEX["red"]),
            ("% de incidencias", pct(G["pct_inc"]), mp_HEX["gold"]),
        ], ncols=4)
        if H is not None:
            nm = MESES[mes - 1].title()
            story.append(subtitulo(f"Comparativo contra {nm} {ANIO_HISTORIAL}  //  flete y volumen"))
            if H["sin_dato"]:
                story += [sin_datos(H["motivo"]), Spacer(1, 8)]
            else:
                def _var(v, dif, dec, bueno_si_sube, money_=False):
                    if v is None:
                        return "-", mp_HEX["gray"]
                    d_txt = ("+" if dif >= 0 else "-") + (money(abs(dif)) if money_ else num(abs(dif)))
                    sube_es_bien = (v >= 0) == bueno_si_sube
                    return f"{v:+.1f}%  ({d_txt})", (mp_HEX["green"] if sube_es_bien else mp_HEX["red"])
                v_f, c_f = _var(H["var_flete"], H["dif_flete"], 1, False, money_=True)
                v_c, c_c = _var(H["var_cajas"], H["dif_cajas"], 1, True)
                story += kpi_row([(f"Flete {nm} {ANIO_HISTORIAL}", money(H["flete25"], 2), mp_HEX["slate"]),
                                  (f"Flete {anio} vs {ANIO_HISTORIAL}", v_f, c_f),
                                  (f"Cajas {nm} {ANIO_HISTORIAL}", num(H["cajas25"]), mp_HEX["slate"]),
                                  (f"Cajas {anio} vs {ANIO_HISTORIAL}", v_c, c_c)], ncols=4)
        story.append(subtitulo("Costo logístico y desglose por concepto"))
        story += kpi_row([(f"Costo logístico con extras (target {TARGET_COSTO_LOG}%)", f"{G['costo_log_total']:.2f}%", col_log_t),
                          ("Costo logístico solo pedidos", f"{G['costo_log']:.2f}%", mp_HEX["slate"]),
                          ("Consignas (informativo)", money(G["consignas"], 2), mp_HEX["purple"]),
                          ("F nacional (informativo)", money(G["fnacional"], 2), mp_HEX["purple"])], ncols=4)

        res = G["res"]
        izq = chart_barras_h([(r.FLETERA, r.flete) for r in res.itertuples()], width=262, label_w=92, color=C_ORANGE, fmt=money)
        cc = res[res["costo_caja"].notna()].sort_values("costo_caja", ascending=True)
        der = (chart_barras_h([(r.FLETERA, r.costo_caja) for r in cc.itertuples()], width=262, label_w=92, color=mp_C_GOLD,
                              fmt=lambda v: money(v, 2)) if not cc.empty else sin_datos("Sin cajas registradas."))
        story.append(lado_a_lado(izq, der, "Gasto de flete por fletera", "Costo promedio por caja"))
        story.append(Spacer(1, 10))

        rows = [[trunc(r.FLETERA, 24), num(r.envios), num(r.cajas), money(r.flete, 2), pct(r.pct_gasto, 0),
                 money(r.costo_caja, 2) if pd.notna(r.costo_caja) else "-", pct(r.pct_log, 2) if pd.notna(r.pct_log) else "-",
                 pct(r.efic, 0) if pd.notna(r.efic) else "-", money(r.val, 2)] for r in res.itertuples()]
        story.append(KeepTogether([subtitulo("Detalle por fletera (cobro regreso)"),
                                   tabla(["FLETERA", "ENVÍOS", "CAJAS", "COSTO FLETE", "% GASTO", "COSTO/CAJA", "% LOGÍSTICO", "% A TIEMPO", "VALUACIÓN INC."],
                                         rows, [100, 40, 42, 70, 48, 58, 52, 56, 74],
                                         ["L", "R", "R", "R", "R", "R", "R", "R", "R"],
                                         total_row=["TOTAL", num(G["registros"]), num(G["cajas"]), money(G["flete"], 2), "100%",
                                                    money(G["costo_caja"], 2), pct(G["costo_log"], 2),
                                                    pct(G["efic"], 0) if G["efic"] is not None else "-", money(G["val"], 2)])]))
        story += [Spacer(1, 4), Paragraph(
            f"Solo incluye envíos con forma de envío COBRO REGRESO. Los registros con concepto de recolecciones o maniobras no se cuentan como pedidos: "
            "su costo (guía + adicionales) se muestra como costo extra y no entra en el detalle por fletera ni en el costo por caja. "
            f"Costo logístico solo pedidos = flete / facturación; con extras = (flete + extras) / facturación (target {TARGET_COSTO_LOG}%). "
            "A diferencia de la sección 03, aquí se incluyen los adicionales y no se mezclan otras modalidades. "
            "% a tiempo = entregas con fecha real menor o igual a la promesa, sobre envíos con ambas fechas. "
            "Consignas y F nacional usan solo el costo de guía y son informativos.", mp_ST["nota"])]
        if H is not None and not H["sin_dato"]:
            story += [Spacer(1, 3), Paragraph(
                f"Comparativo anual: flete {anio} (pedidos + costos extras, guía + adicionales) y sus cajas contra el costo de guía y cajas de {ARCHIVO_HIST} para el mismo mes; "
                "para flete, una variación negativa es favorable; para cajas, una positiva.", mp_ST["nota"])]
    else:
        story.append(sin_datos(f"No se encontraron envíos de COBRO REGRESO en {titulo_mes.title()}."))

    # ---------------- 07 RANKING DE FLETERAS ----------------
    story += [PageBreak()] + seccion(7, "RANKING DE FLETERAS", "Efectividad de entregas, tiempos de tránsito y costo promedio (fleteras principales)")
    if not K.get("vacio"):
        res = K["res"]
        # Tabla resumen
        story.append(subtitulo("Resumen comparativo"))
        orden = res.sort_values("PCT_A_TIEMPO", ascending=False, na_position="last")
        rows = [[r.FLETERA, num(r.ENVIOS), num(r.ENTREGADOS), pct(r.PCT_A_TIEMPO, 0), pct(r.INCIDENCIAS_PCT, 0),
                 f"{r.DIAS_PROM:.1f}" if pd.notna(r.DIAS_PROM) else "-", money(r.COSTO_PROM_ENVIO),
                 money(r.COSTO_PROM_CAJA, 2) if pd.notna(r.COSTO_PROM_CAJA) else "-"] for r in orden.itertuples()]
        story.append(tabla(["FLETERA", "ENVÍOS", "ENTREGADOS", "% A TIEMPO", "% INCIDENCIAS", "DÍAS PROM.", "COSTO/ENVÍO", "COSTO/CAJA"], rows,
                           [96, 48, 62, 58, 66, 56, 74, 80], ["L", "R", "R", "R", "R", "R", "R", "R"]))
        story.append(Spacer(1, 12))

        # 6.1 Efectividad de entregas
        story.append(CondPageBreak(300))
        story.append(subtitulo("Efectividad de entregas  //  cumplimiento de promesa"))
        tp, ti = K["top_puntual"], K["top_incid"]
        story += kpi_row([("Efectividad global", f"{K['efectividad']:.0f}%", mp_HEX["green"]),
                          ("Pedidos entregados", num(K["entregados"]), mp_HEX["slate"]),
                          (f"Más puntual: {tp['FLETERA']}" if tp is not None else "Más puntual", pct(tp["PCT_A_TIEMPO"], 0) if tp is not None else "-", mp_HEX["gold"]),
                          (f"Más incidencias: {ti['FLETERA']}" if ti is not None else "Más incidencias", pct(ti["INCIDENCIAS_PCT"], 0) if ti is not None else "-", mp_HEX["orange"])], ncols=4)
        con_ent = res[res["ENTREGADOS"] > 0]
        if not con_ent.empty:
            a = con_ent.sort_values("PCT_A_TIEMPO", ascending=False)
            izq = chart_barras_h([(r.FLETERA, r.PCT_A_TIEMPO) for r in a.itertuples()], width=262, label_w=92, color=mp_C_GREEN, fmt=lambda v: f"{v:.0f}%")
            der = chart_apiladas_v([r.FLETERA for r in a.itertuples()],
                                   [("A tiempo", a["A_TIEMPO"].astype(int).tolist(), mp_C_GREEN), ("Con retraso", a["RETRASO"].astype(int).tolist(), C_ORANGE)],
                                   width=262, height=max(120, 28 + len(a) * 14))
            story.append(lado_a_lado(izq, der, "Ranking de efectividad (%)", "Entregas a tiempo vs. con retraso"))
        story.append(Spacer(1, 12))

        # 6.2 Tiempos
        story.append(CondPageBreak(260))
        story.append(subtitulo("Tiempos de tránsito  //  días de envío a entrega"))
        ra, le = K["rapida"], K["lenta"]
        story += kpi_row([("Tiempo promedio general", f"{K['dias_prom']:.1f} días", mp_HEX["teal"]),
                          ("Envíos analizados", num(K["dias_n"]), mp_HEX["slate"]),
                          (f"Más rápida: {ra['FLETERA']}" if ra is not None else "Más rápida", f"{ra['DIAS_PROM']:.1f} días" if ra is not None else "-", mp_HEX["green"]),
                          (f"Más lenta: {le['FLETERA']}" if le is not None else "Más lenta", f"{le['DIAS_PROM']:.1f} días" if le is not None else "-", mp_HEX["red"])], ncols=4)
        con_t = res[res["DIAS_PROM"].notna()].sort_values("DIAS_PROM", ascending=True)
        if not con_t.empty:
            story.append(KeepTogether([subtitulo("Días promedio de tránsito por fletera"),
                                       chart_barras_h([(r.FLETERA, r.DIAS_PROM) for r in con_t.itertuples()], label_w=130, color=C_BLUE, fmt=lambda v: f"{v:.1f} d")]))
        story.append(Spacer(1, 12))

        # 6.3 Costos
        story.append(CondPageBreak(330))
        story.append(subtitulo("Costo promedio por paquetería  //  guía + adicionales"))
        ba, ca = K["barata"], K["cara"]
        story += kpi_row([("Costo prom. por envío", money(K["costo_envio"]), mp_HEX["orange"]),
                          ("Costo prom. por caja", money(K["costo_caja"], 2), mp_HEX["gold"]),
                          (f"Más económica: {ba['FLETERA']}" if ba is not None else "Más económica", money(ba["COSTO_PROM_ENVIO"]) if ba is not None else "-", mp_HEX["green"]),
                          (f"Más cara: {ca['FLETERA']}" if ca is not None else "Más cara", money(ca["COSTO_PROM_ENVIO"]) if ca is not None else "-", mp_HEX["red"])], ncols=4)
        costo = res[res["ENVIOS"] > 0].sort_values("COSTO_PROM_ENVIO", ascending=True)
        caja = res[res["COSTO_PROM_CAJA"].notna()].sort_values("COSTO_PROM_CAJA", ascending=True)
        izq = chart_barras_h([(r.FLETERA, r.COSTO_PROM_ENVIO) for r in costo.itertuples()], width=262, label_w=92, color=C_ORANGE, fmt=money)
        der = chart_barras_h([(r.FLETERA, r.COSTO_PROM_CAJA) for r in caja.itertuples()], width=262, label_w=92, color=mp_C_GOLD, fmt=lambda v: money(v, 2)) if not caja.empty else sin_datos("Sin cajas registradas.")
        story.append(lado_a_lado(izq, der, "Costo promedio por envío", "Costo promedio por caja"))
        story.append(Spacer(1, 10))
        mapa = res[res["PCT_A_TIEMPO"].notna() & (res["ENVIOS"] > 0)]
        story.append(KeepTogether([subtitulo("Mapa de valor: costo promedio vs. efectividad por fletera"),
                                   chart_scatter([(r.FLETERA, r.COSTO_PROM_ENVIO, r.PCT_A_TIEMPO, r.ENVIOS) for r in mapa.itertuples()],
                                                 xtitle="COSTO PROMEDIO POR ENVÍO ($)", ytitle="% ENTREGAS A TIEMPO"),
                                   Paragraph("Más arriba = más puntual; más a la izquierda = más económica. El tamaño del círculo es el volumen de envíos.", mp_ST["nota"])]))
    else:
        story.append(sin_datos())

    # ---------------- 08 COSTOS DE MUESTRAS ----------------
    story += [PageBreak()] + seccion(8, "COSTOS DE MUESTRAS", "Inversión en producto y fletes de muestras del periodo")
    if M is None:
        story.append(sin_datos(muestras_msg or "No fue posible cargar la información de muestras."))
    elif not M["total"]:
        story.append(sin_datos(f"No hay envíos de muestras registrados en {titulo_mes.title()}."))
    else:
        story += kpi_row([("Total de envíos", num(M["total"]), mp_HEX["slate"]), ("Costo productos", money(M["prod"]), mp_HEX["blue"]),
                          ("Costo fletes", money(M["flete"]), mp_HEX["purple"]), ("Inversión total", money(M["inv"]), mp_HEX["green"])], ncols=4)
        story += kpi_row([("Costo promedio por envío", money(M["prom"]), mp_HEX["orange"]), ("% despachado", pct(M["pct_desp"], 0), mp_HEX["gold"]),
                          ("Agente con más envíos", trunc(M["top_agente"], 26), mp_HEX["red"])], ncols=3)
        story.append(KeepTogether([subtitulo("Inversión por mes (últimos 12 meses)"),
                                   chart_barras_v([c for c, _ in M["tendencia"]], [v for _, v in M["tendencia"]], height=140, color=mp_C_TEAL, fmt=lambda v: money(v))]))
        story.append(Spacer(1, 8))
        if not M["por_sol"].empty:
            rows = [[trunc(r.SOL, 34), num(r.envios), money(r.prod), money(r.flete), money(r.inversion)] for r in M["por_sol"].itertuples()]
            story.append(KeepTogether([subtitulo("Solicitantes / agentes (top 10 por inversión)"),
                                       tabla(["SOLICITANTE", "ENVÍOS", "PRODUCTOS", "FLETE", "INVERSIÓN"], rows, [200, 55, 95, 90, 100], ["L", "R", "R", "R", "R"])]))
        story.append(CondPageBreak(250))
        story.append(Spacer(1, 8))
        izq = chart_barras_h(list(M["top_dest"].items()), width=262, label_w=120, color=C_BLUE) if len(M["top_dest"]) else sin_datos("Sin datos.")
        der = chart_barras_h(list(M["flete_paq"].items()), width=262, label_w=100, color=C_PURPLE, fmt=money) if len(M["flete_paq"]) else sin_datos("Sin datos.")
        story.append(lado_a_lado(izq, der, "Top destinos / hoteles (envíos)", "Costo de flete por paquetería"))
        story.append(Spacer(1, 8))
        if len(M["prod_top"]):
            story.append(KeepTogether([subtitulo("Productos más solicitados (piezas)"),
                                       chart_barras_h([(str(k)[:40].upper(), v) for k, v in M["prod_top"].items()], label_w=210, color=mp_C_GOLD)]))

    # ---------------- CONSTRUCCIÓN ----------------
    buf = BytesIO()
    doc = BaseDocTemplate(buf, pagesize=letter, leftMargin=MARGIN, rightMargin=MARGIN, topMargin=48, bottomMargin=44,
                          title=f"Reporte Mensual de Logística - {titulo_mes}", author="JYPESA Logística")
    doc.addPageTemplates(_hacer_paginas(titulo_mes, f"Generado el {ahora}  ·  Fuente: Matriz de envíos y registro de muestras", logo_bytes))

    class _Canvas(_NumberedCanvas):
        titulo_pie = f"Reporte mensual {titulo_mes.title()}"
    doc.build(story, canvasmaker=_Canvas)
    buf.seek(0)
    return buf


# ============================================================
# 10. REPORTE INDEPENDIENTE: % LOGÍSTICO POR CONCEPTO
# ============================================================
MAX_BARRAS_CONCEPTO = 15   # solo para las gráficas; la tabla muestra TODOS los conceptos


def _sin_acentos(t):
    return "".join(c for c in unicodedata.normalize("NFD", str(t)) if unicodedata.category(c) != "Mn")


def _filtrar_periodo_modalidad(df_base, anio, mes, modalidad):
    """Periodo igual que Análisis Mensual (columna MES + validación de año; mes=0 -> todo el año) y modalidad por FORMA DE ENVIO."""
    d = df_base.copy()
    f_env = d["FECHA DE ENVÍO"]
    ok_anio = f_env.isna() | (f_env.dt.year == anio)
    d = d[ok_anio & ((d["MES"] == MESES[mes - 1]) if mes else (d["MES"] != ""))]
    if modalidad == "COBRO REGRESO":
        d = d[d["FORMA DE ENVIO"].str.contains("REGRESO", case=False, na=False)]
    elif modalidad == "COBRO DESTINO":
        d = d[d["FORMA DE ENVIO"].str.contains("DESTINO", case=False, na=False)]
    return d.copy()


def calc_logistico_concepto(df_base, anio, mes, modalidad="COBRO REGRESO", incluir_adic=True):
    """% logístico por CONCEPTO.  mes = 0 -> todo el año.  modalidad: 'COBRO REGRESO' | 'COBRO DESTINO' | 'TODAS'.
    Periodo igual que Análisis Mensual / sección 06 (columna MES + validación de año).
    % logístico = costo (guía [+ adicionales]) / facturación del concepto."""
    d_all = _filtrar_periodo_modalidad(df_base, anio, mes, modalidad)
    d, d_ext = _separar_extras(d_all)
    E = _resumen_extras(d_ext)
    if d.empty:
        return dict(vacio=True, motivo=("No hay pedidos con esos filtros; solo hay registros de recolecciones/maniobras (costos extras)."
                                        if E["n"] else "No hay envíos con esos filtros (periodo y modalidad)."))

    col_c = next((c for c in d.columns if "CONCEPTO" in str(c).upper()), None)
    if col_c is None:
        return dict(vacio=True, motivo="La matriz de envíos no trae la columna CONCEPTO.")

    d["_CONC"] = d[col_c].fillna("").astype(str).map(lambda x: _sin_acentos(x).strip().upper()).replace("", "SIN CONCEPTO")
    d["_COSTO"] = d["COSTO DE LA GUÍA"] + (d["COSTOS ADICIONALES"] if incluir_adic else 0.0)

    pc = d.groupby("_CONC").agg(
        envios=("_CONC", "size"), cajas=("CAJAS", "sum"), fact=("FACTURACION", "sum"),
        guia=("COSTO DE LA GUÍA", "sum"), adic=("COSTOS ADICIONALES", "sum"), costo=("_COSTO", "sum"),
    ).reset_index().rename(columns={"_CONC": "CONCEPTO"})
    pc = pc.sort_values("costo", ascending=False).reset_index(drop=True)

    costo, fact, cajas = d["_COSTO"].sum(), d["FACTURACION"].sum(), d["CAJAS"].sum()
    pc["pct_gasto"] = (pc["costo"] / costo * 100) if costo else 0.0
    pc["costo_caja"] = pc["costo"] / pc["cajas"].replace(0, np.nan)
    pc["pct_log"] = pc["costo"] / pc["fact"].replace(0, np.nan) * 100
    pc["vs_target"] = pc["pct_log"] - TARGET_COSTO_LOG

    return dict(vacio=False, registros=len(d), pc=pc,
                guia=d["COSTO DE LA GUÍA"].sum(), adic=d["COSTOS ADICIONALES"].sum(),
                costo=costo, fact=fact, cajas=cajas,
                pct_log=(costo / fact * 100) if fact else 0.0,
                costo_caja=(costo / cajas) if cajas else 0.0,
                extras=E, costo_total=costo + E["costo"],
                pct_log_total=((costo + E["costo"]) / fact * 100) if fact else 0.0)


def _encabezado_simple(titulo, subtitulo_=""):
    cont = [Paragraph(esc(titulo), mp_ST["sec_t"])]
    if subtitulo_:
        cont.append(Paragraph(esc(subtitulo_), mp_ST["sec_s"]))
    t = Table([[cont]], colWidths=[CONTENT_W])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), mp_C_SLATE), ("LINEBEFORE", (0, 0), (0, 0), 5, mp_C_TEAL),
        ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 12),
        ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 9)]))
    return [t, Spacer(1, 10)]


def _hacer_paginas_simple(titulo_hdr, periodo_txt):
    def normal(c, doc):
        c.saveState()
        c.setFillColor(mp_C_NAVY)
        c.rect(0, PAGE_H - 30, PAGE_W, 30, stroke=0, fill=1)
        c.setFillColor(mp_C_TEAL)
        c.rect(0, PAGE_H - 30, PAGE_W, 2.5, stroke=0, fill=1)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 8.5)
        c.drawString(MARGIN, PAGE_H - 19, titulo_hdr)
        c.setFillColor(mp_C_GOLD)
        c.drawRightString(PAGE_W - MARGIN, PAGE_H - 19, periodo_txt)
        c.restoreState()
    f = Frame(MARGIN, 44, CONTENT_W, PAGE_H - 44 - 48, id="fn", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    return [PageTemplate(id="normal", frames=[f], onPage=normal)]


def generar_reporte_concepto_pdf(df_base, anio, mes, modalidad="COBRO REGRESO", incluir_adic=True):
    """PDF del % logístico por concepto.  df_base: salida de preparar_base().  mes=0 -> todo el año.  Devuelve BytesIO."""
    periodo = f"{MESES[mes - 1]} {anio}" if mes else f"AÑO {anio} (ACUMULADO)"
    ahora = datetime.now().strftime("%d/%m/%Y %H:%M")
    base_txt = "guía + costos adicionales" if incluir_adic else "solo costo de guía"
    G = calc_logistico_concepto(df_base, anio, mes, modalidad, incluir_adic)

    story = _encabezado_simple("% LOGÍSTICO POR CONCEPTO",
                               f"{modalidad.title() if modalidad != 'TODAS' else 'Todas las modalidades'}  ·  {periodo.title()}  ·  costo = {base_txt}")
    if G["vacio"]:
        story.append(sin_datos(G["motivo"]))
    else:
        pc = G["pc"]
        col_log = mp_HEX["green"] if G["pct_log"] <= TARGET_COSTO_LOG else mp_HEX["red"]
        story += kpi_row([
            ("Facturación", money(G["fact"], 2), mp_HEX["green"]),
            (f"Costo logístico de pedidos ({base_txt})", money(G["costo"], 2), mp_HEX["slate"]),
            ("Cajas enviadas", num(G["cajas"]), mp_HEX["teal"]),
            (f"% logístico de pedidos (target {TARGET_COSTO_LOG}%)", f"{G['pct_log']:.2f}%", col_log),
        ], ncols=4)

        con_pct = pc[pc["pct_log"].notna()]
        fuera = int((con_pct["pct_log"] > TARGET_COSTO_LOG).sum())
        if not con_pct.empty:
            peor = con_pct.sort_values("pct_log", ascending=False).iloc[0]
            peor_txt = f"{trunc(peor['CONCEPTO'], 16)}  {peor['pct_log']:.2f}%"
        else:
            peor, peor_txt = None, "-"
        story += kpi_row([
            ("Costo por caja", money(G["costo_caja"], 2), mp_HEX["orange"]),
            ("Envíos analizados", num(G["registros"]), mp_HEX["blue"]),
            ("Concepto con mayor % logístico", peor_txt, mp_HEX["red"]),
            ("Conceptos fuera de target", f"{fuera} de {len(con_pct)}", mp_HEX["gold"]),
        ], ncols=4)

        # Gráficas (máx. MAX_BARRAS_CONCEPTO barras; la tabla trae todos los conceptos)
        if not con_pct.empty:
            ord_pct = con_pct.sort_values("pct_log", ascending=False).head(MAX_BARRAS_CONCEPTO)
            izq = chart_barras_h([(r.CONCEPTO, r.pct_log) for r in ord_pct.itertuples()], width=262, label_w=100,
                                 fmt=lambda v: f"{v:.2f}%",
                                 colors_list=[mp_C_GREEN if v <= TARGET_COSTO_LOG else mp_C_RED for v in ord_pct["pct_log"]])
        else:
            izq = sin_datos("Ningún concepto tiene facturación.")
        der = chart_barras_h([(r.CONCEPTO, r.costo) for r in pc.head(MAX_BARRAS_CONCEPTO).itertuples()],
                             width=262, label_w=100, color=C_ORANGE, fmt=money)
        story.append(lado_a_lado(izq, der, f"% logístico por concepto (verde = dentro de {TARGET_COSTO_LOG}%)", "Costo logístico por concepto"))
        story.append(Spacer(1, 10))

        def _vs(v):
            return "-" if pd.isna(v) else f"{v:+.2f} pp"
        rows = [[trunc(r.CONCEPTO, 24), num(r.envios), num(r.cajas), money(r.fact, 2), money(r.guia, 2), money(r.adic, 2),
                 pct(r.pct_log, 2), _vs(r.vs_target), money(r.costo_caja, 2) if pd.notna(r.costo_caja) else "-"]
                for r in pc.itertuples()]
        t_conc = tabla(["CONCEPTO", "ENVÍOS", "CAJAS", "FACTURACIÓN", "COSTO GUÍA", "ADICIONALES", "% LOGÍSTICO", f"VS TARGET {TARGET_COSTO_LOG}%", "COSTO/CAJA"],
                       rows, [92, 38, 42, 78, 68, 62, 52, 52, 56],
                       ["L", "R", "R", "R", "R", "R", "R", "R", "R"],
                       total_row=["TOTAL", num(G["registros"]), num(G["cajas"]), money(G["fact"], 2), money(G["guia"], 2), money(G["adic"], 2),
                                  pct(G["pct_log"], 2), _vs(G["pct_log"] - TARGET_COSTO_LOG), money(G["costo_caja"], 2)])
        if len(pc) <= 8:   # tabla corta: completa en una hoja
            story.append(KeepTogether([subtitulo("Detalle por concepto"), t_conc]))
        else:              # tabla larga: puede continuar en la siguiente hoja (repite encabezado)
            story += [CondPageBreak(120), subtitulo("Detalle por concepto"), t_conc]

        # ---- costos extras (recolecciones / maniobras) ----
        E = G["extras"]
        if E["n"]:
            col_t = mp_HEX["green"] if G["pct_log_total"] <= TARGET_COSTO_LOG else mp_HEX["red"]
            story += [CondPageBreak(170), Spacer(1, 8), subtitulo("Costos extras (recolecciones y maniobras)")]
            story += kpi_row([
                ("Costos extras", money(E["costo"], 2), mp_HEX["gold"]),
                ("Costo logístico total (pedidos + extras)", money(G["costo_total"], 2), mp_HEX["slate"]),
                (f"% logístico con extras (target {TARGET_COSTO_LOG}%)", f"{G['pct_log_total']:.2f}%", col_t),
                ("Extras como % del costo total", pct(E["costo"] / G["costo_total"] * 100 if G["costo_total"] else 0, 1), mp_HEX["purple"]),
            ], ncols=4)
            pce = E["por_concepto"]
            rows_e = [[trunc(r.CONCEPTO, 30), num(r.n), money(r.guia, 2), money(r.adic, 2), money(r.costo, 2),
                       pct(r.costo / G["costo_total"] * 100 if G["costo_total"] else 0, 1)] for r in pce.itertuples()]
            story.append(tabla(["CONCEPTO", "REGISTROS", "COSTO GUÍA", "ADICIONALES", "COSTO TOTAL", "% DEL COSTO TOTAL"],
                               rows_e, [150, 64, 86, 80, 86, 74], ["L", "R", "R", "R", "R", "R"],
                               total_row=(["TOTAL", num(E["n"]), money(E["guia"], 2), money(E["adic"], 2), money(E["costo"], 2),
                                           pct(E["costo"] / G["costo_total"] * 100 if G["costo_total"] else 0, 1)] if len(pce) > 1 else None)))

        story += [Spacer(1, 8), Paragraph("LECTURA RÁPIDA", ParagraphStyle("lr", fontName="Helvetica-Bold", fontSize=10, textColor=mp_C_SLATE, spaceAfter=4))]
        if E["n"]:
            story.append(Paragraph(f"Los costos extras (recolecciones y maniobras) suman <b>{money(E['costo'])}</b> "
                                   f"({E['costo'] / G['costo_total'] * 100 if G['costo_total'] else 0:.0f}% del costo total); con ellos el % logístico pasa de "
                                   f"<b>{G['pct_log']:.2f}%</b> (solo pedidos) a <b>{G['pct_log_total']:.2f}%</b>.", mp_ST["bullet"], bulletText="•"))
        if peor is not None:
            story.append(Paragraph(f"El concepto con mayor % logístico es <b>{esc(peor['CONCEPTO'])}</b> con <b>{peor['pct_log']:.2f}%</b> "
                                   f"({money(peor['costo'])} de costo sobre {money(peor['fact'])} facturados).", mp_ST["bullet"], bulletText="•"))
            story.append(Paragraph(f"<b>{fuera}</b> de {len(con_pct)} conceptos con facturación están por encima del target de {TARGET_COSTO_LOG}%.",
                                   mp_ST["bullet"], bulletText="•"))
        top_gasto = pc.iloc[0]
        story.append(Paragraph(f"El mayor gasto está en <b>{esc(top_gasto['CONCEPTO'])}</b>: {money(top_gasto['costo'])} "
                               f"({top_gasto['pct_gasto']:.0f}% del costo total).", mp_ST["bullet"], bulletText="•"))
        sin_fact = pc[(pc["fact"] <= 0) & (pc["costo"] > 0)]
        if not sin_fact.empty:
            nombres = ", ".join(esc(trunc(x, 22)) for x in sin_fact["CONCEPTO"].head(4))
            story.append(Paragraph(f"Sin facturación pero con costo: <b>{nombres}</b> ({money(sin_fact['costo'].sum())} en total). "
                                   "No tienen % logístico propio, pero su costo sí está incluido en el total.", mp_ST["bullet"], bulletText="•"))

        story += [Spacer(1, 6), Paragraph(
            f"% logístico = costo ({base_txt}) / facturación del concepto. Vs target = diferencia en puntos porcentuales (pp) contra la meta de {TARGET_COSTO_LOG}%. "
            "Conceptos sin facturación aparecen con '-'. Los registros con concepto de recolecciones o maniobras no se cuentan como pedidos: van aparte como costos extras. "
            "Periodo según la columna MES de la matriz y el año seleccionado."
            + (f" Las gráficas muestran los {MAX_BARRAS_CONCEPTO} conceptos principales; la tabla incluye todos." if len(pc) > MAX_BARRAS_CONCEPTO else ""),
            mp_ST["nota"])]

    buf = BytesIO()
    doc = BaseDocTemplate(buf, pagesize=letter, leftMargin=MARGIN, rightMargin=MARGIN, topMargin=48, bottomMargin=44,
                          title=f"% Logístico por concepto - {periodo}", author="JYPESA Logística")
    doc.addPageTemplates(_hacer_paginas_simple("REPORTE DE % LOGÍSTICO POR CONCEPTO", periodo))

    class _Canvas(_NumberedCanvas):
        titulo_pie = f"% logístico por concepto  -  generado {ahora}"
    doc.build(story, canvasmaker=_Canvas)
    buf.seek(0)
    return buf


# ============================================================
# 11. REPORTE INDEPENDIENTE: SALUD DE PEDIDOS PEQUEÑOS (1 A 4 CAJAS)
# ============================================================
# Costo de flete = guía; costo de distribución = costos adicionales (maniobras, reexpediciones, etc.).
# Semáforo por pedido, según su % logístico individual (costo / facturación del pedido):
#   SALUDABLE <= target | EN ALERTA > target | CRÍTICO > 2 x target | PÉRDIDA >= 100% | SIN FACTURACIÓN
CLASES_SALUD = [   # (clave, etiqueta corta, color)
    ("SALUDABLE", "Saludable", mp_C_GREEN),
    ("EN ALERTA", "En alerta", mp_C_GOLD),
    ("CRÍTICO", "Crítico", C_ORANGE),
    ("PÉRDIDA", "Pérdida", mp_C_RED),
    ("SIN FACTURACIÓN", "Sin fact.", mp_C_GRAY),
]
CLASES_FUERA_TARGET = ["EN ALERTA", "CRÍTICO", "PÉRDIDA", "SIN FACTURACIÓN"]


def _m(v, dec=0):
    """money() que muestra '-' cuando no hay dato."""
    return "-" if v is None or pd.isna(v) else money(v, dec)


def _vs_target(v):
    return "-" if v is None or pd.isna(v) else f"{v - TARGET_COSTO_LOG:+.2f} pp"


def _color_pct(v):
    """Verde si está dentro del target, rojo si no (gris si no hay dato)."""
    if v is None or pd.isna(v):
        return mp_C_GRAY
    return mp_C_GREEN if v <= TARGET_COSTO_LOG else mp_C_RED


def _etq_tam(n, max_cajas):
    if n > max_cajas:
        return f"{max_cajas + 1}+ CAJAS"
    return f"{int(n)} CAJA" + ("S" if n > 1 else "")


def _bloque(g):
    """Indicadores de un grupo de pedidos (usa columnas _CAJ, _COSTO, _EXC, _CLASE preparadas)."""
    n = len(g)
    fact, costo, cajas = g["FACTURACION"].sum(), g["_COSTO"].sum(), g["_CAJ"].sum()
    costo_ped = (costo / n) if n else np.nan
    ticket_min = (costo_ped / (TARGET_COSTO_LOG / 100)) if n else np.nan   # facturación mínima para estar en target
    return dict(
        n=n, cajas=cajas, fact=fact, guia=g["COSTO DE LA GUÍA"].sum(), adic=g["COSTOS ADICIONALES"].sum(), costo=costo,
        pct_log=(costo / fact * 100) if fact > 0 else np.nan,
        costo_caja=(costo / cajas) if cajas else np.nan,
        costo_ped=costo_ped, ticket=(fact / n) if n else np.nan, ticket_min=ticket_min,
        fact_caja=(fact / cajas) if cajas else np.nan,
        pct_fuera=(g["_CLASE"].isin(CLASES_FUERA_TARGET).sum() / n * 100) if n else np.nan,
        pct_bajo_min=((g["FACTURACION"] < ticket_min).sum() / n * 100) if n else np.nan,
        exceso=g["_EXC"].sum())


def _df_bloques(d, col):
    filas = []
    for k, g in d.groupby(col, sort=False):
        b = _bloque(g)
        b["K"] = k   # clave del grupo (nombre sin guion bajo para poder usar itertuples)
        filas.append(b)
    return pd.DataFrame(filas)


def _preparar_pequenos(df_base, anio, modalidad, incluir_adic, max_cajas):
    """Filtra año y modalidad (sin filtrar mes) y agrega columnas de análisis. Devuelve (df, sin_cajas)."""
    d = df_base.copy()
    f_env = d["FECHA DE ENVÍO"]
    d = d[(f_env.isna() | (f_env.dt.year == anio)) & (d["MES"] != "")].copy()
    if modalidad == "COBRO REGRESO":
        d = d[d["FORMA DE ENVIO"].str.contains("REGRESO", case=False, na=False)]
    elif modalidad == "COBRO DESTINO":
        d = d[d["FORMA DE ENVIO"].str.contains("DESTINO", case=False, na=False)]
    d = d.copy()

    # Cajas del pedido: CANTIDAD DE CAJAS; si viene en 0 se usa CAJAS (columna de cobro regreso)
    d["_CAJ"] = np.where(d["CANTIDAD DE CAJAS"] > 0, d["CANTIDAD DE CAJAS"], d["CAJAS"]).astype(float).round()
    sin_cajas = int((d["_CAJ"] < 1).sum())
    d = d[d["_CAJ"] >= 1].copy()
    d["_PEQ"] = d["_CAJ"] <= max_cajas
    d["_TAMN"] = np.minimum(d["_CAJ"], max_cajas + 1).astype(int)

    d["_COSTO"] = d["COSTO DE LA GUÍA"] + (d["COSTOS ADICIONALES"] if incluir_adic else 0.0)
    fact, costo = d["FACTURACION"], d["_COSTO"]
    d["_PCT"] = costo / fact.where(fact > 0) * 100          # NaN si el pedido no tiene facturación
    conds = [(fact <= 0) & (costo <= 0), fact <= 0, d["_PCT"] >= 100,
             d["_PCT"] > 2 * TARGET_COSTO_LOG, d["_PCT"] > TARGET_COSTO_LOG]
    d["_CLASE"] = np.select(conds, ["SIN DATOS", "SIN FACTURACIÓN", "PÉRDIDA", "CRÍTICO", "EN ALERTA"], default="SALUDABLE")
    d["_EXC"] = (costo - TARGET_COSTO_LOG / 100 * fact).clip(lower=0)   # lo que se paga por encima del target
    return d, sin_cajas


def calc_pequenos(df_base, anio, mes, modalidad="TODAS", incluir_adic=True, max_cajas=CAJAS_PEQUENO_MAX):
    """Salud de pedidos pequeños (1 a max_cajas cajas) vs el resto. mes = 0 -> todo el año."""
    df_base, df_ext = _separar_extras(df_base)          # recolecciones / maniobras no son pedidos
    E = calc_extras(df_ext, anio, mes, modalidad)
    anual, sin_cajas = _preparar_pequenos(df_base, anio, modalidad, incluir_adic, max_cajas)
    d = anual[anual["MES"] == MESES[mes - 1]] if mes else anual
    if d.empty:
        return dict(vacio=True, motivo="No hay envíos con cajas registradas para ese periodo y modalidad.")
    peq, resto = d[d["_PEQ"]], d[~d["_PEQ"]]
    if peq.empty:
        return dict(vacio=True, motivo=f"No hay pedidos de 1 a {max_cajas} cajas en ese periodo y modalidad.")

    B_peq, B_res, B_tot = _bloque(peq), (_bloque(resto) if not resto.empty else None), _bloque(d)

    por_tam = _df_bloques(d.sort_values("_TAMN"), "_TAMN")
    por_clase = _df_bloques(peq, "_CLASE").set_index("K")
    cruce = pd.crosstab(peq["_TAMN"], peq["_CLASE"])

    def _ranking(col, n=10):
        x = peq.assign(_K=peq[col].replace("", pd.NA)).dropna(subset=["_K"])
        if x.empty:
            return pd.DataFrame()
        return _df_bloques(x, "_K").sort_values("exceso", ascending=False).head(n).reset_index(drop=True)

    pf = peq.assign(_F=peq["FLETERA"].replace("", "SIN ASIGNAR"))
    por_f = _df_bloques(pf, "_F").sort_values("costo", ascending=False).reset_index(drop=True)
    if not resto.empty:
        rf = _df_bloques(resto.assign(_F=resto["FLETERA"].replace("", "SIN ASIGNAR")), "_F").set_index("K")["pct_log"]
        por_f["pct_log_resto"] = por_f["K"].map(rf)
    else:
        por_f["pct_log_resto"] = np.nan
    por_forma = _df_bloques(peq.assign(_K=peq["FORMA DE ENVIO"].replace("", "SIN FORMA")), "_K").sort_values("costo", ascending=False)

    # Bandas de facturación por pedido (cuartiles de los pedidos pequeños con facturación)
    bandas = pd.DataFrame()
    v = peq[peq["FACTURACION"] > 0]
    if len(v) >= 8:
        try:
            cat, edges = pd.qcut(v["FACTURACION"], 4, retbins=True, duplicates="drop")
            etiquetas = {iv: f"{money(iv.left if i else 0)} a {money(iv.right)}" for i, iv in enumerate(cat.cat.categories)}
            v = v.assign(_B=cat.map(etiquetas).astype(str), _O=cat.cat.codes)
            bandas = _df_bloques(v.sort_values("_O"), "_B")
        except Exception:
            bandas = pd.DataFrame()

    # Tendencia del año (hasta el mes elegido): % logístico de pequeños vs resto
    def _pl(g):
        f = g["FACTURACION"].sum()
        return (g["_COSTO"].sum() / f * 100) if f > 0 else np.nan
    tendencia = []
    for i, m in enumerate(MESES):
        if mes and i + 1 > mes:
            break
        g = anual[anual["MES"] == m]
        if g.empty or not g["_PEQ"].any():
            continue
        tendencia.append((m[:3].title(), _pl(g[g["_PEQ"]]), _pl(g[~g["_PEQ"]]), int(g["_PEQ"].sum())))

    peores = peq[peq["_EXC"] > 0].sort_values("_EXC", ascending=False).head(15)
    exceso_total = peq["_EXC"].sum()
    clientes, destinos = _ranking("NOMBRE DEL CLIENTE"), _ranking("DESTINO")
    conc_top = (clientes["exceso"].head(10).sum() / exceso_total * 100) if (exceso_total > 0 and not clientes.empty) else np.nan

    return dict(vacio=False, d=d, peq=peq, B=B_peq, R=B_res, T=B_tot, max_cajas=max_cajas, sin_cajas=sin_cajas,
                sin_costo=int((peq["_COSTO"] <= 0).sum()), por_tam=por_tam, por_clase=por_clase, cruce=cruce,
                por_f=por_f, por_forma=por_forma, bandas=bandas, tendencia=tendencia, peores=peores,
                clientes=clientes, destinos=destinos, exceso_total=exceso_total, conc_top=conc_top,
                sh_ped=B_peq["n"] / B_tot["n"] * 100, sh_fact=(B_peq["fact"] / B_tot["fact"] * 100) if B_tot["fact"] else np.nan,
                sh_costo=(B_peq["costo"] / B_tot["costo"] * 100) if B_tot["costo"] else np.nan,
                sh_cajas=B_peq["cajas"] / B_tot["cajas"] * 100, extras=E)


def generar_reporte_pequenos_pdf(df_base, anio, mes, modalidad="TODAS", incluir_adic=True, max_cajas=CAJAS_PEQUENO_MAX):
    """PDF de salud de pedidos pequeños (1 a max_cajas cajas): facturación, flete y distribución. Devuelve BytesIO."""
    periodo = f"{MESES[mes - 1]} {anio}" if mes else f"AÑO {anio} (ACUMULADO)"
    ahora = datetime.now().strftime("%d/%m/%Y %H:%M")
    base_txt = "guía + costos adicionales" if incluir_adic else "solo costo de guía"
    rng = f"1 a {max_cajas} cajas" if max_cajas > 1 else "1 caja"
    P = calc_pequenos(df_base, anio, mes, modalidad, incluir_adic, max_cajas)

    story = _encabezado_simple(f"SALUD DE PEDIDOS PEQUEÑOS  ·  {rng.upper()}",
                               f"{modalidad.title() if modalidad != 'TODAS' else 'Todas las modalidades'}  ·  {periodo.title()}  ·  costo = {base_txt}")
    if P["vacio"]:
        story.append(sin_datos(P["motivo"]))
    else:
        B, R, T = P["B"], P["R"], P["T"]
        mc = max_cajas
        # ---------------- RESUMEN EJECUTIVO ----------------
        story.append(Paragraph("RESUMEN EJECUTIVO", ParagraphStyle("re", fontName="Helvetica-Bold", fontSize=11, textColor=mp_C_SLATE, spaceAfter=8)))
        col_l = mp_HEX["green"] if B["pct_log"] <= TARGET_COSTO_LOG else mp_HEX["red"]
        story += kpi_row([
            (f"Pedidos de {rng}", f"{num(B['n'])}  ·  {P['sh_ped']:.0f}%", mp_HEX["slate"]),
            ("Facturación pequeños", f"{money(B['fact'])}  ·  {pct(P['sh_fact'], 0)}", mp_HEX["green"]),
            (f"Costo logístico ({base_txt})", f"{money(B['costo'])}  ·  {pct(P['sh_costo'], 0)}", mp_HEX["orange"]),
            (f"% logístico pequeños (target {TARGET_COSTO_LOG}%)", pct(B["pct_log"], 2), col_l),
        ], ncols=4)
        res_txt = pct(R["pct_log"], 2) if R else "-"
        story += kpi_row([
            (f"% logístico de {mc + 1}+ cajas", res_txt, mp_HEX["gray"] if not R else (mp_HEX["green"] if R["pct_log"] <= TARGET_COSTO_LOG else mp_HEX["red"])),
            ("Costo por caja (pequeños)", _m(B["costo_caja"], 2) + (f"  vs {_m(R['costo_caja'], 2)}" if R else ""), mp_HEX["gold"]),
            ("Costo por pedido (pequeños)", _m(B["costo_ped"], 2), mp_HEX["purple"]),
            ("Sobrecosto vs target", money(P["exceso_total"]), mp_HEX["red"]),
        ], ncols=4)
        story += kpi_row([
            ("Ticket promedio", _m(B["ticket"]), mp_HEX["teal"]),
            ("Ticket mínimo para estar en target", _m(B["ticket_min"]), mp_HEX["blue"]),
            ("Pedidos fuera de target", pct(B["pct_fuera"], 0), mp_HEX["red"]),
            ("Pedidos bajo el ticket mínimo", pct(B["pct_bajo_min"], 0), mp_HEX["orange"]),
        ], ncols=4)

        hall = []
        hall.append(f"Los pedidos de {rng} son <b>{P['sh_ped']:.0f}%</b> de los envíos y <b>{P['sh_cajas']:.0f}%</b> de las cajas, pero generan solo "
                    f"<b>{pct(P['sh_fact'], 0)}</b> de la facturación y cargan <b>{pct(P['sh_costo'], 0)}</b> del costo logístico.")
        if R and not pd.isna(B["pct_log"]) and not pd.isna(R["pct_log"]):
            dif = B["pct_log"] - R["pct_log"]
            hall.append(f"Su % logístico es <b>{B['pct_log']:.2f}%</b> contra <b>{R['pct_log']:.2f}%</b> de los pedidos de {mc + 1}+ cajas "
                        f"({dif:+.2f} pp) y {'rebasa' if B['pct_log'] > TARGET_COSTO_LOG else 'cumple'} el target de {TARGET_COSTO_LOG}%.")
        pt = P["por_tam"][P["por_tam"]["K"] <= mc]
        pt_f = pt[pt["pct_log"].notna()]
        if not pt_f.empty:
            w = pt_f.sort_values("pct_log", ascending=False).iloc[0]
            hall.append(f"El tamaño más pesado es <b>{_etq_tam(w['K'], mc)}</b> con <b>{w['pct_log']:.2f}%</b> logístico "
                        f"({_m(w['costo_ped'], 2)} de costo por pedido sobre un ticket de {_m(w['ticket'])}).")
        hall.append(f"<b>{B['pct_fuera']:.0f}%</b> de los pedidos pequeños supera el target; el sobrecosto acumulado frente a la meta es de <b>{money(P['exceso_total'])}</b>. "
                    f"Los pedidos con costo igual o mayor a su facturación (pérdida) o sin facturación suman "
                    f"<b>{num(P['por_clase']['n'].get('PÉRDIDA', 0) + P['por_clase']['n'].get('SIN FACTURACIÓN', 0))}</b>.")
        if not pd.isna(B["ticket_min"]):
            hall.append(f"Con el costo promedio actual ({_m(B['costo_ped'], 2)} por pedido), se necesita facturar al menos <b>{money(B['ticket_min'])}</b> por pedido "
                        f"para quedar en target; el ticket promedio es {_m(B['ticket'])} y <b>{B['pct_bajo_min']:.0f}%</b> de los pedidos está por debajo de ese mínimo.")
        pf_ = P["por_f"][(P["por_f"]["n"] >= 3) & P["por_f"]["pct_log"].notna()]
        if len(pf_) >= 2:
            a, z = pf_.sort_values("pct_log").iloc[0], pf_.sort_values("pct_log").iloc[-1]
            hall.append(f"Por fletera (mín. 3 pedidos): la más eficiente en pedidos pequeños es <b>{esc(a['K'])}</b> ({a['pct_log']:.2f}%) y la más costosa "
                        f"<b>{esc(z['K'])}</b> ({z['pct_log']:.2f}%).")
        if not pd.isna(P["conc_top"]):
            hall.append(f"Los 10 clientes con mayor sobrecosto concentran <b>{P['conc_top']:.0f}%</b> del exceso: ahí está la mayor oportunidad de negociación o de pedido mínimo.")
        story.append(Paragraph("HALLAZGOS CLAVE", ParagraphStyle("hk", fontName="Helvetica-Bold", fontSize=11, textColor=mp_C_SLATE, spaceBefore=4, spaceAfter=6)))
        if P["extras"]["n"]:
            hall.append(f"No se incluyeron <b>{num(P['extras']['n'])}</b> registros de recolecciones y maniobras ({money(P['extras']['costo'])}): "
                        "no son pedidos, se tratan como costos extras.")
        for h in hall:
            story.append(Paragraph(h, mp_ST["bullet"], bulletText="•"))

        # ---------------- 01 PEQUEÑOS VS RESTO, POR TAMAÑO ----------------
        story += [PageBreak()] + seccion(1, "PEQUEÑOS VS RESTO, POR TAMAÑO DE PEDIDO",
                                         f"Facturación, flete y distribución de pedidos de {rng} frente a {mc + 1}+ cajas")
        W9 = [66, 40, 36, 76, 60, 44, 48, 54, 56, 60]
        H9 = ["TAMAÑO", "PEDIDOS", "CAJAS", "FACTURACIÓN", "COSTO", "% LOG.", f"VS {TARGET_COSTO_LOG}%", "COSTO / CAJA", "COSTO / PEDIDO", "TICKET PROM."]
        A9 = ["L"] + ["R"] * 9

        def _fila(lab, b):
            return [lab, num(b["n"]), num(b["cajas"]), money(b["fact"]), money(b["costo"]), pct(b["pct_log"], 2), _vs_target(b["pct_log"]),
                    _m(b["costo_caja"], 2), _m(b["costo_ped"], 2), _m(b["ticket"])]
        filas = [_fila(f"PEQUEÑOS (1-{mc})", B)] + ([_fila(f"RESTO ({mc + 1}+)", R)] if R else [])
        story.append(tabla(H9, filas, W9, A9, total_row=_fila("TOTAL", T)))
        story.append(Spacer(1, 10))
        pt_all = P["por_tam"]
        cats = [_etq_tam(n, mc).replace(" CAJAS", " cj").replace(" CAJA", " cj") for n in pt_all["K"]]
        izq = chart_barras_v(cats, [0 if pd.isna(v) else v for v in pt_all["pct_log"]], width=262, height=150, fmt=lambda v: f"{v:.1f}%",
                             colors_list=[_color_pct(v) for v in pt_all["pct_log"]])
        der = chart_barras_v(cats, [0 if pd.isna(v) else v for v in pt_all["costo_caja"]], width=262, height=150, color=mp_C_GOLD, fmt=lambda v: money(v, 0))
        story.append(lado_a_lado(izq, der, f"% logístico por tamaño (verde = dentro de {TARGET_COSTO_LOG}%)", "Costo por caja ($)"))
        story.append(Spacer(1, 8))
        izq = chart_barras_h([("% de los pedidos", P["sh_ped"]), ("% de las cajas", P["sh_cajas"]), ("% de la facturación", P["sh_fact"] if not pd.isna(P["sh_fact"]) else 0),
                              ("% del costo logístico", P["sh_costo"] if not pd.isna(P["sh_costo"]) else 0)],
                             width=262, label_w=104, fmt=lambda v: f"{v:.0f}%", colors_list=[mp_C_SLATE, mp_C_TEAL, mp_C_GREEN, C_ORANGE])
        der = chart_barras_v(cats, [0 if pd.isna(v) else v for v in pt_all["ticket"]], width=262, height=130, color=mp_C_GREEN, fmt=lambda v: money(v, 0))
        story.append(lado_a_lado(izq, der, f"Peso de los pedidos de {rng} en el total", "Ticket promedio por pedido ($)"))
        story.append(Spacer(1, 10))
        rows = [_fila(_etq_tam(r["K"], mc), r) for _, r in pt_all.iterrows()]
        story.append(KeepTogether([subtitulo("Detalle por número de cajas"), tabla(H9, rows, W9, A9)]))
        story += [Spacer(1, 4), Paragraph(
            f"Costo = {base_txt}. % logístico = costo / facturación. Ticket = facturación promedio por pedido. "
            f"Pedidos sin cajas registradas ({num(P['sin_cajas'])}) quedan fuera del análisis.", mp_ST["nota"])]

        # ---------------- 02 SEMÁFORO DE SALUD ----------------
        story += [Spacer(1, 12), CondPageBreak(340)] + seccion(2, "SEMÁFORO DE SALUD DEL PEDIDO",
                                         f"Cada pedido de {rng} clasificado por su % logístico individual (target {TARGET_COSTO_LOG}%)")
        pc = P["por_clase"]
        presentes = [(k, e, c) for k, e, c in CLASES_SALUD if k in pc.index]
        n_pq = B["n"]
        story += kpi_row([(e, f"{num(pc.loc[k, 'n'])}  ·  {pc.loc[k, 'n'] / n_pq * 100:.0f}%", c.hexval().replace("0x", "#")) for k, e, c in presentes],
                         ncols=len(presentes))
        izq = chart_donut([(e, pc.loc[k, "n"], c) for k, e, c in presentes])
        der = chart_barras_h([(e, pc.loc[k, "costo"]) for k, e, c in presentes], width=262, label_w=70, fmt=money,
                             colors_list=[c for _, _, c in presentes])
        story.append(lado_a_lado(izq, der, "Pedidos por estado de salud", "Costo logístico por estado ($)"))
        story.append(Spacer(1, 10))
        rows = [[e, num(pc.loc[k, "n"]), pct(pc.loc[k, "n"] / n_pq * 100, 0), money(pc.loc[k, "fact"]),
                 pct(pc.loc[k, "fact"] / B["fact"] * 100, 0) if B["fact"] else "-", money(pc.loc[k, "costo"]),
                 pct(pc.loc[k, "costo"] / B["costo"] * 100, 0) if B["costo"] else "-", money(pc.loc[k, "exceso"]),
                 _m(pc.loc[k, "costo_ped"], 2)] for k, e, c in presentes]
        story.append(KeepTogether([subtitulo("Detalle por estado"),
                                   tabla(["ESTADO", "PEDIDOS", "% PED.", "FACTURACIÓN", "% FACT.", "COSTO", "% COSTO", "SOBRECOSTO", "COSTO / PEDIDO"], rows,
                                         [78, 48, 44, 78, 48, 70, 50, 70, 54], ["L"] + ["R"] * 8,
                                         total_row=["TOTAL", num(n_pq), "100%", money(B["fact"]), "100%", money(B["costo"]), "100%", money(P["exceso_total"]), _m(B["costo_ped"], 2)])]))
        story.append(Spacer(1, 10))
        cr = P["cruce"]
        cats_t = [_etq_tam(n, mc).replace(" CAJAS", " cj").replace(" CAJA", " cj") for n in cr.index]
        series = [(e, [int(cr.loc[n, k]) if k in cr.columns else 0 for n in cr.index], c) for k, e, c in CLASES_SALUD]
        story.append(KeepTogether([subtitulo("Semáforo por tamaño de pedido (número de pedidos)"), chart_apiladas_v(cats_t, series)]))
        story += [Spacer(1, 4), Paragraph(
            f"Saludable: hasta {TARGET_COSTO_LOG}% · En alerta: más de {TARGET_COSTO_LOG}% · Crítico: más de {2 * TARGET_COSTO_LOG:g}% · Pérdida: el costo iguala o supera la facturación · "
            "Sin fact.: tiene costo pero no facturación. Sobrecosto = costo del pedido menos lo que costaría estando en target.", mp_ST["nota"])]

        # ---------------- 03 FLETERA Y MODALIDAD ----------------
        story += [Spacer(1, 12), CondPageBreak(340)] + seccion(3, "POR FLETERA Y MODALIDAD", f"Quién cuesta más en pedidos de {rng} y cómo se compara contra sus envíos de {mc + 1}+ cajas")
        pf = P["por_f"].head(12)
        pfp = pf[pf["pct_log"].notna()].sort_values("pct_log", ascending=False)
        izq = chart_barras_h([(r.K, r.pct_log) for r in pfp.itertuples()], width=262, label_w=92, fmt=lambda v: f"{v:.2f}%",
                             colors_list=[_color_pct(v) for v in pfp["pct_log"]]) if not pfp.empty else sin_datos("Sin facturación.")
        pcj = pf[pf["costo_caja"].notna()].sort_values("costo_caja", ascending=False)
        der = chart_barras_h([(r.K, r.costo_caja) for r in pcj.itertuples()], width=262, label_w=92, color=mp_C_GOLD,
                             fmt=lambda v: money(v, 2)) if not pcj.empty else sin_datos("Sin cajas.")
        story.append(lado_a_lado(izq, der, "% logístico en pedidos pequeños", "Costo por caja ($)"))
        story.append(Spacer(1, 10))
        rows = [[trunc(r.K, 24), num(r.n), num(r.cajas), money(r.fact), money(r.costo), pct(r.pct_log, 2), pct(r.pct_log_resto, 2),
                 _m(r.costo_caja, 2), _m(r.costo_ped, 2), pct(r.pct_fuera, 0)] for r in pf.itertuples()]
        story.append(KeepTogether([subtitulo("Detalle por fletera (top 12 por costo)"),
                                   tabla(["FLETERA", "PEDIDOS", "CAJAS", "FACTURACIÓN", "COSTO", "% LOG. PEQ.", f"% LOG. {mc + 1}+", "COSTO / CAJA", "COSTO / PEDIDO", "% FUERA"],
                                         rows, [96, 42, 36, 66, 60, 46, 46, 50, 54, 44], ["L"] + ["R"] * 9)]))
        story.append(Spacer(1, 10))
        pfm = P["por_forma"]
        rows = [[trunc(r.K, 28), num(r.n), money(r.fact), money(r.costo), pct(r.pct_log, 2), _m(r.costo_caja, 2), pct(r.pct_fuera, 0)] for r in pfm.itertuples()]
        story.append(KeepTogether([subtitulo("Por forma de envío"),
                                   tabla(["FORMA DE ENVÍO", "PEDIDOS", "FACTURACIÓN", "COSTO", "% LOG.", "COSTO / CAJA", "% FUERA DE TARGET"], rows,
                                         [130, 50, 80, 70, 60, 70, 80], ["L"] + ["R"] * 6)]))
        story += [Spacer(1, 4), Paragraph("% Log. peq. = pedidos pequeños; % Log. resto = los de más cajas de la misma fletera (si es mucho menor, el problema es el tamaño del pedido y no la fletera). "
                                          "% fuera = pedidos que superan el target.", mp_ST["nota"])]

        # ---------------- 04 TICKET Y PUNTO DE EQUILIBRIO ----------------
        story += [Spacer(1, 12), CondPageBreak(340)] + seccion(4, "TICKET Y PUNTO DE EQUILIBRIO", "¿Desde qué facturación por pedido el flete deja de pesar?")
        bd = P["bandas"]
        if not bd.empty:
            izq = chart_barras_h([(r.K, r.pct_log) for r in bd.itertuples()], width=262, label_w=104, fmt=lambda v: f"{v:.2f}%",
                                 colors_list=[_color_pct(v) for v in bd["pct_log"]])
            der = chart_barras_h([(r.K, r.costo_ped) for r in bd.itertuples()], width=262, label_w=104, color=C_ORANGE, fmt=lambda v: money(v, 2))
            story.append(lado_a_lado(izq, der, "% logístico por rango de facturación", "Costo promedio por pedido ($)"))
            story.append(Spacer(1, 10))
            rows = [[r.K, num(r.n), money(r.fact), money(r.costo), pct(r.pct_log, 2), _m(r.costo_ped, 2), pct(r.pct_fuera, 0)] for r in bd.itertuples()]
            story.append(KeepTogether([subtitulo("Pedidos pequeños por rango de facturación (cuartiles)"),
                                       tabla(["RANGO POR PEDIDO", "PEDIDOS", "FACTURACIÓN", "COSTO", "% LOG.", "COSTO / PEDIDO", "% FUERA"], rows,
                                             [140, 50, 80, 70, 56, 74, 70], ["L"] + ["R"] * 6)]))
            story.append(Spacer(1, 10))
        else:
            story += [sin_datos("No hay suficientes pedidos pequeños con facturación para armar rangos."), Spacer(1, 10)]
        ptm = P["por_tam"][P["por_tam"]["K"] <= mc]
        rows = [[_etq_tam(r.K, mc), _m(r.ticket), _m(r.costo_ped, 2), _m(r.ticket_min), _m(r.ticket - r.ticket_min) if pd.notna(r.ticket_min) else "-",
                 pct(r.pct_bajo_min, 0), _m(r.fact_caja)] for r in ptm.itertuples()]
        story.append(KeepTogether([subtitulo("Ticket mínimo para estar en target, por número de cajas"),
                                   tabla(["TAMAÑO", "TICKET PROM.", "COSTO / PEDIDO", "TICKET MÍNIMO", "BRECHA", "% PEDIDOS BAJO MÍNIMO", "FACT. POR CAJA"], rows,
                                         [76, 76, 76, 76, 76, 90, 70], ["L"] + ["R"] * 6)]))
        story += [Spacer(1, 4), Paragraph(
            f"Ticket mínimo = costo promedio por pedido / {TARGET_COSTO_LOG}%. Brecha = ticket promedio menos ticket mínimo (negativa = en promedio el pedido no alcanza el target). "
            "Sirve como referencia para definir pedido mínimo, cargos por envío chico o consolidación.", mp_ST["nota"])]

        # ---------------- 05 DÓNDE SE FUGA EL COSTO ----------------
        story += [Spacer(1, 12), CondPageBreak(340)] + seccion(5, "DÓNDE SE FUGA EL COSTO", "Clientes, destinos y pedidos con mayor sobrecosto frente al target")
        cl, de = P["clientes"], P["destinos"]
        izq = chart_barras_h([(r.K, r.exceso) for r in cl.head(8).itertuples()], width=262, label_w=110, color=mp_C_RED, fmt=money) if not cl.empty else sin_datos("Sin datos.")
        der = chart_barras_h([(r.K, r.exceso) for r in de.head(8).itertuples()], width=262, label_w=90, color=C_ORANGE, fmt=money) if not de.empty else sin_datos("Sin datos.")
        story.append(lado_a_lado(izq, der, "Sobrecosto por cliente (top 8)", "Sobrecosto por destino (top 8)"))
        story.append(Spacer(1, 10))
        if not cl.empty:
            rows = [[str(i + 1), trunc(r.K, 40), num(r.n), num(r.cajas), money(r.fact), money(r.costo), pct(r.pct_log, 2), money(r.exceso)]
                    for i, r in enumerate(cl.itertuples())]
            story.append(KeepTogether([subtitulo("Top 10 clientes por sobrecosto (pedidos pequeños)"),
                                       tabla(["#", "CLIENTE", "PEDIDOS", "CAJAS", "FACTURACIÓN", "COSTO", "% LOG.", "SOBRECOSTO"], rows,
                                             [22, 188, 46, 38, 70, 62, 52, 62], ["C", "L"] + ["R"] * 6)]))
            story.append(Spacer(1, 10))
        pe = P["peores"]
        if not pe.empty:
            rows = [[trunc(str(r["NÚMERO DE PEDIDO"]), 11), trunc(r["NOMBRE DEL CLIENTE"], 24), trunc(r["FLETERA"] or "S/A", 14), trunc(r["DESTINO"], 11),
                     num(r["_CAJ"]), money(r["FACTURACION"]), money(r["_COSTO"]), "S/F" if r["FACTURACION"] <= 0 else pct(r["_PCT"], 1), money(r["_EXC"])]
                    for _, r in pe.iterrows()]
            story.append(KeepTogether([subtitulo("15 pedidos pequeños con mayor sobrecosto"),
                                       tabla(["PEDIDO", "CLIENTE", "FLETERA", "DESTINO", "CJ.", "FACTURACIÓN", "COSTO", "% LOG.", "SOBRECOSTO"], rows,
                                             [52, 108, 72, 62, 28, 62, 50, 44, 62], ["L", "L", "L", "L", "R", "R", "R", "R", "R"])]))
            story += [Spacer(1, 3), Paragraph("S/F = pedido sin facturación registrada (todo su costo cuenta como sobrecosto).", mp_ST["nota"])]

        # ---------------- 06 TENDENCIA ----------------
        story += [Spacer(1, 12), CondPageBreak(340)] + seccion(6, "TENDENCIA DEL AÑO", f"% logístico mensual: pedidos de {rng} vs {mc + 1}+ cajas")
        tr = P["tendencia"]
        if len(tr) >= 2:
            story.append(chart_lineas([t[0] for t in tr], [(f"{rng.capitalize()}", [t[1] for t in tr], C_ORANGE),
                                                           (f"{mc + 1}+ cajas", [t[2] for t in tr], mp_C_TEAL)], target=TARGET_COSTO_LOG))
            story.append(Spacer(1, 8))
            rows = [[t[0], num(t[3]), pct(t[1], 2), pct(t[2], 2), (f"{t[1] - t[2]:+.2f} pp" if pd.notna(t[1]) and pd.notna(t[2]) else "-")] for t in tr]
            story.append(KeepTogether([subtitulo("Detalle mensual"),
                                       tabla(["MES", "PEDIDOS PEQUEÑOS", "% LOG. PEQUEÑOS", f"% LOG. {mc + 1}+ CAJAS", "DIFERENCIA"], rows,
                                             [90, 110, 110, 120, 110], ["L", "R", "R", "R", "R"])]))
        else:
            story.append(sin_datos("Se necesitan al menos dos meses con pedidos pequeños para mostrar la tendencia."))

        # ---------------- NOTAS ----------------
        notas = (f"Periodo según la columna MES de la matriz y el año seleccionado. Se toman las cajas de CANTIDAD DE CAJAS (o CAJAS si viene en 0) y se redondean al entero. "
                 f"Costo de flete = costo de la guía; costo de distribución = costos adicionales; se usa {base_txt}. Target logístico: {TARGET_COSTO_LOG}%.")
        if P["sin_costo"]:
            notas += f" Hay {num(P['sin_costo'])} pedido(s) pequeños sin costo registrado; si la guía aún no se captura, el % logístico se verá mejor de lo real."
        if P["extras"]["n"]:
            notas += " Los registros con concepto de recolecciones o maniobras se excluyen del análisis porque no son pedidos; su costo se reporta como costo extra."
        story += [Spacer(1, 10), Paragraph(notas, mp_ST["nota"])]

    buf = BytesIO()
    doc = BaseDocTemplate(buf, pagesize=letter, leftMargin=MARGIN, rightMargin=MARGIN, topMargin=48, bottomMargin=44,
                          title=f"Salud de pedidos pequeños - {periodo}", author="JYPESA Logística")
    doc.addPageTemplates(_hacer_paginas_simple("SALUD DE PEDIDOS PEQUEÑOS", periodo))

    class _Canvas(_NumberedCanvas):
        titulo_pie = f"Pedidos pequeños  -  generado {ahora}"
    doc.build(story, canvasmaker=_Canvas)
    buf.seek(0)
    return buf


# ============================================================
# 12. REPORTE INDEPENDIENTE: DETALLE PEDIDO POR PEDIDO (1 A 4 CAJAS)  ·  PDF u EXCEL
# ============================================================
# Factura por factura de los pedidos chicos: cliente, destino, transporte, valor de factura, costo logístico,
# % logístico, variación contra el target y CUÁNTO SE PERDIÓ.
#   PERDIMOS ($)  = lo que se pagó de logística por encima del target = max(costo - TARGET% x valor de factura, 0)
#   VS TARGET     = % logístico del pedido menos el target, en puntos porcentuales (pp)
#   ESTADO        = mismo semáforo del reporte de salud de pedidos pequeños (Sano / Alerta / Crítico / Pérdida / Sin facturación)
# Filtros: mes (o todo el año), año, modalidad (cobro regreso / cobro destino), número de cajas, destino, transporte,
# solo pedidos fuera de target y orden. Los registros de recolecciones / maniobras NO son pedidos (se excluyen, igual que en los demás reportes).
ORDENES_DETALLE = ["MAYOR PÉRDIDA PRIMERO", "FECHA DE ENVÍO", "FACTURA", "CLIENTE", "DESTINO"]
CAJAS_DETALLE_OPC = list(range(1, 11))
CAJAS_DETALLE_DEF = [1, 2, 3, 4]

PAGE_W_H, PAGE_H_H = PAGE_H, PAGE_W            # carta horizontal (792 x 612): el detalle trae muchas columnas
CONTENT_W_H = PAGE_W_H - 2 * MARGIN            # 720

TINTE_CLASE = {"SALUDABLE": "#DDF1E7", "EN ALERTA": "#FFF3CC", "CRÍTICO": "#FCE0CB", "PÉRDIDA": "#F8D4D4",
               "SIN FACTURACIÓN": "#E4E8EB", "SIN DATOS": "#E4E8EB"}
ETQ_CLASE_CORTA = {"SALUDABLE": "SANO", "EN ALERTA": "ALERTA", "CRÍTICO": "CRÍTICO", "PÉRDIDA": "PÉRDIDA",
                   "SIN FACTURACIÓN": "S/FACT.", "SIN DATOS": "S/DATOS"}


def mp_col_factura(df):
    """Columna que identifica la factura: una columna FACTURA si la matriz la trae; si no, NÚMERO DE PEDIDO (igual que 'facturas' en la sección 02)."""
    exactos = {"FACTURA", "NUMERO DE FACTURA", "NUMERO FACTURA", "NO FACTURA", "NO. FACTURA", "NUM FACTURA", "NUM. FACTURA",
               "FOLIO FACTURA", "FOLIO DE FACTURA", "N FACTURA"}
    for c in df.columns:
        if _sin_acentos(c).strip().upper() in exactos:
            return c
    return "NÚMERO DE PEDIDO"


def _etq_cajas(n):
    n = int(n)
    return f"{n} CAJA" + ("S" if n > 1 else "")


def _txt_cajas(cajas):
    c = sorted(int(x) for x in cajas)
    if len(c) == 1:
        return _etq_cajas(c[0])
    if c == list(range(c[0], c[-1] + 1)):
        return f"{c[0]} A {c[-1]} CAJAS"
    return ", ".join(str(x) for x in c[:-1]) + f" Y {c[-1]} CAJAS"


def _opciones_detalle(df_base, anio):
    """Destinos y transportes disponibles del año (para los filtros de la página)."""
    d, _ = _separar_extras(df_base)
    f_env = d["FECHA DE ENVÍO"]
    d = d[(f_env.isna() | (f_env.dt.year == anio)) & (d["MES"] != "")]
    dest = sorted(set(d["DESTINO"].str.upper().replace("", "SIN DESTINO")))
    trans = sorted(set(d["TRANSPORTE"].str.upper().replace("", "SIN TRANSPORTE")))
    return dest, trans


def _texto_filtros_detalle(periodo, modalidad, cajas, destinos, transportes, solo_fuera, base_txt):
    def _lst(x, vacio):
        if not x:
            return vacio
        return ", ".join(x) if len(x) <= 3 else f"{len(x)} seleccionados"
    partes = ["Todas las modalidades" if modalidad == "TODAS" else modalidad.title(), periodo.title(), _txt_cajas(cajas).lower(),
              f"Destino: {_lst(destinos, 'todos')}", f"Transporte: {_lst(transportes, 'todos')}"]
    if solo_fuera:
        partes.append("solo pedidos fuera de target")
    partes.append(f"costo = {base_txt}")
    return "  ·  ".join(partes)


def _ordenar_detalle(d, orden):
    if orden == "FECHA DE ENVÍO":
        return d.sort_values(["FECHA DE ENVÍO", "_FAC"], na_position="last")
    if orden == "FACTURA":
        return d.sort_values("_FAC")
    if orden == "CLIENTE":
        return d.sort_values(["_CLI", "_EXC"], ascending=[True, False])
    if orden == "DESTINO":
        return d.sort_values(["_DE", "_EXC"], ascending=[True, False])
    return d.sort_values(["_EXC", "_PCT"], ascending=[False, False], na_position="last")


def calc_detalle_pequenos(df_base, anio, mes, modalidad="TODAS", incluir_adic=True, cajas=None, destinos=None,
                          transportes=None, solo_fuera=False, orden=ORDENES_DETALLE[0]):
    """Pedido por pedido (1 a N cajas) con sus indicadores + resúmenes por cajas / transporte / destino / cliente / modalidad / mes.
    Periodo y modalidad igual que los demás reportes (columna MES + año; modalidad por FORMA DE ENVIO). mes = 0 -> todo el año."""
    cajas = sorted({int(c) for c in (cajas or CAJAS_DETALLE_DEF)})
    df_base, df_ext = _separar_extras(df_base)                     # recolecciones / maniobras no son pedidos
    E = calc_extras(df_ext, anio, mes, modalidad)
    anual, sin_cajas = _preparar_pequenos(df_base, anio, modalidad, incluir_adic, max(cajas))
    d = (anual[anual["MES"] == MESES[mes - 1]] if mes else anual).copy()
    d["_DE"] = d["DESTINO"].str.upper().replace("", "SIN DESTINO")
    d["_TR"] = d["TRANSPORTE"].str.upper().replace("", "SIN TRANSPORTE")
    d = d[d["_CAJ"].isin(cajas)]
    if destinos:
        d = d[d["_DE"].isin([str(x).upper() for x in destinos])]
    if transportes:
        d = d[d["_TR"].isin([str(x).upper() for x in transportes])]
    if solo_fuera:
        d = d[d["_CLASE"].isin(CLASES_FUERA_TARGET)]
    if d.empty:
        return dict(vacio=True, extras=E, motivo=f"No hay pedidos de {_txt_cajas(cajas).lower()} con esos filtros (periodo, modalidad, destino, transporte).")

    d = d.copy()
    col_f = mp_col_factura(d)
    d["_FAC"] = d[col_f].fillna("").astype(str).str.strip()
    d.loc[d["_FAC"].str.lower() == "nan", "_FAC"] = ""
    d["_MOD"] = d["FORMA DE ENVIO"].str.upper().map(
        lambda x: "COBRO REGRESO" if "REGRESO" in x else ("COBRO DESTINO" if "DESTINO" in x else (x or "SIN FORMA")))
    d["_CLI"] = d["NOMBRE DEL CLIENTE"].replace("", "SIN CLIENTE")
    d["_FLE"] = d["FLETERA"].replace("", "SIN ASIGNAR")
    d = _ordenar_detalle(d, orden)

    det = pd.DataFrame({
        "FACTURA": d["_FAC"].to_numpy(), "FECHA": d["FECHA DE ENVÍO"].to_numpy(), "MES": d["MES"].to_numpy(),
        "ANIO": d["FECHA DE ENVÍO"].dt.year.fillna(anio).astype(int).to_numpy(), "CLIENTE": d["_CLI"].to_numpy(),
        "DESTINO": d["_DE"].to_numpy(), "CAJAS": d["_CAJ"].astype(int).to_numpy(), "MODALIDAD": d["_MOD"].to_numpy(),
        "TRANSPORTE": d["_TR"].to_numpy(), "FLETERA": d["_FLE"].to_numpy(), "GUIA": d["COSTO DE LA GUÍA"].to_numpy(),
        "ADIC": d["COSTOS ADICIONALES"].to_numpy(), "COSTO": d["_COSTO"].to_numpy(), "FACT": d["FACTURACION"].to_numpy(),
        "PCT": d["_PCT"].to_numpy(), "VS": (d["_PCT"] - TARGET_COSTO_LOG).to_numpy(), "PERDIMOS": d["_EXC"].to_numpy(),
        "CLASE": d["_CLASE"].to_numpy()})

    def _grupo(col):
        return _df_bloques(d, col).sort_values(["exceso", "costo"], ascending=False).reset_index(drop=True)
    por_mes = []
    if mes == 0:
        for m in MESES:
            g = d[d["MES"] == m]
            if not g.empty:
                b = _bloque(g)
                b["K"] = m.title()
                por_mes.append(b)

    return dict(vacio=False, d=d, det=det, T=_bloque(d), cajas=cajas, sin_cajas=sin_cajas, extras=E, col_factura=col_f,
                por_cajas=_df_bloques(d.sort_values("_CAJ"), "_CAJ"), por_transp=_grupo("_TR"), por_dest=_grupo("_DE"),
                por_cli=_grupo("_CLI"), por_mod=_grupo("_MOD"), por_mes=pd.DataFrame(por_mes),
                peores=det[det["PERDIMOS"] > 0].sort_values("PERDIMOS", ascending=False).head(15),
                exceso_total=float(d["_EXC"].sum()), n_fuera=int(d["_CLASE"].isin(CLASES_FUERA_TARGET).sum()),
                n_perdida=int((d["_CLASE"] == "PÉRDIDA").sum()),
                n_sin_fact=int(d["_CLASE"].isin(["SIN FACTURACIÓN", "SIN DATOS"]).sum()),
                sin_costo=int((d["_COSTO"] <= 0).sum()))


# ---------------- PDF horizontal ----------------
def _kpi_row_h(items, ncols=4, gap=8):
    """Igual que kpi_row pero a todo el ancho de la hoja horizontal."""
    cw = (CONTENT_W_H - gap * (ncols - 1)) / ncols
    out = []
    for i in range(0, len(items), ncols):
        chunk = items[i:i + ncols]
        row, widths, style = [], [], []
        for j, (lab, val, col) in enumerate(chunk):
            c = j * 2
            v_style = mp_ST["kpi_v"] if len(str(val)) <= 17 else mp_ST["kpi_v_s"]
            row.append([Paragraph(esc(lab).upper(), mp_ST["kpi_l"]), Spacer(1, 3),
                        Paragraph(f'<font color="{col}">{esc(val)}</font>', v_style)])
            widths.append(cw)
            style += [("BACKGROUND", (c, 0), (c, 0), mp_C_LIGHT), ("LINEABOVE", (c, 0), (c, 0), 3, colors.HexColor(col)),
                      ("BOX", (c, 0), (c, 0), 0.5, mp_C_BORDER)]
            if j < len(chunk) - 1:
                row.append("")
                widths.append(gap)
        t = Table([row], colWidths=widths, hAlign="LEFT")
        t.setStyle(TableStyle(style + [("VALIGN", (0, 0), (-1, -1), "TOP"),
                                       ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                                       ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        out += [t, Spacer(1, 8)]
    return out


def _seccion_h(num_, titulo, subtitulo_=""):
    """Banda de título a todo el ancho de la hoja horizontal (num_=None -> sin número)."""
    cab = f"{num_:02d}&nbsp;&nbsp;{esc(titulo)}" if num_ else esc(titulo)
    cont = [Paragraph(cab, mp_ST["sec_t"])]
    if subtitulo_:
        cont.append(Paragraph(esc(subtitulo_), mp_ST["sec_s"]))
    t = Table([[cont]], colWidths=[CONTENT_W_H])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), mp_C_SLATE), ("LINEBEFORE", (0, 0), (0, 0), 5, mp_C_TEAL),
        ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 12),
        ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 9)]))
    return [t, Spacer(1, 10)]


class _NumberedCanvasH(_NumberedCanvas):
    """Numeración de páginas para la hoja horizontal."""

    def save(self):
        n = len(self._saved)
        for s_ in self._saved:
            self.__dict__.update(s_)
            self.setStrokeColor(mp_C_BORDER)
            self.setLineWidth(0.5)
            self.line(MARGIN, 34, PAGE_W_H - MARGIN, 34)
            self.setFont("Helvetica", 7)
            self.setFillColor(mp_C_GRAY)
            self.drawString(MARGIN, 23, f"JYPESA | Logística  -  {self.titulo_pie}  -  Generado e impreso con Nexion Smart Logistic")
            self.drawRightString(PAGE_W_H - MARGIN, 23, f"Página {self._pageNumber} de {n}")
            rl_canvas.Canvas.showPage(self)
        rl_canvas.Canvas.save(self)


def _hacer_paginas_horizontal(titulo_hdr, periodo_txt):
    def normal(c, doc):
        c.saveState()
        c.setFillColor(mp_C_NAVY)
        c.rect(0, PAGE_H_H - 30, PAGE_W_H, 30, stroke=0, fill=1)
        c.setFillColor(mp_C_TEAL)
        c.rect(0, PAGE_H_H - 30, PAGE_W_H, 2.5, stroke=0, fill=1)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 8.5)
        c.drawString(MARGIN, PAGE_H_H - 19, titulo_hdr)
        c.setFillColor(mp_C_GOLD)
        c.drawRightString(PAGE_W_H - MARGIN, PAGE_H_H - 19, periodo_txt)
        c.restoreState()
    f = Frame(MARGIN, 44, CONTENT_W_H, PAGE_H_H - 44 - 48, id="fh", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    return [PageTemplate(id="horizontal", frames=[f], onPage=normal)]


W_RES = [124, 52, 46, 88, 76, 52, 58, 68, 56, 100]          # suma 720
H_RES = ["", "PEDIDOS", "CAJAS", "VALOR FACTURA", "COSTO LOG.", "% LOG.", f"VS {TARGET_COSTO_LOG}%", "COSTO / CAJA", "% FUERA", "PERDIMOS ($)"]
W_DET = [56, 40, 118, 64, 22, 48, 82, 58, 50, 38, 42, 52, 50]   # suma 720
H_DET = ["FACTURA", "FECHA", "CLIENTE", "DESTINO", "CJ.", "COBRO", "TRANSPORTE", "VALOR FACT.", "COSTO", "% LOG.", f"VS {TARGET_COSTO_LOG}%",
         "PERDIMOS", "ESTADO"]
A_DET = ["L", "C", "L", "L", "R", "L", "L", "R", "R", "R", "R", "R", "C"]


def _fila_res(lab, b):
    return [lab, num(b["n"]), num(b["cajas"]), money(b["fact"]), money(b["costo"]), pct(b["pct_log"], 2), _vs_target(b["pct_log"]),
            _m(b["costo_caja"], 2), pct(b["pct_fuera"], 0), money(b["exceso"])]


def _tabla_res(titulo_col, filas, total=None):
    return tabla([titulo_col] + H_RES[1:], filas, W_RES, ["L"] + ["R"] * 9, total_row=total)


def _tabla_detalle(g, total_row=None):
    """Tabla pedido por pedido; la columna ESTADO lleva color de semáforo."""
    filas = []
    for r in g.itertuples():
        filas.append([trunc(r.FACTURA, 10), r.FECHA.strftime("%d/%m/%y") if pd.notna(r.FECHA) else "S/F", trunc(r.CLIENTE, 25), trunc(r.DESTINO, 12),
                      num(r.CAJAS), {"COBRO REGRESO": "REGRESO", "COBRO DESTINO": "DESTINO"}.get(r.MODALIDAD, trunc(r.MODALIDAD, 8)),
                      trunc(r.TRANSPORTE, 16), money(r.FACT, 2), money(r.COSTO, 2), pct(r.PCT, 2), _vs_target(r.PCT), money(r.PERDIMOS, 2),
                      ETQ_CLASE_CORTA.get(r.CLASE, r.CLASE)])
    t = tabla(H_DET, filas, W_DET, A_DET, font=6.8, total_row=total_row)
    t.setStyle(TableStyle([("BACKGROUND", (12, i + 1), (12, i + 1), colors.HexColor(TINTE_CLASE.get(r.CLASE, "#E4E8EB")))
                           for i, r in enumerate(g.itertuples())]))
    return t


def generar_detalle_pequenos_pdf(df_base, anio, mes, modalidad="TODAS", incluir_adic=True, cajas=None, destinos=None,
                                 transportes=None, solo_fuera=False, orden=ORDENES_DETALLE[0]):
    """PDF horizontal pedido por pedido (1 a N cajas). mes=0 -> todo el año. Devuelve BytesIO."""
    cajas = sorted({int(c) for c in (cajas or CAJAS_DETALLE_DEF)})
    periodo = f"{MESES[mes - 1]} {anio}" if mes else f"AÑO {anio} (ACUMULADO)"
    ahora = datetime.now().strftime("%d/%m/%Y %H:%M")
    base_txt = "guía + costos adicionales" if incluir_adic else "solo costo de guía"
    rng = _txt_cajas(cajas)
    P = calc_detalle_pequenos(df_base, anio, mes, modalidad, incluir_adic, cajas, destinos, transportes, solo_fuera, orden)
    story = _seccion_h(None, f"DETALLE DE PEDIDOS PEQUEÑOS  ·  {rng}",
                       _texto_filtros_detalle(periodo, modalidad, cajas, destinos, transportes, solo_fuera, base_txt))
    if P["vacio"]:
        story.append(sin_datos(P["motivo"]))
    else:
        T, det = P["T"], P["det"]
        col_l = mp_HEX["green"] if (pd.notna(T["pct_log"]) and T["pct_log"] <= TARGET_COSTO_LOG) else mp_HEX["red"]
        story.append(Paragraph("RESUMEN EJECUTIVO", ParagraphStyle("re", fontName="Helvetica-Bold", fontSize=11, textColor=mp_C_SLATE, spaceAfter=8)))
        story += _kpi_row_h([
            ("Pedidos (facturas)", num(T["n"]), mp_HEX["slate"]),
            ("Cajas", num(T["cajas"]), mp_HEX["teal"]),
            ("Valor facturado", money(T["fact"], 2), mp_HEX["green"]),
            (f"Costo logístico ({base_txt})", money(T["costo"], 2), mp_HEX["orange"]),
        ])
        story += _kpi_row_h([
            (f"% logístico (target {TARGET_COSTO_LOG}%)", pct(T["pct_log"], 2), col_l),
            ("Perdimos vs target", money(P["exceso_total"], 2), mp_HEX["red"]),
            ("Pedidos fuera de target", f"{num(P['n_fuera'])}  ·  {pct(P['n_fuera'] / T['n'] * 100, 0)}", mp_HEX["gold"]),
            ("Costo ≥ factura / sin facturación", f"{num(P['n_perdida'])}  /  {num(P['n_sin_fact'])}", mp_HEX["purple"]),
        ])

        hall = [f"Se analizaron <b>{num(T['n'])}</b> pedidos de {rng.lower()} ({num(T['cajas'])} cajas) con <b>{money(T['fact'])}</b> de valor de factura; "
                f"el costo logístico fue <b>{money(T['costo'])}</b> ({base_txt}), es decir <b>{pct(T['pct_log'], 2)}</b> frente al target de {TARGET_COSTO_LOG}%.",
                f"Contra ese target <b>perdimos {money(P['exceso_total'])}</b>: <b>{num(P['n_fuera'])}</b> de {num(T['n'])} pedidos ({pct(P['n_fuera'] / T['n'] * 100, 0)}) lo superan; "
                f"<b>{num(P['n_perdida'])}</b> costaron igual o más de lo que se facturó y <b>{num(P['n_sin_fact'])}</b> no tienen facturación registrada."]
        for nombre, df_g, etq in (("número de cajas", P["por_cajas"], lambda k: _etq_cajas(k)), ("transporte", P["por_transp"], str),
                                  ("destino", P["por_dest"], str), ("cliente", P["por_cli"], str)):
            top = df_g.sort_values("exceso", ascending=False).iloc[0]
            if top["exceso"] > 0:
                hall.append(f"Por {nombre}, donde más perdimos es <b>{esc(etq(top['K']))}</b>: <b>{money(top['exceso'])}</b> "
                            f"({pct(top['exceso'] / P['exceso_total'] * 100, 0)} del total) con {pct(top['pct_log'], 2)} logístico en {num(top['n'])} pedido(s).")
        if P["extras"]["n"]:
            hall.append(f"No se incluyeron <b>{num(P['extras']['n'])}</b> registros de recolecciones y maniobras ({money(P['extras']['costo'])}): no son pedidos, se reportan como costos extras.")
        story.append(Paragraph("HALLAZGOS CLAVE", ParagraphStyle("hk", fontName="Helvetica-Bold", fontSize=11, textColor=mp_C_SLATE, spaceBefore=4, spaceAfter=6)))
        for h in hall:
            story.append(Paragraph(h, mp_ST["bullet"], bulletText="•"))

        # ---------------- 01 POR NÚMERO DE CAJAS ----------------
        story += [Spacer(1, 6), CondPageBreak(330)] + _seccion_h(1, "RESUMEN POR NÚMERO DE CAJAS", "Cuánto factura, cuánto cuesta y cuánto perdimos según el tamaño del pedido")
        pc = P["por_cajas"]
        story.append(_tabla_res("NÚMERO DE CAJAS", [_fila_res(_etq_cajas(r["K"]), r) for _, r in pc.iterrows()], total=_fila_res("TOTAL", T)))
        story.append(Spacer(1, 10))
        cats = [f"{int(k)} cj" for k in pc["K"]]
        izq = chart_barras_v(cats, [0 if pd.isna(v) else v for v in pc["pct_log"]], width=352, height=150, fmt=lambda v: f"{v:.1f}%",
                             colors_list=[_color_pct(v) for v in pc["pct_log"]])
        der = chart_barras_v(cats, [float(v) for v in pc["exceso"]], width=352, height=150, color=mp_C_RED, fmt=lambda v: money(v, 0))
        story.append(lado_a_lado(izq, der, f"% logístico por número de cajas (verde = dentro de {TARGET_COSTO_LOG}%)", "Cuánto perdimos por número de cajas ($)", w=(352, 352)))

        # ---------------- 02 POR TRANSPORTE ----------------
        story += [Spacer(1, 12), CondPageBreak(330)] + _seccion_h(2, "RESUMEN POR TRANSPORTE", "Qué transporte se llevó los pedidos y cuánto nos costó de más")
        pt = P["por_transp"]
        story.append(_tabla_res("TRANSPORTE", [_fila_res(trunc(r["K"], 32), r) for _, r in pt.iterrows()], total=_fila_res("TOTAL", T)))
        story.append(Spacer(1, 10))
        ptc = pt[pt["pct_log"].notna()].sort_values("pct_log", ascending=False).head(10)
        izq = chart_barras_h([(r["K"], r["pct_log"]) for _, r in ptc.iterrows()], width=352, label_w=110, fmt=lambda v: f"{v:.2f}%",
                             colors_list=[_color_pct(v) for v in ptc["pct_log"]]) if not ptc.empty else sin_datos("Sin facturación.")
        der = chart_barras_h([(r["K"], r["exceso"]) for _, r in pt.head(10).iterrows()], width=352, label_w=110, color=mp_C_RED, fmt=money)
        story.append(lado_a_lado(izq, der, "% logístico por transporte", "Cuánto perdimos por transporte ($)", w=(352, 352)))

        # ---------------- 03 POR DESTINO ----------------
        story += [Spacer(1, 12), CondPageBreak(330)] + _seccion_h(3, "RESUMEN POR DESTINO", "Destinos con mayor pérdida frente al target (top 20)")
        pdst = P["por_dest"]
        story.append(_tabla_res("DESTINO", [_fila_res(trunc(r["K"], 32), r) for _, r in pdst.head(20).iterrows()], total=_fila_res("TOTAL", T)))
        if len(pdst) > 20:
            story += [Spacer(1, 3), Paragraph(f"Se muestran los 20 destinos con mayor pérdida de {len(pdst)}; el total incluye todos. El Excel trae la lista completa.", mp_ST["nota"])]

        # ---------------- 04 POR MODALIDAD / MES ----------------
        if len(P["por_mod"]) > 1 or len(P["por_mes"]) > 0:
            story += [Spacer(1, 12), CondPageBreak(250)] + _seccion_h(4, "POR MODALIDAD Y POR MES", "Cobro regreso vs cobro destino y evolución mensual")
            if len(P["por_mod"]) > 1:
                story.append(KeepTogether([subtitulo("Por modalidad (forma de envío)"),
                                           _tabla_res("MODALIDAD", [_fila_res(trunc(r["K"], 32), r) for _, r in P["por_mod"].iterrows()], total=_fila_res("TOTAL", T))]))
                story.append(Spacer(1, 10))
            if len(P["por_mes"]) > 0:
                story.append(KeepTogether([subtitulo("Por mes"),
                                           _tabla_res("MES", [_fila_res(r["K"], r) for _, r in P["por_mes"].iterrows()], total=_fila_res("TOTAL", T))]))

        # ---------------- 05 CLIENTES Y PEDIDOS CON MAYOR PÉRDIDA ----------------
        story += [Spacer(1, 12), CondPageBreak(440)] + _seccion_h(5, "DÓNDE PERDIMOS MÁS", "Clientes y pedidos con mayor sobrecosto frente al target")
        pcl = P["por_cli"][P["por_cli"]["exceso"] > 0].head(15)
        if not pcl.empty:
            story.append(KeepTogether([subtitulo("Top 15 clientes por pérdida"),
                                       _tabla_res("CLIENTE", [_fila_res(trunc(r["K"], 32), r) for _, r in pcl.iterrows()])]))
            story.append(Spacer(1, 10))
        pe = P["peores"]
        if not pe.empty:
            story.append(KeepTogether([subtitulo("Los 15 pedidos donde más perdimos"), _tabla_detalle(pe)]))
        else:
            story.append(sin_datos("Ningún pedido superó el target en este periodo."))

        # ---------------- 06 DETALLE PEDIDO POR PEDIDO ----------------
        story += [PageBreak()] + _seccion_h(6, "DETALLE PEDIDO POR PEDIDO", f"Agrupado por número de cajas  ·  orden: {orden.lower()}  ·  perdimos = costo menos {TARGET_COSTO_LOG}% del valor de factura")
        d_all = P["d"]
        for n in sorted(det["CAJAS"].unique()):
            g = det[det["CAJAS"] == n]
            tot = _bloque(d_all[d_all["_CAJ"] == n])
            story += [CondPageBreak(110),
                      subtitulo(f"{_etq_cajas(n)}  ·  {num(tot['n'])} pedidos  ·  % logístico {pct(tot['pct_log'], 2)}  ·  perdimos {money(tot['exceso'], 2)}")]
            story.append(_tabla_detalle(g, total_row=["TOTAL", "", f"{num(tot['n'])} pedidos", "", num(tot["cajas"]), "", "", money(tot["fact"], 2),
                                                      money(tot["costo"], 2), pct(tot["pct_log"], 2), _vs_target(tot["pct_log"]), money(tot["exceso"], 2), ""]))
            story.append(Spacer(1, 12))

        notas = (f"FACTURA = columna '{P['col_factura']}' de la matriz. COSTO = {base_txt}. % LOG. = costo / valor de factura. VS {TARGET_COSTO_LOG}% = diferencia en puntos porcentuales contra el target. "
                 f"PERDIMOS = lo pagado por encima del target (costo menos {TARGET_COSTO_LOG}% del valor de factura; si no hay factura, todo el costo). "
                 f"ESTADO: Sano hasta {TARGET_COSTO_LOG}% · Alerta más de {TARGET_COSTO_LOG}% · Crítico más de {2 * TARGET_COSTO_LOG:g}% · Pérdida si el costo iguala o supera la factura · S/Fact. sin factura registrada. "
                 "COBRO: REGRESO / DESTINO según la forma de envío. Periodo según la columna MES de la matriz y el año seleccionado. Las cajas salen de CANTIDAD DE CAJAS (o CAJAS si viene en 0). "
                 "Los registros con concepto de recolecciones o maniobras no son pedidos y se excluyen.")
        if P["sin_costo"]:
            notas += f" Hay {num(P['sin_costo'])} pedido(s) sin costo registrado; si la guía aún no se captura, su % logístico se ve mejor de lo real."
        if P["sin_cajas"]:
            notas += f" {num(P['sin_cajas'])} registro(s) sin cajas quedan fuera del análisis."
        story += [Spacer(1, 4), Paragraph(notas, mp_ST["nota"])]

    buf = BytesIO()
    doc = BaseDocTemplate(buf, pagesize=(PAGE_W_H, PAGE_H_H), leftMargin=MARGIN, rightMargin=MARGIN, topMargin=48, bottomMargin=44,
                          title=f"Detalle de pedidos pequeños - {periodo}", author="JYPESA Logística")
    doc.addPageTemplates(_hacer_paginas_horizontal("DETALLE DE PEDIDOS PEQUEÑOS (PEDIDO POR PEDIDO)", periodo))

    class _Canvas(_NumberedCanvasH):
        titulo_pie = f"Detalle pedidos pequeños  -  generado {ahora}"
    doc.build(story, canvasmaker=_Canvas)
    buf.seek(0)
    return buf


# ---------------- Excel ----------------
def generar_detalle_pequenos_excel(df_base, anio, mes, modalidad="TODAS", incluir_adic=True, cajas=None, destinos=None,
                                   transportes=None, solo_fuera=False, orden=ORDENES_DETALLE[0]):
    """Excel con hoja 'Detalle' (filtros automáticos; % logístico, vs target, perdimos y estado con fórmulas, totales que siguen al filtro)
    y hojas de resumen. El target de la celda B3 de 'Detalle' se puede cambiar y todo se recalcula. Devuelve BytesIO."""
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import CellIsRule, FormulaRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as e:   # pragma: no cover
        raise RuntimeError("Falta la librería 'openpyxl' para generar Excel: agrégala a requirements.txt (openpyxl).") from e

    cajas = sorted({int(c) for c in (cajas or CAJAS_DETALLE_DEF)})
    periodo = f"{MESES[mes - 1]} {anio}" if mes else f"AÑO {anio} (ACUMULADO)"
    base_txt = "guía + costos adicionales" if incluir_adic else "solo costo de guía"
    P = calc_detalle_pequenos(df_base, anio, mes, modalidad, incluir_adic, cajas, destinos, transportes, solo_fuera, orden)
    filtros = _texto_filtros_detalle(periodo, modalidad, cajas, destinos, transportes, solo_fuera, base_txt)

    FN = "Arial"
    f_n, f_b = Font(name=FN, size=10), Font(name=FN, size=10, bold=True)
    f_h, f_t = Font(name=FN, size=10, bold=True, color="FFFFFF"), Font(name=FN, size=14, bold=True, color="2B343B")
    f_s, f_in = Font(name=FN, size=9, italic=True, color="6B7A86"), Font(name=FN, size=10, bold=True, color="0000FF")
    fill = lambda hx: PatternFill("solid", start_color=hx, end_color=hx)
    lado = Side(style="thin", color="D5DDE2")
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)
    MONEY, PCT_F, PP, FECHA = '"$"#,##0.00', "0.00%", '+0.00" pp";-0.00" pp";0.00" pp"', "dd/mm/yyyy"

    def _v(x):
        return None if x is None or (isinstance(x, float) and np.isnan(x)) else x

    wb = Workbook()
    wr = wb.active
    wr.title = "Resumen"
    wr["A1"], wr["A1"].font = f"DETALLE DE PEDIDOS PEQUEÑOS  ·  {_txt_cajas(cajas)}", f_t
    wr["A2"], wr["A2"].font = filtros, f_s
    wr["A3"], wr["A3"].font = f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}  ·  JYPESA Logística  ·  Nexion Smart Logistic", f_s
    if P["vacio"]:
        wr["A5"], wr["A5"].font = P["motivo"], f_b
        wr.column_dimensions["A"].width = 110
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    T, det = P["T"], P["det"]
    N = len(det)

    # ======== hoja DETALLE (viva) ========
    wd = wb.create_sheet("Detalle")
    DATA0, HDR, TOT = 6, 5, 4
    LAST = DATA0 + N - 1
    wd["A1"], wd["A1"].font = f"DETALLE PEDIDO POR PEDIDO  ·  {_txt_cajas(cajas)}  ·  {periodo.title()}", f_t
    wd["A2"], wd["A2"].font = filtros, f_s
    wd["A3"], wd["A3"].font = "TARGET % LOG. →", f_b
    wd["B3"] = TARGET_COSTO_LOG / 100
    wd["B3"].number_format, wd["B3"].font, wd["B3"].fill, wd["B3"].border = "0.0%", f_in, fill("FFF2CC"), borde
    wd["C3"], wd["C3"].font = "Cambia este valor y se recalculan % vs target, PERDIMOS y ESTADO. Los totales de la fila 4 siguen a los filtros.", f_s

    cols = [("FACTURA", 16), ("FECHA ENVÍO", 12), ("MES", 12), ("AÑO", 7), ("CLIENTE", 40), ("DESTINO", 20), ("CAJAS", 8),
            ("COBRO (MODALIDAD)", 18), ("TRANSPORTE", 20), ("FLETERA", 20), ("COSTO GUÍA", 14), ("COSTOS ADIC.", 14),
            ("COSTO LOGÍSTICO", 16), ("VALOR FACTURA", 16), ("% LOGÍSTICO", 12), ("VS TARGET (pp)", 14), ("PERDIMOS ($)", 15), ("ESTADO", 17)]
    for j, (h, w) in enumerate(cols, start=1):
        c = wd.cell(row=HDR, column=j, value=h)
        c.font, c.fill, c.border = f_h, fill("2B343B"), borde
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        wd.column_dimensions[get_column_letter(j)].width = w
    wd.row_dimensions[HDR].height = 30

    for i, r in enumerate(det.itertuples(index=False)):
        x = DATA0 + i
        costo_f = f"=K{x}+L{x}" if incluir_adic else f"=K{x}"
        vals = [r.FACTURA, r.FECHA.to_pydatetime() if pd.notna(r.FECHA) else None, r.MES.title(), int(r.ANIO), r.CLIENTE, r.DESTINO, int(r.CAJAS),
                r.MODALIDAD, r.TRANSPORTE, r.FLETERA, float(r.GUIA), float(r.ADIC), costo_f, float(r.FACT),
                f'=IF(N{x}>0,M{x}/N{x},"")', f'=IF(N{x}>0,(M{x}/N{x}-$B$3)*100,"")', f"=MAX(M{x}-$B$3*N{x},0)",
                f'=IF(AND(N{x}<=0,M{x}<=0),"SIN DATOS",IF(N{x}<=0,"SIN FACTURACIÓN",IF(M{x}/N{x}>=1,"PÉRDIDA",'
                f'IF(M{x}/N{x}>2*$B$3,"CRÍTICO",IF(M{x}/N{x}>$B$3,"EN ALERTA","SALUDABLE")))))']
        fmts = [None, FECHA, None, "0", None, None, "0", None, None, None, MONEY, MONEY, MONEY, MONEY, PCT_F, PP, MONEY, None]
        for j, (v, fm) in enumerate(zip(vals, fmts), start=1):
            c = wd.cell(row=x, column=j, value=v)
            c.font, c.border = f_n, borde
            if fm:
                c.number_format = fm
            if j in (2, 4, 7, 18):
                c.alignment = Alignment(horizontal="center")

    # fila de totales (SUBTOTAL: solo suma lo que está visible con el filtro)
    wd.cell(row=TOT, column=1, value=f'="TOTAL FILTRADO ("&SUBTOTAL(103,A{DATA0}:A{LAST})&" pedidos)"')
    for j, letra in ((7, "G"), (11, "K"), (12, "L"), (13, "M"), (14, "N"), (17, "Q")):
        wd.cell(row=TOT, column=j, value=f"=SUBTOTAL(109,{letra}{DATA0}:{letra}{LAST})")
    wd.cell(row=TOT, column=15, value=f'=IF(N{TOT}>0,M{TOT}/N{TOT},"")')
    wd.cell(row=TOT, column=16, value=f'=IF(N{TOT}>0,(M{TOT}/N{TOT}-$B$3)*100,"")')
    for j in range(1, len(cols) + 1):
        c = wd.cell(row=TOT, column=j)
        c.font, c.fill, c.border = f_b, fill("E3EAEE"), borde
        c.number_format = {7: "0", 11: MONEY, 12: MONEY, 13: MONEY, 14: MONEY, 15: PCT_F, 16: PP, 17: MONEY}.get(j, "General")

    wd.freeze_panes = f"B{DATA0}"
    wd.auto_filter.ref = f"A{HDR}:{get_column_letter(len(cols))}{LAST}"
    for txt, hx in TINTE_CLASE.items():
        wd.conditional_formatting.add(f"R{DATA0}:R{LAST}", CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=fill(hx.lstrip("#"))))
    wd.conditional_formatting.add(f"O{DATA0}:O{LAST}", FormulaRule(formula=[f"AND(ISNUMBER($O{DATA0}),$O{DATA0}>$B$3)"], font=Font(name=FN, color="D64545", bold=True)))
    wd.conditional_formatting.add(f"Q{DATA0}:Q{LAST}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FN, color="D64545", bold=True)))
    wd.page_setup.orientation = "landscape"
    wd.page_setup.fitToWidth, wd.page_setup.fitToHeight = 1, 0
    wd.sheet_properties.pageSetUpPr.fitToPage = True
    wd.print_title_rows = f"{HDR}:{HDR}"

    # ======== hojas de RESUMEN (valores al momento de generar, con el target indicado) ========
    heads = ["", "PEDIDOS", "CAJAS", "VALOR FACTURA", "COSTO LOGÍSTICO", "% LOGÍSTICO", "VS TARGET (pp)", "COSTO / CAJA", "% FUERA DE TARGET", "PERDIMOS ($)"]
    fmt_g = [None, "#,##0", "#,##0", MONEY, MONEY, PCT_F, PP, MONEY, "0%", MONEY]

    def _fila_x(b, etq):
        pl = _v(b["pct_log"])
        return [etq, int(b["n"]), float(b["cajas"]), float(b["fact"]), float(b["costo"]), None if pl is None else pl / 100,
                None if pl is None else pl - TARGET_COSTO_LOG, _v(b["costo_caja"]), _v(b["pct_fuera"]) / 100 if _v(b["pct_fuera"]) is not None else None,
                float(b["exceso"])]

    def _tabla_x(ws, fila, titulo, titulo_col, df, etq_fn, con_total=True):
        ws.cell(row=fila, column=1, value=titulo).font = Font(name=FN, size=11, bold=True, color="2B343B")
        fila += 1
        for j, h in enumerate([titulo_col] + heads[1:], start=1):
            c = ws.cell(row=fila, column=j, value=h)
            c.font, c.fill, c.border = f_h, fill("2B343B"), borde
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        filas = [_fila_x(r, etq_fn(r["K"])) for _, r in df.iterrows()] + ([_fila_x(T, "TOTAL")] if con_total else [])
        for k, vals in enumerate(filas):
            es_tot = con_total and k == len(filas) - 1
            for j, (v, fm) in enumerate(zip(vals, fmt_g), start=1):
                c = ws.cell(row=fila + 1 + k, column=j, value=v)
                c.font, c.border = (f_b if es_tot else f_n), borde
                if es_tot:
                    c.fill = fill("E3EAEE")
                if fm:
                    c.number_format = fm
        return fila + len(filas) + 3

    # KPIs
    wr["A5"], wr["A5"].font = "INDICADORES", Font(name=FN, size=11, bold=True, color="2B343B")
    kpis = [("Pedidos (facturas)", T["n"], "#,##0"), ("Cajas", T["cajas"], "#,##0"), ("Valor facturado", T["fact"], MONEY),
            (f"Costo logístico ({base_txt})", T["costo"], MONEY), ("% logístico", None if pd.isna(T["pct_log"]) else T["pct_log"] / 100, PCT_F),
            (f"Target % logístico", TARGET_COSTO_LOG / 100, "0.0%"), ("PERDIMOS vs target ($)", P["exceso_total"], MONEY),
            ("Pedidos fuera de target", P["n_fuera"], "#,##0"), ("Pedidos con costo ≥ factura", P["n_perdida"], "#,##0"),
            ("Pedidos sin facturación", P["n_sin_fact"], "#,##0"), ("Costo por caja", _v(T["costo_caja"]), MONEY),
            ("Ticket promedio por pedido", _v(T["ticket"]), MONEY), ("Ticket mínimo para estar en target", _v(T["ticket_min"]), MONEY)]
    for k, (lab, val, fm) in enumerate(kpis):
        a, b = wr.cell(row=6 + k, column=1, value=lab), wr.cell(row=6 + k, column=2, value=_v(float(val)) if val is not None else None)
        a.font, b.font, a.border, b.border = f_n, f_b, borde, borde
        b.number_format = fm
    fila = 6 + len(kpis) + 2
    fila = _tabla_x(wr, fila, "POR NÚMERO DE CAJAS", "NÚMERO DE CAJAS", P["por_cajas"], lambda k: _etq_cajas(k))
    fila = _tabla_x(wr, fila, "POR TRANSPORTE", "TRANSPORTE", P["por_transp"], str)
    if len(P["por_mod"]) > 1:
        fila = _tabla_x(wr, fila, "POR MODALIDAD (COBRO REGRESO / COBRO DESTINO)", "MODALIDAD", P["por_mod"], str)
    if len(P["por_mes"]) > 0:
        fila = _tabla_x(wr, fila, "POR MES", "MES", P["por_mes"], str)
    notas = ["NOTAS",
             f"· FACTURA = columna '{P['col_factura']}' de la matriz. Costo logístico = {base_txt}. % logístico = costo / valor de factura.",
             f"· PERDIMOS = lo pagado por encima del target = máx(costo − {TARGET_COSTO_LOG}% × valor de factura, 0). Si no hay factura, todo el costo cuenta como pérdida.",
             f"· ESTADO: SALUDABLE hasta {TARGET_COSTO_LOG}% · EN ALERTA más de {TARGET_COSTO_LOG}% · CRÍTICO más de {2 * TARGET_COSTO_LOG:g}% · PÉRDIDA si costo ≥ factura · SIN FACTURACIÓN.",
             f"· Las tablas de resumen son valores calculados al generar el archivo con target {TARGET_COSTO_LOG}%. La hoja 'Detalle' es la hoja viva (fórmulas y filtros).",
             "· Periodo según la columna MES de la matriz y el año; las cajas salen de CANTIDAD DE CAJAS (o CAJAS si viene en 0).",
             "· Los registros con concepto de recolecciones o maniobras no son pedidos y se excluyen"
             + (f" ({num(P['extras']['n'])} registros, {money(P['extras']['costo'], 2)})." if P["extras"]["n"] else ".")]
    for k, t_ in enumerate(notas):
        c = wr.cell(row=fila + k, column=1, value=t_)
        c.font = f_b if k == 0 else f_s
    wr.column_dimensions["A"].width = 42
    for j in range(2, 11):
        wr.column_dimensions[get_column_letter(j)].width = 17
    wr.sheet_view.showGridLines = False

    # hojas largas: destino y cliente (lista completa)
    for nombre, clave, col_t, tit in (("Por destino", "por_dest", "DESTINO", "POR DESTINO (mayor pérdida primero)"),
                                      ("Por cliente", "por_cli", "CLIENTE", "POR CLIENTE (mayor pérdida primero)")):
        w = wb.create_sheet(nombre)
        _tabla_x(w, 1, tit, col_t, P[clave], str)
        w.column_dimensions["A"].width = 44
        for j in range(2, 11):
            w.column_dimensions[get_column_letter(j)].width = 17
        w.freeze_panes = "B3"
        w.sheet_view.showGridLines = False

    from openpyxl.worksheet.properties import PageSetupProperties
    for w_ in wb.worksheets:                # impresión: horizontal y a un ancho de página
        w_.page_setup.orientation = "landscape"
        w_.page_setup.fitToWidth, w_.page_setup.fitToHeight = 1, 0
        w_.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    wb.calculation.fullCalcOnLoad = True    # Excel recalcula las fórmulas al abrir
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------- Interfaz de este reporte ----------------
def _ui_detalle_pequenos(df_base, anios, mes_prev, idx_anio):
    c1, c2, c3, c4 = st.columns(4, vertical_alignment="bottom")
    mes_sel = c1.selectbox("MES", ["TODO EL AÑO"] + MESES, index=mes_prev, key="det_mes")
    anio_sel = int(c2.selectbox("AÑO", anios, index=idx_anio, key="det_anio"))
    modalidad = c3.selectbox("MODALIDAD (COBRO REGRESO / COBRO DESTINO)", ["TODAS", "COBRO REGRESO", "COBRO DESTINO"], key="det_mod")
    costo_sel = c4.selectbox("COSTO A CONSIDERAR", ["Guía + adicionales", "Solo guía"], key="det_costo")

    opc_dest, opc_trans = _opciones_detalle(df_base, anio_sel)
    d1, d2, d3, d4 = st.columns([1, 1.4, 1.4, 1], vertical_alignment="bottom")
    cajas_sel = d1.multiselect("NÚMERO DE CAJAS", CAJAS_DETALLE_OPC, default=CAJAS_DETALLE_DEF, key="det_cajas")
    dest_sel = d2.multiselect("DESTINO (vacío = todos)", opc_dest, key=f"det_dest_{anio_sel}")
    trans_sel = d3.multiselect("TRANSPORTE (vacío = todos)", opc_trans, key=f"det_trans_{anio_sel}")
    orden = d4.selectbox("ORDENAR POR", ORDENES_DETALLE, key="det_orden")
    solo_fuera = st.checkbox("Solo pedidos fuera de target (en alerta, críticos o en pérdida)", key="det_fuera")

    st.caption("Lista pedido por pedido (factura, cliente, destino, transporte, valor de factura, costo, % logístico, variación contra el target y "
               "cuánto perdimos) de los pedidos de 1 a 4 cajas (o las que elijas), agrupados por número de cajas. Sale en PDF horizontal o en Excel con filtros.")
    if not cajas_sel:
        st.warning("Elige al menos un número de cajas.")
        return

    mes_num = 0 if mes_sel == "TODO EL AÑO" else MESES.index(mes_sel) + 1
    incluir_adic = (costo_sel == "Guía + adicionales")
    kw = dict(incluir_adic=incluir_adic, cajas=cajas_sel, destinos=dest_sel, transportes=trans_sel, solo_fuera=solo_fuera, orden=orden)
    vista = calc_detalle_pequenos(df_base, anio_sel, mes_num, modalidad, **kw)
    if vista["vacio"]:
        st.warning(vista["motivo"])
        return

    b1, b2 = st.columns(2)
    gen_pdf = b1.button("GENERAR PDF", use_container_width=True, key="det_btn_pdf")
    gen_xls = b2.button("GENERAR EXCEL", use_container_width=True, key="det_btn_xls")
    per = "ANUAL" if mes_num == 0 else mes_sel
    base_nom = f"Detalle_Pedidos_Pequenos_{min(cajas_sel)}a{max(cajas_sel)}_Cajas_{modalidad.replace(' ', '_')}_{per}_{anio_sel}"

    if gen_pdf:
        with st.spinner("Armando el PDF..."):
            try:
                pdf = generar_detalle_pequenos_pdf(df_base, anio_sel, mes_num, modalidad, **kw)
                st.session_state["rep_det_pdf"] = {"bytes": pdf.getvalue(), "nombre": base_nom + ".pdf"}
                st.success("¡PDF generado!")
            except Exception as e:
                st.error(f"No se pudo generar el PDF: {e}")
    if gen_xls:
        with st.spinner("Armando el Excel..."):
            try:
                xls = generar_detalle_pequenos_excel(df_base, anio_sel, mes_num, modalidad, **kw)
                st.session_state["rep_det_xlsx"] = {"bytes": xls.getvalue(), "nombre": base_nom + ".xlsx"}
                st.success("¡Excel generado!")
            except Exception as e:
                st.error(f"No se pudo generar el Excel: {e}")

    r_pdf, r_xls = st.session_state.get("rep_det_pdf"), st.session_state.get("rep_det_xlsx")
    if r_pdf or r_xls:
        x1, x2 = st.columns(2)
        if r_pdf:
            x1.download_button("DESCARGAR PDF", data=r_pdf["bytes"], file_name=r_pdf["nombre"], mime="application/pdf",
                               use_container_width=True, key="rep_det_pdf_dl")
        if r_xls:
            x2.download_button("DESCARGAR EXCEL", data=r_xls["bytes"], file_name=r_xls["nombre"],
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               use_container_width=True, key="rep_det_xlsx_dl")


# ============================================================
# 13. INTERFAZ (elegir reporte, periodo y generar)
# ============================================================
REPORTE_MENSUAL = "REPORTE MENSUAL COMPLETO"
REPORTE_CONCEPTO = "% LOGÍSTICO POR CONCEPTO"
REPORTE_PEQUENOS = "SALUD DE PEDIDOS PEQUEÑOS (1 A 4 CAJAS)"
REPORTE_DETALLE = "DETALLE PEDIDO POR PEDIDO (1 A 4 CAJAS) - PDF / EXCEL"


@st.cache_data(ttl=300)
def mp_datos_base():
    """Base preparada del reporte mensual (en caché para no recalcularla en cada clic)."""
    return mp_preparar_base(cargar_matriz())


def mp_ui_predefinidos():
    try:
        import pytz
        hoy = pd.Timestamp(datetime.now(pytz.timezone("America/Mexico_City")).date())
    except Exception:
        hoy = pd.Timestamp(datetime.now().date())

    with st.spinner("Cargando base de envíos..."):
        try:
            df_base = mp_datos_base()
        except Exception as e:
            st.error(f"No se pudo cargar la matriz de envíos: {e}")
            return

    anios = sorted({int(y) for y in df_base["FECHA DE ENVÍO"].dt.year.dropna().unique()} | {hoy.year}, reverse=True)
    mes_prev = hoy.month - 1 or 12
    anio_prev = hoy.year if hoy.month > 1 else hoy.year - 1
    idx_anio = anios.index(anio_prev) if anio_prev in anios else 0

    tipo = st.selectbox("¿QUÉ REPORTE QUIERES GENERAR?", [REPORTE_MENSUAL, REPORTE_CONCEPTO, REPORTE_PEQUENOS, REPORTE_DETALLE])
    st.caption("Los registros con CONCEPTO de recolecciones o maniobras no se cuentan como pedidos: se reportan aparte como costos extras "
               "y se suman al costo logístico total.")

    # ---------------- REPORTE MENSUAL COMPLETO ----------------
    if tipo == REPORTE_MENSUAL:
        c1, c2, c3 = st.columns([1.3, 1, 1.6], vertical_alignment="bottom")
        mes_sel = c1.selectbox("MES", MESES, index=mes_prev - 1, key="mens_mes")
        anio_sel = c2.selectbox("AÑO", anios, index=idx_anio, key="mens_anio")
        generar = c3.button("GENERAR REPORTE PDF", use_container_width=True, key="mens_btn")

        st.caption("Incluye: resumen del mes, efectividad de envíos, inteligencia de negocio, top 20 clientes, distribución de carga (cobro regreso), "
                   "costo por fletera en cobro regreso, ranking de fleteras (efectividad, tiempos y costo promedio) y costos de muestras.")

        if generar:
            mes_num = MESES.index(mes_sel) + 1
            with st.spinner("Armando el reporte..."):
                df_muestras, msg_muestras = None, ""
                if MUESTRAS_OK:
                    try:
                        df_m, _sha = obtener_datos_github()
                        if df_m is not None and not df_m.empty:
                            df_muestras = preparar_muestras(df_m)
                        else:
                            msg_muestras = "No hay registros de muestras cargados."
                    except Exception as e:
                        msg_muestras = f"No se pudo leer el registro de muestras: {e}"
                else:
                    msg_muestras = f"No se pudo importar muestras_common: {MUESTRAS_ERR}"
                try:
                    pdf = generar_reporte_pdf(df_base, df_muestras, int(anio_sel), mes_num, hoy,
                                              precios=PRECIOS_MUESTRAS, logo_bytes=mp_obtener_logo_bytes(), muestras_msg=msg_muestras,
                                              df_hist=cargar_historial_2025())
                    st.session_state["rep_mensual_pdf"] = {"bytes": pdf.getvalue(), "nombre": f"Reporte_Mensual_Logistica_{mes_sel}_{anio_sel}.pdf"}
                    if msg_muestras:
                        st.warning(f"Reporte generado, pero la sección de muestras quedó sin datos: {msg_muestras}")
                    else:
                        st.success("¡Reporte generado!")
                except Exception as e:
                    st.error(f"No se pudo generar el reporte: {e}")

        rep = st.session_state.get("rep_mensual_pdf")
        if rep:
            st.download_button("DESCARGAR REPORTE PDF", data=rep["bytes"], file_name=rep["nombre"], mime="application/pdf",
                               use_container_width=True, key="rep_mensual_dl")

    # ---------------- % LOGÍSTICO POR CONCEPTO ----------------
    elif tipo == REPORTE_CONCEPTO:
        c1, c2, c3, c4 = st.columns(4, vertical_alignment="bottom")
        mes_sel = c1.selectbox("MES", ["TODO EL AÑO"] + MESES, index=mes_prev, key="conc_mes")   # +1 por "TODO EL AÑO"
        anio_sel = c2.selectbox("AÑO", anios, index=idx_anio, key="conc_anio")
        modalidad = c3.selectbox("MODALIDAD", ["COBRO REGRESO", "COBRO DESTINO", "TODAS"], key="conc_mod")
        costo_sel = c4.selectbox("COSTO A CONSIDERAR", ["Guía + adicionales", "Solo guía"], key="conc_costo")
        generar = st.button("GENERAR REPORTE PDF", use_container_width=True, key="conc_btn")

        st.caption("Calcula el % logístico (costo / facturación) de cada concepto, con facturación, costo de guía, adicionales, cajas, "
                   f"costo por caja y comparación contra el target de {TARGET_COSTO_LOG}%. 'Guía + adicionales' es como Análisis Mensual; "
                   "'Solo guía' es como la sección 03 del reporte mensual.")

        if generar:
            mes_num = 0 if mes_sel == "TODO EL AÑO" else MESES.index(mes_sel) + 1
            with st.spinner("Armando el reporte..."):
                try:
                    pdf = generar_reporte_concepto_pdf(df_base, int(anio_sel), mes_num, modalidad,
                                                       incluir_adic=(costo_sel == "Guía + adicionales"))
                    per = "ANUAL" if mes_num == 0 else mes_sel
                    st.session_state["rep_log_concepto_pdf"] = {
                        "bytes": pdf.getvalue(),
                        "nombre": f"Porcentaje_Logistico_por_Concepto_{modalidad.replace(' ', '_')}_{per}_{anio_sel}.pdf"}
                    st.success("¡Reporte generado!")
                except Exception as e:
                    st.error(f"No se pudo generar el reporte: {e}")

        rep = st.session_state.get("rep_log_concepto_pdf")
        if rep:
            st.download_button("DESCARGAR REPORTE PDF", data=rep["bytes"], file_name=rep["nombre"], mime="application/pdf",
                               use_container_width=True, key="rep_log_concepto_dl")

    # ---------------- SALUD DE PEDIDOS PEQUEÑOS ----------------
    elif tipo == REPORTE_PEQUENOS:
        c1, c2, c3, c4, c5 = st.columns(5, vertical_alignment="bottom")
        mes_sel = c1.selectbox("MES", ["TODO EL AÑO"] + MESES, index=mes_prev, key="peq_mes")
        anio_sel = c2.selectbox("AÑO", anios, index=idx_anio, key="peq_anio")
        modalidad = c3.selectbox("MODALIDAD", ["TODAS", "COBRO REGRESO", "COBRO DESTINO"], key="peq_mod")
        costo_sel = c4.selectbox("COSTO A CONSIDERAR", ["Guía + adicionales", "Solo guía"], key="peq_costo")
        max_cj = c5.selectbox("PEDIDO PEQUEÑO = HASTA (CAJAS)", [2, 3, 4, 5, 6, 8, 10], index=[2, 3, 4, 5, 6, 8, 10].index(CAJAS_PEQUENO_MAX), key="peq_max")
        generar = st.button("GENERAR REPORTE PDF", use_container_width=True, key="peq_btn")

        st.caption("Salud de los pedidos pequeños: facturación vs costo de flete (guía) y de distribución (adicionales), comparativo contra pedidos grandes, "
                   "semáforo por pedido, fletera, ticket mínimo para estar en target, clientes/destinos con mayor sobrecosto y tendencia del año.")

        if generar:
            mes_num = 0 if mes_sel == "TODO EL AÑO" else MESES.index(mes_sel) + 1
            with st.spinner("Armando el reporte..."):
                try:
                    pdf = generar_reporte_pequenos_pdf(df_base, int(anio_sel), mes_num, modalidad,
                                                       incluir_adic=(costo_sel == "Guía + adicionales"), max_cajas=int(max_cj))
                    per = "ANUAL" if mes_num == 0 else mes_sel
                    st.session_state["rep_pequenos_pdf"] = {
                        "bytes": pdf.getvalue(),
                        "nombre": f"Salud_Pedidos_Pequenos_1a{int(max_cj)}_Cajas_{modalidad.replace(' ', '_')}_{per}_{anio_sel}.pdf"}
                    st.success("¡Reporte generado!")
                except Exception as e:
                    st.error(f"No se pudo generar el reporte: {e}")

        rep = st.session_state.get("rep_pequenos_pdf")
        if rep:
            st.download_button("DESCARGAR REPORTE PDF", data=rep["bytes"], file_name=rep["nombre"], mime="application/pdf",
                               use_container_width=True, key="rep_pequenos_dl")

    # ---------------- DETALLE PEDIDO POR PEDIDO (PDF / EXCEL) ----------------
    else:
        _ui_detalle_pequenos(df_base, anios, mes_prev, idx_anio)



# ============================================================
# 14. REPORTEADOR: ESTILO GLOBAL Y PESTAÑAS
# ============================================================
CSS_REPORTEADOR = """<style>
div.stButton > button, div.stDownloadButton > button{background-color:#628290!important;color:#fff!important;border:1px solid #628290!important;
border-radius:6px!important;transition:all .25s ease!important;width:100%!important;box-shadow:none!important;font-weight:700!important}
div.stButton > button:hover, div.stDownloadButton > button:hover{background-color:#4E6772!important;border-color:#4E6772!important;color:#fff!important}
div.stButton > button:focus, div.stDownloadButton > button:focus, div.stButton > button:active, div.stDownloadButton > button:active{
background-color:#4E6772!important;border-color:#4E6772!important;color:#fff!important;box-shadow:none!important}
div.stButton > button p, div.stDownloadButton > button p{color:inherit!important}
div[data-testid="stButton"], div[data-testid="stDownloadButton"], div.stButton, div.stDownloadButton{width:100%!important}
div[data-testid="stElementContainer"]:has(> div.stButton), div[data-testid="stElementContainer"]:has(> div[data-testid="stButton"]), div[data-testid="stElementContainer"]:has(> div[data-testid="stDownloadButton"]){width:100%!important}
div[data-testid="stButton"] > button, div[data-testid="stDownloadButton"] > button{width:100%!important;min-height:38px!important;height:38px!important;padding-top:0!important;padding-bottom:0!important;font-size:13px!important}
button[data-baseweb="tab"] p{font-weight:800!important;letter-spacing:1.5px!important;font-size:13px!important}
button[data-baseweb="tab"][aria-selected="true"] p{color:#00A3A3!important}
div[data-baseweb="tab-highlight"]{background-color:#00A3A3!important}
</style>"""


def main():
    st.markdown(CSS_REPORTEADOR, unsafe_allow_html=True)
    st.markdown("<div style='padding:6px 0 14px 0;border-bottom:1px solid rgba(255,255,255,0.08);margin-bottom:18px;'>"
                "<span style='color:#FFFFFF;font-size:13px;font-weight:800;letter-spacing:2.5px;text-transform:uppercase;'>"
                "REPORTES // REPORTEADOR</span></div>", unsafe_allow_html=True)
    tab_din, tab_pre = st.tabs(["REPORTES DINÁMICOS", "REPORTES PREDEFINIDOS"])
    with tab_din:
        ui_dinamicos()
    with tab_pre:
        mp_ui_predefinidos()


if __name__ == "__main__":
    main()
